"""ToolHub - single-file Flask app (website + backend + admin panel).
Env vars: MONGO_URI, GEMINI_API_KEY, ADMIN_PASSWORD
Optional: SECRET_KEY, GEMINI_MODEL, GEMINI_IMAGE_MODEL, MONGO_DB
Install: pip install flask "pymongo[srv]" requests gunicorn
Run: python app.py   |   Production: gunicorn app:app --timeout 120
"""
import os, io, re, json, time, uuid, hmac, base64, hashlib, socket, ipaddress
from urllib.parse import urljoin, urlparse
import requests
from flask import Flask, request, jsonify, Response, session
from werkzeug.exceptions import HTTPException

from urllib3.util import connection as _conn

_orig_cc = _conn.create_connection


def _safe_cc(address, *a, **k):
    """Only allow outgoing connections to public internet addresses (blocks localhost/private networks)."""
    host, port = address[0], address[1]
    last = OSError("Blocked address")
    for _f, _t, _p, _c, sa in socket.getaddrinfo(host, port, 0, socket.SOCK_STREAM):
        if not ipaddress.ip_address(sa[0].split("%")[0]).is_global:
            continue
        try:
            return _orig_cc((sa[0], port), *a, **k)
        except OSError as e:
            last = e
    raise last


_conn.create_connection = _safe_cc

try:
    from dotenv import load_dotenv

    load_dotenv()  # optional: reads a local .env file
except ImportError:
    pass

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 12 * 1024 * 1024  # 12 MB uploads
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = bool(os.environ.get("RENDER"))
app.secret_key = os.environ.get("SECRET_KEY") or hashlib.sha256(
    ("toolhub" + os.environ.get("ADMIN_PASSWORD", "")).encode()).hexdigest()

TEXT_MODEL = os.environ.get("GEMINI_MODEL", "gemini-2.5-flash")
IMAGE_MODEL = os.environ.get("GEMINI_IMAGE_MODEL", "gemini-2.5-flash-image")
_db = None


class Err(Exception):
    def __init__(self, msg, code=400):
        self.msg, self.code = msg, code


def get_db():
    global _db
    if _db is None:
        uri = os.environ.get("MONGO_URI")
        if not uri:
            return None
        from pymongo import MongoClient
        _db = MongoClient(uri, serverSelectionTimeoutMS=5000)[os.environ.get("MONGO_DB", "toolhub")]
    return _db


def need_db():
    db = get_db()
    if db is None:
        raise Err("MONGO_URI is not set on the server", 503)
    return db


hits = {}  # small per-IP limits to protect your Gemini quota


def limit(kind, n):
    ip = (request.headers.get("X-Forwarded-For") or request.remote_addr or "x").split(",")[0].strip()
    now = time.time()
    h = hits[(kind, ip)] = [t for t in hits.get((kind, ip), []) if now - t < 3600]
    if len(h) >= n:
        raise Err("Limit reached (%d per hour). Try again later." % n, 429)
    h.append(now)


def gemini(model, parts, cfg=None, tools=None):
    key = os.environ.get("GEMINI_API_KEY")
    if not key:
        raise Err("Server has no GEMINI_API_KEY set", 500)
    body = {"contents": [{"parts": parts}]}
    if cfg:
        body["generationConfig"] = cfg
    if tools:
        body["tools"] = tools
    r = requests.post("https://generativelanguage.googleapis.com/v1beta/models/%s:generateContent" % model,
                      headers={"x-goog-api-key": key, "Content-Type": "application/json"}, json=body, timeout=120)
    try:
        j = r.json()
    except ValueError:
        j = {}
    if r.status_code != 200:
        raise Err("Gemini: " + str(j.get("error", {}).get("message", "request failed"))[:160], 502)
    cands = j.get("candidates") or [{}]
    return (cands[0].get("content") or {}).get("parts") or []


# ---------- public pages & stats ----------
@app.get("/")
def home():
    return Response(HTML, mimetype="text/html")


@app.get("/health")
def health():
    e = os.environ.get
    return jsonify(ok=True, mongo=bool(e("MONGO_URI")), gemini=bool(e("GEMINI_API_KEY")), admin=bool(e("ADMIN_PASSWORD")))


@app.post("/api/visit")
def visit():
    db = need_db()
    vid = request.cookies.get("vid") or uuid.uuid4().hex
    db.visitors.update_one({"_id": vid}, {"$setOnInsert": {"first": time.time()}, "$inc": {"visits": 1}}, upsert=True)
    db.stats.update_one({"_id": "views"}, {"$inc": {"total": 1}}, upsert=True)
    total = (db.stats.find_one({"_id": "views"}) or {}).get("total", 0)
    resp = jsonify(unique=db.visitors.count_documents({}), total=total)
    resp.set_cookie("vid", vid, max_age=31536000, httponly=True, samesite="Lax")
    return resp


# ---------- site config (editable from the admin panel) ----------
DEFAULT = {"name": "ToolHub", "c1": "#2563eb", "c2": "#7c3aed",
           "tk": ["Welcome to ToolHub", "New: AI Photo Generator & Prompt Generator", "Question Paper Solver is live"],
           "lk": [{"t": "Home", "u": "#"}], "ov": {}, "ct": [], "logo": "/logo.png", "tag": "All-in-one tools for everyone"}


def s(v, n=80):
    return str(v if v is not None else "")[:n]


def lst(v):
    return v if isinstance(v, list) else []


def url_ok(u):
    return u if re.match(r"^(https?://\S+|/f/[0-9a-f]{24}|/logo\.png|#)$", u) else "#"


def hexc(c, d):
    return c if re.match(r"^#[0-9a-fA-F]{6}$", c) else d


def grad(g):
    return g if re.match(r"^#[0-9a-fA-F]{6},#[0-9a-fA-F]{6}$", g) else "#bfdbfe,#3b82f6"


def clean(d):
    d = d if isinstance(d, dict) else {}
    ov = {}
    for k, v in (d["ov"] if isinstance(d.get("ov"), dict) else {}).items():
        if re.match(r"^\w{1,12}$", k) and isinstance(v, dict):
            ov[k] = {"n": s(v.get("n"), 40), "d": s(v.get("d"), 90), "i": s(v.get("i"), 8),
                     "c": s(v.get("c"), 20), "hide": bool(v.get("hide"))}
    ct = []
    for t in lst(d.get("ct"))[:60]:
        if isinstance(t, dict):
            ct.append({"id": re.sub(r"\W", "", s(t.get("id"), 20)) or uuid.uuid4().hex[:8],
                       "n": s(t.get("n"), 40) or "Tile", "d": s(t.get("d"), 90), "i": s(t.get("i"), 8) or "🔗",
                       "c": s(t.get("c"), 20) or "Links", "g": grad(s(t.get("g"), 20)), "h": url_ok(s(t.get("h"), 300))})
    lg = s(d.get("logo", "/logo.png"), 40)
    lk = [{"t": s(l.get("t"), 30), "u": url_ok(s(l.get("u"), 300))} for l in lst(d.get("lk"))[:12] if isinstance(l, dict)]
    return {"name": s(d.get("name"), 40) or "ToolHub", "c1": hexc(s(d.get("c1"), 7), DEFAULT["c1"]),
            "c2": hexc(s(d.get("c2"), 7), DEFAULT["c2"]), "tk": [s(x, 150) for x in lst(d.get("tk"))[:20]],
            "lk": lk, "ov": ov, "ct": ct, "tag": s(d.get("tag"), 80).strip() or DEFAULT["tag"],
            "logo": lg if re.match(r"^(/f/[0-9a-f]{24}|/logo\.png)$", lg) else ""}


@app.get("/api/config")
def get_config():
    return jsonify(cfg_doc())


@app.get("/logo.png")
def default_logo():
    return Response(base64.b64decode(LOGO_B64), mimetype="image/png", headers={"Cache-Control": "public, max-age=86400"})


# ---------- admin ----------
def need_admin():
    if not session.get("admin"):
        raise Err("Not logged in", 401)


@app.post("/api/admin/login")
def login():
    pw = os.environ.get("ADMIN_PASSWORD")
    if not pw:
        raise Err("Set ADMIN_PASSWORD on the server first", 503)
    limit("login", 10)
    given = str((request.get_json(silent=True) or {}).get("pw", ""))
    if not hmac.compare_digest(given.encode(), pw.encode()):
        raise Err("Wrong password", 403)
    session["admin"] = True
    session.permanent = True
    return jsonify(ok=True)


@app.get("/api/admin/me")
def me():
    need_admin()
    return jsonify(ok=True)


@app.post("/api/admin/logout")
def logout():
    session.clear()
    return jsonify(ok=True)


@app.post("/api/admin/config")
def save_config():
    need_admin()
    db = need_db()
    cfg = clean(request.get_json(silent=True))
    db.config.replace_one({"_id": "site"}, {**cfg, "_id": "site"}, upsert=True)
    refs = set(re.findall(r"/f/([0-9a-f]{24})", repr((cfg["ct"], cfg["lk"], cfg["logo"]))))
    for f in list(db.files.find({}, {"_id": 1})):  # delete PDFs no tile uses any more
        if str(f["_id"]) not in refs:
            db.files.delete_one({"_id": f["_id"]})
    return jsonify(cfg)


@app.post("/api/admin/upload")
def upload():
    need_admin()
    db = need_db()
    f = request.files.get("file")
    data = f.read() if f else b""
    if not data.startswith(b"%PDF-"):
        raise Err("Please upload a PDF file", 400)
    from bson import Binary
    r = db.files.insert_one({"data": Binary(data), "name": s(f.filename, 80), "t": time.time()})
    return jsonify(url="/f/" + str(r.inserted_id))


@app.post("/api/admin/logo")
def upload_logo():
    need_admin()
    db = need_db()
    f = request.files.get("file")
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(f.read() if f else b""))
        im.load()
        if im.format not in ("PNG", "JPEG", "WEBP"):
            raise ValueError
    except Exception:
        raise Err("Please upload a PNG, JPG or WebP image")
    im = im.convert("RGBA")
    im.thumbnail((512, 512))
    bio = io.BytesIO()
    im.save(bio, "PNG", optimize=True)
    from bson import Binary
    r = db.files.insert_one({"data": Binary(bio.getvalue()), "name": "logo.png", "mime": "image/png", "t": time.time()})
    return jsonify(url="/f/" + str(r.inserted_id))


@app.get("/f/<fid>")
def pdf(fid):
    from bson import ObjectId
    db = need_db()
    if not re.match(r"^[0-9a-f]{24}$", fid):
        raise Err("Not found", 404)
    d = db.files.find_one({"_id": ObjectId(fid)})
    if not d:
        raise Err("Not found", 404)
    nm = re.sub(r"[^\w.-]", "_", d.get("name") or "file.pdf")
    return Response(bytes(d["data"]), mimetype=d.get("mime") or "application/pdf", headers={
        "Content-Disposition": 'inline; filename="%s"' % nm, "X-Content-Type-Options": "nosniff",
        "Cache-Control": "public, max-age=3600"})


# ---------- AI tools (Gemini) ----------
SOLVE = ("This is a question paper. First detect its language. Then answer EVERY question and give a clear "
         "step-by-step explanation. Write the answers and explanations in the SAME language as the question paper "
         "(Telugu paper -> Telugu, Hindi -> Hindi, English -> English). Keep the original question numbers. "
         "For maths show the steps. If something is unreadable, say so instead of guessing. "
         "Use plain text only: no markdown symbols like ** or #.")


@app.post("/api/solve")
def solve():
    limit("solve", 10)
    f = request.files.get("file")
    if not f:
        raise Err("No file uploaded")
    mt = f.mimetype
    if mt not in ("application/pdf", "image/jpeg", "image/png", "image/webp"):
        raise Err("Upload a PDF or an image (JPG/PNG)")
    data = base64.b64encode(f.read()).decode()
    parts = gemini(TEXT_MODEL, [{"inline_data": {"mime_type": mt, "data": data}}, {"text": SOLVE}])
    return jsonify(answer="".join(p.get("text", "") for p in parts).strip() or "No answer returned")


# ---------- Response Sheet Calculator (no AI: regex + PyPDF2 + BeautifulSoup) ----------
Q_RE = re.compile(r"Question\s*ID\s*[:\-]?\s*(\d+)", re.I)
_VAL = r"(--+|-?\d+(?:\.\d+)?(?:\s*,\s*\d+)*|[A-Da-d](?:\s*,\s*[A-Da-d])*(?![A-Za-z])|—|–|-)"
CH_RE = re.compile(r"(?:Chosen|Given|Marked)\s*(?:Option|Answer)\s*[:\-]?\s*" + _VAL, re.I)
CO_RE = re.compile(r"(?:Correct|Right)\s*(?:Option|Answer)\s*[:\-]?\s*" + _VAL, re.I)
NAME_RE = re.compile(r"(?:Candidate|Applicant|Student)\s*Name\s*[:\-]?[ \t]*\n?[ \t]*([A-Za-z][A-Za-z .'\-]{1,50})", re.I)
HDR1 = re.compile(r"^(?:Subject|Section|Part)\s*[:\-]\s*(.*)$", re.I)
HDR0 = re.compile(r"^(?:Subject|Section|Part)\s*[:\-]?$", re.I)
NOISE = set("""QUESTION QUESTIONS STATUS ANSWERED ANSWER ANSWERS OPTION OPTIONS CHOSEN CORRECT WRONG RESPONSE RESPONSES
SHEET CANDIDATE ROLL APPLICATION CENTRE CENTER INSTRUCTION INSTRUCTIONS NATIONAL TESTING AGENCY EXAM EXAMINATION TOTAL
SCORE MARK MARKS KEY PAGE DATE TIME SHIFT SESSION NAME NO NUMBER FORM TEST TYPE ID SECTION SUBJECT PART MCQ NOT ATTEMPTED
UNANSWERED PHOTO SIGNATURE SUMMARY RESULT DETAILS""".split())


def norm_text(t):
    t = t.replace("\xa0", " ").replace("\r", "\n")
    return re.sub(r"[ \t\f\v]+", " ", t)


def fetch(url, maxb=8 * 1024 * 1024):
    """Download a public link (HTML or PDF). Redirects are followed manually; private addresses are blocked."""
    sess = requests.Session()
    sess.trust_env = False
    for _ in range(4):
        p = urlparse(url)
        if p.scheme not in ("http", "https") or not p.hostname or p.port not in (None, 80, 443, 8080, 8443):
            raise Err("Please paste a valid http(s) link")
        try:
            r = sess.get(url, stream=True, timeout=(6, 15), allow_redirects=False,
                         headers={"User-Agent": "Mozilla/5.0 (compatible; ToolHub)"})
        except requests.RequestException:
            raise Err("Could not open the link. Make sure it is public (no login needed).", 422)
        if r.status_code in (301, 302, 303, 307, 308) and r.headers.get("Location"):
            url = urljoin(url, r.headers["Location"])
            r.close()
            continue
        if r.status_code != 200:
            raise Err("The link returned error %d" % r.status_code, 422)
        buf = b""
        for ch in r.iter_content(65536):
            buf += ch
            if len(buf) > maxb:
                raise Err("That page is too large (max 8 MB)", 422)
        return buf, url
    raise Err("Too many redirects", 422)


def pdf_text(data):
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        raise Err("Server is missing PyPDF2 (add it to requirements)", 500)
    try:
        rd = PdfReader(io.BytesIO(data))
        if rd.is_encrypted:
            rd.decrypt("")
        return "\n".join((pg.extract_text() or "") for pg in rd.pages[:300])
    except Exception:
        raise Err("Could not read this PDF", 422)


def html_text(data):
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        raise Err("Server is missing beautifulsoup4 (add it to requirements)", 500)
    soup = BeautifulSoup(data, "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    return soup.get_text("\n")


def clean_name(n):
    n = re.sub(r"\s+", " ", n).strip(" :-–|")
    n = re.sub(r"\s*[-–(]?\s*Section\s*[A-Za-z0-9]+\)?$", "", n, flags=re.I).strip(" :-–|")
    return n[:40]


def find_candidate(text):
    m = NAME_RE.search(text)
    if not m:
        return ""
    n = re.split(r"\s+(?:Roll|Application|Registration|Reg|Test|Date|Exam|Father|Mother|Mobile|Email|Candidate)\b",
                 m.group(1), flags=re.I)[0].strip(" .'-")
    if not n or any(w.upper() in NOISE for w in re.findall(r"[A-Za-z]+", n)):
        return ""
    return n.title() if n.isupper() else n


def find_headers(text, cand):
    """Detect subject headings dynamically: 'Subject: X', 'Section: X', 'Part: X' or standalone heading lines."""
    lines, offs, pos, out = text.split("\n"), [], 0, []
    for ln in lines:
        offs.append(pos)
        pos += len(ln) + 1
    for i, raw in enumerate(lines):
        ln, name = raw.strip(), None
        if not ln or len(ln) > 60:
            continue
        m = HDR1.match(ln)
        if m and m.group(1).strip():
            name = clean_name(m.group(1))
        elif HDR0.match(ln):
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines):
                name = clean_name(lines[j])
        elif re.fullmatch(r"[A-Za-z][A-Za-z &/\-]{1,38}[A-Za-z]", ln):
            words = [w.upper() for w in re.findall(r"[A-Za-z]+", ln)]
            if not any(w in NOISE for w in words):
                if ln.isupper():
                    name = ln
                elif len(words) <= 4 and ln[0].isupper() and re.search(
                        r"Q\.?\s*\d+|Question\s*ID", " ".join(lines[i + 1:i + 6]), re.I):
                    name = ln
        if name and len(re.findall(r"[A-Za-z]", name)) >= 2 and name.lower() != cand.lower() \
                and not any(w.upper() in NOISE for w in re.findall(r"[A-Za-z]+", name)):
            out.append((offs[i], name.title() if name.isupper() else name))
    return out


def is_na(v):
    return v is None or re.fullmatch(r"-+|—|–", v.strip()) is not None


def norm_ans(v):
    out = []
    for t in re.split(r"\s*,\s*", v.strip().upper()):
        if t:
            try:
                out.append(repr(float(t)))
            except ValueError:
                out.append(t)
    return tuple(sorted(out))


def num(v, d):
    try:
        return max(0.0, min(100.0, float(v)))
    except (TypeError, ValueError):
        return d


def parse_sheet(text, pos, neg):
    cand = find_candidate(text)
    qs, hs = list(Q_RE.finditer(text)), find_headers(text, cand)
    subj, seen, ignored = {}, set(), 0
    for i, m in enumerate(qs):
        qid = m.group(1)
        if qid in seen:
            continue
        blk = text[m.end(): qs[i + 1].start() if i + 1 < len(qs) else len(text)]
        ch, co = CH_RE.search(blk), CO_RE.search(blk)
        if not ch and not co:
            ignored += 1
            continue
        seen.add(qid)
        h = [n for o, n in hs if o <= m.start()]
        name = h[-1] if h else "General"
        d = subj.setdefault(name, {"name": name, "total": 0, "correct": 0, "wrong": 0,
                                   "not_attempted": 0, "unverified": 0})
        d["total"] += 1
        c = ch.group(1) if ch else None
        if is_na(c):
            d["not_attempted"] += 1
        elif not co or is_na(co.group(1)):
            d["unverified"] += 1  # answered, but the correct option is not on the sheet
        elif norm_ans(c) == norm_ans(co.group(1)):
            d["correct"] += 1
        else:
            d["wrong"] += 1
    subs = list(subj.values())
    g = {"total": 0, "correct": 0, "wrong": 0, "not_attempted": 0}
    for d in subs:
        d["marks"] = round(d["correct"] * pos - d["wrong"] * neg, 2)
        for k in g:
            g[k] += d[k]
    g["marks"] = round(g["correct"] * pos - g["wrong"] * neg, 2)
    notes = []
    unv = sum(d["unverified"] for d in subs)
    if unv:
        notes.append("%d answered question(s) have no correct option on the sheet, so they are not counted in marks." % unv)
    if subs and list(subj) == ["General"]:
        notes.append("No subject headings were found, so all questions are shown together.")
    return {"candidate": cand, "pos": pos, "neg": neg, "subjects": subs, "grand_total": g, "note": " ".join(notes)}


def to_photo(raw):
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(raw))
        im.load()
        w, h = im.size
        if w < 60 or h < 60 or w > 3000 or not (.5 <= w / h <= 1.3):
            return "", 9
        im = im.convert("RGB")
        im.thumbnail((300, 300))
        bio = io.BytesIO()
        im.save(bio, "JPEG", quality=86)
        return "data:image/jpeg;base64," + base64.b64encode(bio.getvalue()).decode(), abs(w / h - .8)
    except Exception:
        return "", 9


def photo_from_pdf(data):
    best = ("", 9)
    try:
        from PyPDF2 import PdfReader
        for pg in PdfReader(io.BytesIO(data)).pages[:3]:
            for im in getattr(pg, "images", []):
                p = to_photo(im.data)
                if p[0] and p[1] < best[1]:
                    best = p
    except Exception:
        pass
    return best[0]


def photo_from_html(data, base):
    try:
        from bs4 import BeautifulSoup
        cands = []
        for im in BeautifulSoup(data, "html.parser").find_all("img")[:30]:
            src = (im.get("src") or "").strip()
            meta = " ".join([src, im.get("alt") or "", " ".join(im.get("class") or []), im.get("id") or ""]).lower()
            if not src or re.search(r"logo|banner|icon|sprite|header|footer|button|arrow", meta):
                continue
            cands.append((0 if re.search(r"photo|candidate|cand|pic|passport|profile|image", meta) else 1, src))
        cands.sort(key=lambda c: c[0])
        for _s, src in cands[:5]:
            try:
                raw = base64.b64decode(src.split(",", 1)[1]) if src.startswith("data:image") else fetch(urljoin(base, src), 3 * 1024 * 1024)[0]
            except Exception:
                continue
            p = to_photo(raw)[0]
            if p:
                return p
    except Exception:
        pass
    return ""


@app.post("/response-calculator")
def response_calculator():
    limit("calc", 40)
    pos, neg = num(request.form.get("pos"), 4.0), num(request.form.get("neg"), 1.0)
    f, url = request.files.get("file"), s(request.form.get("url"), 600).strip()
    photo = ""
    if f and f.filename:
        data = f.read()
        if not data.startswith(b"%PDF-"):
            raise Err("Please upload a PDF file")
        text, photo = pdf_text(data), photo_from_pdf(data)
    elif url:
        data, final = fetch(url)
        if data.startswith(b"%PDF-"):
            text, photo = pdf_text(data), photo_from_pdf(data)
        else:
            text, photo = html_text(data), photo_from_html(data, final)
    else:
        raise Err("Upload the PDF or paste the link first")
    text = norm_text(text)
    if len(text.strip()) < 20:
        raise Err("No readable text found. A scanned (image-only) PDF cannot be read without OCR.", 422)
    res = parse_sheet(text, pos, neg)
    if not res["subjects"]:
        raise Err("No questions found. The sheet must list 'Question ID', 'Chosen Option' and 'Correct Option'.", 422)
    res["photo"] = photo
    return jsonify(res)


@app.post("/api/prompt")
def prompt_gen():
    limit("prompt", 30)
    j = request.get_json(silent=True) or {}
    idea, style = s(j.get("idea"), 500).strip(), s(j.get("style"), 40)
    if not idea:
        raise Err("Write your idea first")
    ask = ("You are an expert prompt engineer for AI image generators. Turn the idea below into ONE detailed, vivid "
           "image prompt in English (60-120 words), style: %s. Include subject details, setting, lighting, "
           "camera and lens, composition, mood, colour palette and realistic textures. "
           "Output only the prompt text, no quotes, no markdown.\nIdea: %s" % (style, idea))
    parts = gemini(TEXT_MODEL, [{"text": ask}])
    return jsonify(prompt="".join(p.get("text", "") for p in parts).strip())


@app.post("/api/image")
def image_gen():
    limit("image", 10)
    j = request.get_json(silent=True) if request.is_json else None
    p = s((j or request.form).get("prompt"), 1200).strip()
    if not p:
        raise Err("Describe the photo first")
    imgs = []
    for f in request.files.getlist("photos")[:3]:
        if f.mimetype not in ("image/jpeg", "image/png", "image/webp"):
            raise Err("Upload JPG, PNG or WebP photos")
        imgs.append({"inline_data": {"mime_type": f.mimetype, "data": base64.b64encode(f.read()).decode()}})
    if imgs:
        text = ("Use the uploaded photo(s) as the reference. Apply this change while keeping the result photorealistic, "
                "natural and sharp, and keep the people's faces and identity recognizable: " + p)
    else:
        text = "Create a photorealistic, ultra-detailed, natural-looking photograph, sharp focus, realistic lighting. " + p
    for part in gemini(IMAGE_MODEL, imgs + [{"text": text}], {"responseModalities": ["TEXT", "IMAGE"]}):
        d = part.get("inlineData") or part.get("inline_data")
        if d:
            return jsonify(image="data:%s;base64,%s" % (d.get("mimeType") or d.get("mime_type") or "image/png", d["data"]))
    raise Err("No image returned. The prompt may have been blocked, or image generation needs billing enabled on your Gemini key.", 502)


# ---------- Pro Animated PPT Maker (python-pptx only: no AI, no paid API) ----------
W, H = 13.333, 7.5


def an(fade, fly, zoom, wipe, pop):
    return {"fade": fade, "fly": fly, "zoom": zoom, "wipe": wipe, "pop": pop}


def _d(name, tags, bg, ang, tc, bc, acc, acc2, card, calpha, cline, ct, hl, hlt, font, tfont, decor, cardm, head, cov,
       anim, loop, trans, step, bt="#FFFFFF", sh="#000000"):
    return dict(name=name, tags=tags, bg=bg, ang=ang, tc=tc, bc=bc, acc=acc, acc2=acc2, card=card, calpha=calpha,
                cline=cline, ct=ct, hl=hl, hlt=hlt, font=font, tfont=tfont, decor=decor, cardm=cardm, head=head,
                cov=cov, anim=anim, loop=loop, trans=trans, step=step, bt=bt, sh=sh)


DESIGNS = {
    "neon": _d("Neon Tech", "Best for Tech, Startup, AI", ["#0F172A", "#1E1B4B"], 135, "#FFFFFF", "#CBD5E1", "#6366F1", "#22D3EE", "#1E293B", 70, "#6366F1", "#E2E8F0", ["#6366F1", "#22D3EE"], "#FFFFFF", "Poppins", "Poppins", "orbs", "glass", "plain", "orb", an("fade", "fly_b", "zoom", "wipe_l", "pop"), "float", "circle", .2),
    "minimal": _d("Minimal Elite", "Best for Pitch, Education, Portfolio", ["#FFFFFF", "#FFFFFF"], 0, "#000000", "#262626", "#000000", "#6B7280", "#F4F4F5", None, None, "#111111", "#000000", "#FFFFFF", "Inter", "Inter", "swiss", "swiss", "swiss", "swiss", an("appear", "wipe_u", "appear", "wipe_l", "appear"), None, "push_u", .15),
    "morph": _d("Gradient Morph", "Best for Creative, Marketing, Social", ["#7C3AED", "#EC4899"], 45, "#FFFFFF", "#F5F3FF", "#FDE68A", "#FFFFFF", "#FFFFFF", 18, "#FFFFFF", "#FFFFFF", ["#FDE68A", "#FBCFE8"], "#4C1D95", "Montserrat", "Montserrat", "blobs", "glass", "plain", "orb", an("fade", "fly_b", "grow", "wipe_r", "bounce"), "float", "zoom", .2, "#4C1D95"),
    "corp": _d("Corporate Pro", "Best for Business, Reports, Finance", ["#FFFFFF", "#EFF6FF"], 270, "#1E40AF", "#1F2937", "#1E40AF", "#3B82F6", "#FFFFFF", None, "#BFDBFE", "#1F2937", ["#1E40AF", "#3B82F6"], "#FFFFFF", "Inter", "Montserrat", "band", "solid", "band", "band", an("fade", "fade", "fade", "wipe_l", "fade"), None, "wipe", .12),
    "pop": _d("Creative Pop", "Best for Events, Youth, Kids", ["#FACC15", "#FACC15"], 0, "#000000", "#111111", "#000000", "#EF4444", "#FFFFFF", None, "#000000", "#000000", "#FFFFFF", "#000000", "Poppins", "Montserrat", "pop", "pop", "pop", "pop", an("pop", "bounce", "pop", "wipe_d", "grow"), "pulse", "diamond", .18),
    "ocean": _d("Ocean Wave", "Best for Rivers, Nature, Travel", ["#0369A1", "#06B6D4"], 135, "#FFFFFF", "#E0F2FE", "#67E8F9", "#FFFFFF", "#FFFFFF", 16, "#BAE6FD", "#F0F9FF", ["#E0F2FE", "#BAE6FD"], "#075985", "Poppins", "Montserrat", "waves", "glass", "plain", "orb", an("fade", "fly_l", "zoom", "wipe_r", "bounce"), "float", "split", .22, "#075985"),
    "forest": _d("Forest Green", "Best for Trees, Agriculture, Eco", ["#14532D", "#166534"], 135, "#FFFFFF", "#DCFCE7", "#86EFAC", "#FDE68A", "#FFFFFF", 12, "#86EFAC", "#F0FDF4", ["#86EFAC", "#FDE68A"], "#14532D", "Montserrat", "Poppins", "leaves", "glass", "plain", "center", an("fade", "fly_b", "grow", "wipe_u", "grow"), "sway", "cover", .22, "#14532D"),
    "sunset": _d("Sunset Glow", "Best for Events, Travel, Lifestyle", ["#F97316", "#DB2777"], 60, "#FFFFFF", "#FFEDD5", "#FEF08A", "#FFFFFF", "#FFFFFF", 18, "#FFFFFF", "#FFFFFF", ["#FEF08A", "#FDBA74"], "#7C2D12", "Poppins", "Poppins", "sun", "glass", "plain", "orb", an("fade", "fly_r", "zoom", "wipe_l", "pop"), "pulse", "pull", .2, "#7C2D12"),
    "royal": _d("Royal Gold", "Best for Luxury, Awards, Formal", ["#0A0A0A", "#1C1917"], 135, "#E7C468", "#E5E7EB", "#D4AF37", "#F5D77A", "#171717", None, "#D4AF37", "#F3F4F6", "#D4AF37", "#0A0A0A", "Georgia", "Georgia", "frame", "solid", "plain", "center", an("fade", "wipe_u", "fade", "wipe_l", "fade"), None, "fade", .3, "#0A0A0A"),
    "cyber": _d("Cyber Grid", "Best for Gaming, Hackathon, Code", ["#020617", "#0F172A"], 135, "#FFFFFF", "#A7F3D0", "#22FF88", "#FF2BD6", "#0B1220", None, "#22FF88", "#D1FAE5", ["#22FF88", "#FF2BD6"], "#020617", "Courier New", "Courier New", "grid", "solid", "plain", "orb", an("blinds", "fly_l", "checker", "wipe_l", "dissolve"), None, "checker", .12, "#020617"),
    "galaxy": _d("Space Galaxy", "Best for Science, Space, Future", ["#0B1026", "#3B0764"], 120, "#FFFFFF", "#E9D5FF", "#C4B5FD", "#FCD34D", "#1E1B4B", 55, "#A78BFA", "#EDE9FE", ["#8B5CF6", "#EC4899"], "#FFFFFF", "Poppins", "Poppins", "stars", "glass", "plain", "orb", an("fade", "zoom", "grow", "wipe_r", "pop"), "spin", "wheel", .2, "#2E1065"),
    "retro": _d("Retro Paper", "Best for History, Stories, Culture", ["#FFF7E6", "#FDEBC8"], 90, "#7F1D1D", "#431407", "#B91C1C", "#0F766E", "#FFFFFF", None, "#7F1D1D", "#431407", "#B91C1C", "#FFF7E6", "Georgia", "Georgia", "diag", "pop", "pop", "pop", an("fade", "fly_t", "pop", "wipe_d", "bounce"), None, "blinds", .22, "#FFFFFF", "#7F1D1D"),
    "chalk": _d("Chalkboard", "Best for Teachers, Lessons, Classes", ["#14352B", "#1B4332"], 135, "#F8FAFC", "#E2E8F0", "#FDE68A", "#93C5FD", "#FFFFFF", 8, "#F8FAFC", "#F1F5F9", ["#FDE68A", "#FCA5A5"], "#14352B", "Trebuchet MS", "Trebuchet MS", "arc", "glass", "plain", "center", an("wipe_l", "wipe_l", "dissolve", "wipe_l", "fade"), None, "dissolve", .25, "#14352B"),
    "medical": _d("Medical Clean", "Best for Health, Care, Science", ["#FFFFFF", "#ECFEFF"], 270, "#0F766E", "#134E4A", "#0D9488", "#06B6D4", "#FFFFFF", None, "#99F6E4", "#134E4A", ["#0D9488", "#22D3EE"], "#FFFFFF", "Inter", "Montserrat", "cross", "solid", "plain", "orb", an("fade", "fly_b", "circle", "wipe_l", "zoom"), "pulse", "plus", .18),
    "candy": _d("Candy Pastel", "Best for Kids, Fun, Social", ["#FDF2F8", "#E0F2FE"], 45, "#BE185D", "#475569", "#EC4899", "#38BDF8", "#FFFFFF", None, "#F9A8D4", "#334155", ["#F472B6", "#818CF8"], "#FFFFFF", "Poppins", "Poppins", "confetti", "solid", "plain", "orb", an("fade", "bounce", "pop", "wipe_r", "grow"), "float", "randomBar", .15),
    "luxe": _d("Dark Luxe", "Best for Fashion, Branding, Premium", ["#18181B", "#27272A"], 135, "#F4C2C2", "#D4D4D8", "#E8A0A0", "#FBCFE8", "#27272A", None, "#E8A0A0", "#F4F4F5", ["#E8A0A0", "#FBCFE8"], "#18181B", "Montserrat", "Georgia", "arc", "solid", "plain", "center", an("fade", "fly_t", "fade", "wipe_r", "fade"), None, "comb", .3, "#18181B"),
    "blueprint": _d("Blueprint", "Best for Engineering, Design, Plans", ["#1D4ED8", "#1E40AF"], 135, "#FFFFFF", "#DBEAFE", "#FFFFFF", "#93C5FD", "#1E3A8A", 50, "#FFFFFF", "#EFF6FF", "#FFFFFF", "#1E3A8A", "Courier New", "Courier New", "grid", "glass", "plain", "swiss", an("box", "wipe_r", "box", "wipe_l", "diamond"), None, "strips", .15, "#1E3A8A"),
    "fire": _d("Fire Energy", "Best for Sports, Motivation, Events", ["#111111", "#450A0A"], 135, "#FFFFFF", "#FECACA", "#F97316", "#EF4444", "#1F1F1F", None, "#F97316", "#FEE2E2", ["#F97316", "#EF4444"], "#FFFFFF", "Montserrat", "Montserrat", "diag", "solid", "plain", "swiss", an("fade", "fly_l", "zoom", "wipe_l", "bounce"), "pulse", "newsflash", .14),
    "earth": _d("Nature Earth", "Best for Environment, Culture, Craft", ["#F5EBDD", "#E7D5BE"], 90, "#44403C", "#57534E", "#B45309", "#65A30D", "#FFFBF5", None, "#D6B98C", "#44403C", ["#65A30D", "#B45309"], "#FFFFFF", "Georgia", "Montserrat", "leaves", "solid", "plain", "orb", an("fade", "fly_b", "circle", "wipe_u", "grow"), "sway", "wedge", .24),
    "aurora": _d("Aurora Glow", "Best for Tech, Music, Night Events", ["#022C22", "#1E1B4B"], 120, "#FFFFFF", "#CCFBF1", "#5EEAD4", "#C084FC", "#FFFFFF", 10, "#5EEAD4", "#F0FDFA", ["#5EEAD4", "#C084FC"], "#022C22", "Poppins", "Poppins", "blobs", "glass", "plain", "orb", an("fade", "fly_r", "diamond", "wipe_d", "pop"), "float", "push_l", .2, "#022C22"),
}
TYPES = {"title", "bullets", "two", "stats", "timeline", "chart", "conclusion"}
LIM = {"bullets": 4, "two": 4, "stats": 3, "timeline": 4, "conclusion": 3, "title": 0, "chart": 0}
TOOLS = ["AI Photo Generator", "AI Prompt Generator", "Response Sheet Analyser", "Marks Percentage",
         "Image & PDF Compressor", "Voice to Text"]
GENERIC = [
    ("Key Concepts", ["Core ideas behind {t}", "Important terms to know", "How the pieces fit together"], None),
    ("Why It Matters", ["Growing importance of {t}", "Benefits for people and organisations", "Opportunities it creates"], None),
    ("How It Works", ["Understand the need", "Plan and prepare", "Apply and test", "Review and improve"], "timeline"),
    ("Applications", ["Everyday uses of {t}", "Examples from different fields", "Case studies worth exploring"], None),
    ("Key Benefits", ["Saves time and effort", "Improves quality and accuracy", "Supports better decisions"], "stats"),
    ("Challenges", ["Common obstacles to expect", "Risks to manage", "Skills and resources needed"], None),
    ("Best Practices", ["Start with clear goals", "Learn from proven examples", "Measure results regularly", "Keep improving"], None),
    ("Future Outlook", ["Where {t} is heading", "Emerging trends to watch", "How to stay prepared"], None),
    ("Getting Started", ["Define your objective", "Gather the right resources", "Take the first small step", "Review and adjust"], "timeline"),
    ("Key Features", ["What makes {t} different", "Main components", "Practical advantages"], None),
    ("Examples", ["A simple everyday example", "A real-world scenario", "Lessons learned"], None),
    ("Tips for Success", ["Stay curious and keep learning", "Practise regularly", "Ask for feedback"], None),
    ("Common Questions", ["What is {t}?", "Who can use it?", "How do I begin?"], None),
]
NAMES = ["Overview", "Key Points", "Highlights", "Details", "More Details", "Insights"]
ICONS = [
    (("river", "water", "ocean", "sea", "lake", "rain", "flood", "wave", "marine", "dam"), "🌊", ["💧", "🏞️", "🐟"]),
    (("tree", "forest", "plant", "garden", "agricultur", "farm", "crop", "green", "leaf", "nature"), "🌳", ["🌿", "🍃", "🌱"]),
    (("animal", "wildlife", "bird", "zoo", "pet", "dog", "cat"), "🦁", ["🐾", "🐦", "🐘"]),
    (("space", "planet", "galaxy", "astronom", "moon", "rocket", "universe", "star"), "🚀", ["🪐", "🌙", "⭐"]),
    (("health", "medic", "hospital", "doctor", "disease", "nurs", "patient", "pharma"), "🩺", ["💊", "🏥", "❤️"]),
    (("ai", "artificial intelligence", "machine learning", "robot", "chatbot", "neural"), "🤖", ["🧠", "💡", "⚙️"]),
    (("computer", "software", "coding", "programming", "internet", "web", "data", "cloud", "tech", "digital"), "💻", ["🖥️", "📡", "⚙️"]),
    (("money", "financ", "bank", "econom", "budget", "invest", "stock", "tax", "loan"), "💰", ["📈", "🏦", "💳"]),
    (("business", "marketing", "startup", "sales", "company", "brand", "management"), "📈", ["💼", "🎯", "🤝"]),
    (("education", "school", "student", "learn", "exam", "college", "teacher", "study", "universit"), "🎓", ["📚", "✏️", "🏫"]),
    (("science", "chemistry", "physics", "lab", "experiment", "biology"), "🔬", ["⚗️", "🧪", "🧬"]),
    (("math", "algebra", "geometry", "calculus"), "📐", ["➕", "➗", "🔢"]),
    (("history", "ancient", "war", "king", "empire", "freedom", "independence"), "🏛️", ["📜", "⚔️", "🗿"]),
    (("earth", "climate", "environment", "pollution", "global warming", "geography", "world", "sustainab"), "🌍", ["♻️", "🌱", "🌤️"]),
    (("energy", "solar", "electric", "power", "wind", "battery", "fuel"), "⚡", ["☀️", "🔋", "💡"]),
    (("food", "cook", "recipe", "restaurant", "nutrition", "diet", "fruit"), "🍽️", ["🍎", "🥗", "🍳"]),
    (("travel", "tourism", "trip", "hotel", "flight", "holiday", "vacation"), "✈️", ["🧳", "🗺️", "🏖️"]),
    (("sport", "cricket", "football", "game", "olympic", "fitness", "gym", "yoga", "exercise"), "🏆", ["⚽", "🏏", "💪"]),
    (("music", "song", "dance", "singer", "concert"), "🎵", ["🎸", "🎤", "🎧"]),
    (("movie", "film", "cinema", "actor", "tv", "media"), "🎬", ["🎭", "📺", "🍿"]),
    (("art", "paint", "design", "draw", "creative", "photograph"), "🎨", ["🖌️", "📷", "✏️"]),
    (("car", "vehicle", "transport", "road", "traffic", "bike", "train", "bus"), "🚗", ["🚆", "🛣️", "🚦"]),
    (("city", "building", "architect", "construction", "urban", "home", "house"), "🏙️", ["🏗️", "🏠", "🌆"]),
    (("mountain", "hill", "trek", "hiking", "adventure"), "⛰️", ["🏕️", "🌄", "🧗"]),
    (("security", "cyber", "privacy", "hack", "safety", "password"), "🔒", ["🛡️", "🔑", "👁️"]),
    (("mobile", "phone", "app", "social media", "youtube", "instagram"), "📱", ["💬", "📲", "🔔"]),
    (("law", "police", "court", "justice", "crime", "legal", "constitution"), "⚖️", ["📜", "🏛️", "👮"]),
    (("culture", "festival", "temple", "religion", "india", "tradition"), "🪔", ["🎉", "🛕", "🌸"]),
    (("family", "parent", "child", "kids", "women", "society", "community"), "👨‍👩‍👧", ["❤️", "🤝", "🏡"]),
    (("election", "vote", "democracy", "government", "politic", "parliament"), "🗳️", ["🏛️", "📢", "🇮🇳"]),
    (("language", "english", "telugu", "grammar", "communication", "writing", "book", "literature"), "📖", ["✍️", "🔤", "💬"]),
    (("time", "schedule", "productivity", "habit", "goal"), "⏰", ["🎯", "✅", "📅"]),
]


def topic_icons(text):
    t = text.lower()
    for kws, main, ex in ICONS:
        for k in kws:
            if re.search((r"\b%s\b" if len(k) <= 3 else r"\b%s") % re.escape(k), t):
                return main, ex
    return "✨", ["⭐", "💡", "🚀"]


def cut(t, n=150):
    t = re.sub(r"\s+", " ", str(t)).strip()
    return t if len(t) <= n else t[:n].rsplit(" ", 1)[0].rstrip(",;:") + "…"


def parse_matter(desc):
    lines = [re.sub(r"^\s*(?:[-*•▪●]+|\d+[.)])\s*", "", ln).strip() for ln in desc.splitlines()]
    lines = [x for x in lines if x]
    data, rest = [], []
    for ln in lines:
        m = re.match(r"^(.{2,28}?)\s*[:=\-–]\s*(-?\d+(?:\.\d+)?)\s*%?$", ln)
        (data if m else rest).append([m.group(1).strip(), float(m.group(2))] if m else ln)
    if len(data) < 3:
        data, rest = [], lines
    groups, cur = [], None
    for k, ln in enumerate(rest):
        nxt = rest[k + 1] if k + 1 < len(rest) else ""
        words = ln.split()
        if ln.endswith(":") or (2 <= len(words) <= 6 and ln.istitle() and not re.search(r"[.!?,;]$", ln)
                                and nxt and not nxt.istitle()):
            cur = {"title": cut(ln.rstrip(":"), 60), "items": []}
            groups.append(cur)
            continue
        if cur is None:
            cur = {"title": None, "items": []}
            groups.append(cur)
        for sent in (re.split(r"(?<=[.!?])\s+", ln) if len(ln) > 170 else [ln]):
            if sent.strip():
                cur["items"].append(cut(sent))
    return [g for g in groups if g["items"]], data[:8]


def pick(n_items, k):
    ks = {1: ["bullets"], 2: ["bullets", "two"], 3: ["bullets", "two", "stats"]}.get(n_items, ["bullets", "two", "timeline"])
    return ks[k % len(ks)]


def make_outline(topic, desc, n):
    notes, T = [], topic[:40]
    slides = [{"type": "title", "title": topic, "sub": "A presentation on " + topic, "bullets": []}]
    groups, data = parse_matter(desc) if desc.strip() else ([], [])
    mids, k = [], 0
    if groups:
        for g in groups:
            for c in range(0, len(g["items"]), 4):
                chunk = g["items"][c:c + 4]
                title = (g["title"] or NAMES[len(mids) % len(NAMES)]) + (" (cont.)" if c else "")
                mids.append({"type": pick(len(chunk), k), "title": title, "bullets": chunk})
                k += 1
        want = max(1, n - 2 - (1 if data else 0))
        if len(mids) > want:
            notes.append("Your content needs more than %d slides, so the last points were left out. Increase the slide count to include them." % n)
            mids = mids[:want]
        end = [(m["bullets"][0] if len(m["bullets"][0]) <= 70 else m["title"].replace(" (cont.)", "")) for m in mids[:3]]
    else:
        notes.append("No content was given, so a ready-made template outline was used (no AI). Please review and edit it for your topic.")
        mids = [{"type": "bullets", "title": "Introduction", "bullets": [
            "What %s is and why it matters today" % T, "Who benefits and how it is used", "What this presentation covers"]}]
        for j in range(max(0, n - 3)):
            ttl, bl, kind = GENERIC[j % len(GENERIC)]
            mids.append({"type": kind or pick(len(bl), k), "title": ttl + (" (%d)" % (j // len(GENERIC) + 1) if j >= len(GENERIC) else ""),
                         "bullets": [cut(b.replace("{t}", T)) for b in bl]})
            k += 1
        end = ["%s offers real opportunities" % T, "Start small and keep learning", "Thank you - questions are welcome"]
    slides += mids
    if data:
        slides.append({"type": "chart", "title": "Data at a Glance", "data": data, "bullets": []})
    slides.append({"type": "conclusion", "title": "Conclusion" if not groups else "Key Takeaways", "bullets": end})
    return clean_slides(slides), notes


def clean_slides(raw):
    out, imgs = [], 0
    for x in lst(raw)[:20]:
        if not isinstance(x, dict) or x.get("type") not in TYPES:
            continue
        t = x["type"]
        sl = {"type": t, "title": s(x.get("title"), 80).strip() or "Untitled", "sub": s(x.get("sub"), 120).strip(), "bullets": []}
        items = [cut(b) for b in lst(x.get("bullets")) if str(b).strip()][:LIM[t]]
        if t == "chart":
            data = []
            for r in lst(x.get("data"))[:8]:
                try:
                    data.append([s(r[0], 24), float(r[1])])
                except (TypeError, ValueError, IndexError):
                    pass
            if len(data) >= 2:
                sl["data"] = data
            else:
                t, items = "bullets", ["Add your points here"]
        elif t != "title":
            if t in ("stats", "timeline") and len(items) < 3:
                t = "bullets"
            if not items:
                t, items = "bullets", ["Add your points here"]
            if t == "stats":
                sl["nums"] = [s(q, 8) for q in lst(x.get("nums"))[:3]]
        im = x.get("img")
        if t != "chart" and imgs < 8:
            if im is True:
                sl["img"], imgs = True, imgs + 1
            elif isinstance(im, str) and len(im) < 1800000 and re.match(r"^data:image/(png|jpeg|webp);base64,[A-Za-z0-9+/=]+$", im):
                sl["img"], imgs = im, imgs + 1
        sl["type"], sl["bullets"] = t, items
        out.append(sl)
    return out


def S(k, x, y, w, h, **o):
    d = {"k": k, "x": x, "y": y, "w": w, "h": h}
    d.update(o)
    return d


def fit(text, w, h, size, lo=12):
    while size > lo:
        cpl = max(1, int((w - .3) * 72 / (size * .52)))
        if sum(max(1, -(-len(p) // cpl)) for p in text.split("\n")) * size * 1.3 / 72 <= h - .15:
            break
        size -= 1
    return size


# ----- decoration recipes: static shapes + "actors" (named !!aN so PowerPoint Morph moves them between slides) -----
def recipe(D, i, cover=False):
    r, A, B = D["decor"], D["acc"], D["acc2"]
    ac = lambda k, f, ps, **o: dict(k=k, f=f, ps=ps, **o)
    if r == "orbs":
        return [], [ac("oval", A, [(8.8, -2.6, 6.4, 6.4, 0), (9.6, 3.4, 5.6, 5.6, 0), (-2.2, -2.8, 5.8, 5.8, 0), (10.4, -3.2, 7, 7, 0)], al=22, lp=1),
                    ac("oval", B, [(-2.2, 4.8, 5, 5, 0), (7.6, 5.2, 4.4, 4.4, 0), (10.2, -1.6, 3.8, 3.8, 0), (-1.6, 5.4, 4.8, 4.8, 0)], al=14, lp=1)]
    if r == "blobs":
        return [], [ac("oval", "#FFFFFF" if D["bg"][0] != "#022C22" else B, [(-1.6, -1.8, 5.2, 5.2, 0), (9.2, -2, 5.6, 5.6, 0), (3.2, 4.2, 5, 5, 0), (10.6, 3.6, 4.6, 4.6, 0)], al=14, lp=1),
                    ac("oval", A, [(9.6, 4.6, 3.4, 3.4, 0), (-1, 5.2, 3.4, 3.4, 0), (10.4, -1.4, 3.4, 3.4, 0), (1, -2.2, 3.4, 3.4, 0)], al=28, lp=1),
                    ac("oval", B, [(5.5, -3, 3, 3, 0), (5, 5.8, 3, 3, 0), (-2, 1, 2.6, 2.6, 0), (11.5, 1, 2.6, 2.6, 0)], al=18)]
    if r == "swiss":
        return [S("rect", .3, H - .55, W - .3, .02, f=A), S("text", W - 2.0, H - .5, 1.6, .4, t="%02d" % (i + 1), sz=14, c=B, b=True, ha="r", ft=D["tfont"])], \
               [ac("rect", A, [(0, 0, .3, H, 0), (0, 0, .55, H, 0), (0, 0, .3, H, 0), (0, 0, .16, H, 0)])]
    if r == "band":
        if cover:
            return [S("rect", 0, 0, 7.6, H, f=A), S("rect", 7.6, 0, .08, H, f=B)], []
        return [S("rect", 0, 0, W, 1.35, f=A), S("rect", 0, 1.35, W, .07, f=B), S("rect", 0, H - .42, W, .42, f="#DBEAFE"),
                S("text", .7, H - .41, 7, .38, t=D["site"], sz=11, c=A, b=True, va="m", ft=D["font"]),
                S("text", W - 2.0, H - .41, 1.3, .38, t="%02d" % (i + 1), sz=11, c=A, b=True, ha="r", va="m", ft=D["font"])], \
               [ac("diamond", B, [(12.6, 1.9, 1.6, 1.6, 45), (12.9, 2.8, 1.8, 1.8, 45), (12.4, 2.2, 1.5, 1.5, 60), (12.8, 1.6, 1.7, 1.7, 30)], al=30, lp=1),
                ac("tri", A, [(12.5, 5.2, 1.4, 1.2, 0), (12.7, 5.0, 1.4, 1.2, 15), (12.4, 5.4, 1.4, 1.2, -10), (12.8, 5.1, 1.4, 1.2, 0)], al=25)]
    if r == "pop":
        return [S("rect", 0, H - .2, W, .2, f=D["sh"])], \
               [ac("oval", D["sh"], [(W - 2.7, -1.4, 3.8, 3.8, 0), (W - 2.2, -1.8, 3.8, 3.8, 0), (W - 3.1, -1.6, 3.8, 3.8, 0), (W - 2.5, -1.2, 3.6, 3.6, 0)]),
                ac("star", "#FFFFFF", [(.25, H - .92, .7, .7, 0), (.25, H - .92, .7, .7, 36), (.25, H - .92, .7, .7, 72), (.25, H - .92, .7, .7, 18)], l=D["cline"], lw=2.5, lp=1)]
    if r == "waves":
        return [], [ac("oval", "#FFFFFF", [(-3, 6.7, 10, 3.6, 0), (-1, 6.9, 10, 3.6, 0), (3, 6.7, 11, 3.6, 0), (-4, 6.9, 11, 3.6, 0)], al=14),
                    ac("oval", B, [(5, 7.0, 10, 3.6, 0), (1, 7.1, 10, 3.6, 0), (-3, 7.0, 11, 3.6, 0), (6, 7.1, 11, 3.6, 0)], al=12),
                    ac("oval", "#FFFFFF", [(11.2, 1.0, .5, .5, 0), (11.6, 3.4, .5, .5, 0), (12.0, 5.0, .5, .5, 0), (11.4, 2.0, .5, .5, 0)], al=35, lp=1),
                    ac("oval", "#FFFFFF", [(12.2, 2.6, .8, .8, 0), (12.4, .8, .8, .8, 0), (11.8, 4.2, .8, .8, 0), (12.5, 3.6, .8, .8, 0)], al=25, lp=1)]
    if r == "leaves":
        return [], [ac("oval", A, [(-1.4, -.6, 3.6, 1.5, 35), (-1.0, .2, 3.8, 1.6, 20), (-1.6, -.2, 3.6, 1.5, 50), (-1.2, .6, 3.4, 1.4, 28)], al=40, lp=1),
                    ac("oval", B, [(11.2, 5.9, 3.4, 1.4, -35), (11.6, 5.4, 3.4, 1.4, -20), (11.0, 6.1, 3.2, 1.4, -50), (11.4, 5.6, 3.6, 1.5, -28)], al=35, lp=1),
                    ac("oval", A, [(10.8, -.9, 3.0, 1.3, -30), (11.4, -.5, 3.0, 1.3, -45), (10.6, -.7, 3.2, 1.4, -15), (11.2, -1.0, 3.0, 1.3, -38)], al=28),
                    ac("oval", B, [(-1.4, 6.6, 3.0, 1.3, 30), (-1.0, 6.2, 3.0, 1.3, 15), (-1.6, 6.5, 3.2, 1.4, 45), (-1.2, 6.8, 3.0, 1.3, 22)], al=30)]
    if r == "sun":
        return [], [ac("oval", A, [(9.2, -2.2, 5, 5, 0), (10, 3.4, 4.4, 4.4, 0), (-1.8, -2.4, 4.6, 4.6, 0), (10.6, -2.8, 5.4, 5.4, 0)], al=45, lp=1),
                    ac("oval", None, [(8.2, -3.2, 7, 7, 0), (9, 2.5, 6.2, 6.2, 0), (-2.8, -3.4, 6.6, 6.6, 0), (9.6, -3.8, 7.4, 7.4, 0)], l="#FFFFFF", lw=2, lp=1),
                    ac("oval", B, [(-2, 5, 4, 4, 0), (7.8, 5.6, 3.4, 3.4, 0), (10.6, -1, 3, 3, 0), (-1.6, 5.6, 3.6, 3.6, 0)], al=14)]
    if r == "grid":
        st = [S("rect", 1.6 * k, 0, .012, H, f=A, al=10) for k in range(1, 9)] + [S("rect", 0, 1.5 * k, W, .012, f=A, al=10) for k in range(1, 5)]
        return st, [ac("hex", B, [(10.6, -.6, 2.2, 2.0, 0), (.2, 5.4, 2.2, 2.0, 30), (11, 5.2, 2.4, 2.2, 15), (.4, -.8, 2.2, 2.0, 45)], al=30, lp=1),
                    ac("diamond", A, [(11.8, .6, 1.2, 1.2, 0), (-.2, 6.2, 1.2, 1.2, 45), (12, 2.4, 1.2, 1.2, 20), (.5, .4, 1.2, 1.2, 60)], al=28, lp=1)]
    if r == "stars":
        import random
        rnd, out = random.Random(7), []
        for k in range(9):
            bx, by, s0 = rnd.uniform(.3, 12.6), rnd.uniform(.2, 7.0), rnd.uniform(.14, .34)
            if .9 < bx < 11.8 and .35 < by < 6.6:
                by = 6.75 + rnd.uniform(0, .4)
            out.append(ac(("star", "diamond", "oval")[k % 3], "#FFFFFF" if k % 2 else B,
                          [(bx + dx, by + dy, s0, s0, rot) for dx, dy, rot in ((0, 0, 0), (.35, .25, 40), (-.3, .4, 80), (.2, -.3, 20))], al=70, lp=1))
        out.append(ac("oval", None, [(9.8, -1.4, 5, 5, 0), (9.2, -.8, 5.4, 5.4, 0), (10.2, -1.8, 5, 5, 0), (9.6, -1.2, 5.2, 5.2, 0)], l=A, lw=1.5))
        out.append(ac("oval", A, [(10.6, 1.1, 1.4, 1.4, 0), (11.2, 4.6, 1.4, 1.4, 0), (10.0, 5.5, 1.4, 1.4, 0), (11.6, .6, 1.4, 1.4, 0)], al=40))
        return [], out
    if r == "diag":
        return [], [ac("rect", A, [(9.3, -1.5, .55, 11, 28), (8.6, -1.8, .55, 11, 32), (9.6, -1.4, .55, 11, 24), (8.9, -1.6, .55, 11, 30)], al=55, lp=1),
                    ac("rect", B, [(10.4, -1.5, .3, 11, 28), (9.8, -1.8, .3, 11, 32), (10.8, -1.4, .3, 11, 24), (10.1, -1.6, .3, 11, 30)], al=60),
                    ac("rect", A, [(-1.2, 5.9, .4, 9, -62), (-.8, 5.6, .4, 9, -58), (-1.4, 6.2, .4, 9, -66), (-1, 5.8, .4, 9, -60)], al=40)]
    if r == "frame":
        st = [S("rect", .35, .3, W - .7, .03, f=A), S("rect", .35, H - .33, W - .7, .03, f=A),
              S("rect", .35, .3, .03, H - .6, f=A), S("rect", W - .38, .3, .03, H - .6, f=A)]
        dm = lambda x, y: ac("diamond", A, [(x, y, .32, .32, 0)] * 4)
        return st, [dm(.2, .15), dm(W - .52, .15), dm(.2, H - .47), dm(W - .52, H - .47),
                    ac("oval", None, [(8.5, -3, 8, 8, 0), (9.5, -2, 7.4, 7.4, 0), (-3.5, -2.5, 7, 7, 0), (10, -3.4, 8, 8, 0)], l=B, lw=1.2)]
    if r == "cross":
        return [], [ac("cross", A, [(11.6, .5, 1.4, 1.4, 0), (11.9, 5.4, 1.4, 1.4, 20), (-.4, 5.6, 1.4, 1.4, 45), (12.2, 3.0, 1.4, 1.4, -20)], al=22, lp=1),
                    ac("cross", B, [(.2, 6.1, .8, .8, 0), (11.0, .4, .8, .8, 30), (.5, .4, .8, .8, 15), (12.4, 6.2, .8, .8, 45)], al=30, lp=1),
                    ac("oval", B, [(-1.4, -1.4, 3.4, 3.4, 0), (10.8, 5.4, 3.6, 3.6, 0), (11, -1.6, 3.4, 3.4, 0), (-1.6, 5.6, 3.6, 3.6, 0)], al=12)]
    if r == "confetti":
        return [], [ac("tri", A, [(12.5, 2.6, .7, .6, 0), (11.8, 6.0, .7, .6, 40), (.2, 6.2, .7, .6, 80), (12.4, 2.2, .7, .6, 20)], al=85, lp=1),
                    ac("hex", B, [(12.3, 3.5, .6, .55, 0), (12.5, 4.4, .6, .55, 30), (.2, 6.0, .6, .55, 60), (12.4, 5.0, .6, .55, 10)], al=85, lp=1),
                    ac("diamond", "#FACC15", [(.2, 6.7, .5, .5, 0), (12.6, 1.8, .5, .5, 40), (.1, 6.4, .5, .5, 10), (11.2, 6.7, .5, .5, 70)], al=90, lp=1),
                    ac("oval", A, [(10.6, 6.2, .4, .4, 0), (.5, 3.4, .4, .4, 0), (12.4, 4.4, .4, .4, 0), (11.7, 1.9, .4, .4, 0)], al=70, lp=1),
                    ac("star", B, [(.05, 3.0, .5, .5, 0), (12.3, 6.4, .6, .6, 25), (12.5, 3.0, .5, .5, 50), (.1, 6.3, .5, .5, 10)], al=80, lp=1),
                    ac("oval", A, [(10.4, -1.2, 3, 3, 0), (-1.4, 5.8, 3, 3, 0), (10, 5.8, 3.4, 3.4, 0), (-1.5, -1.4, 3, 3, 0)], al=10)]
    if r == "arc":
        return [], [ac("oval", None, [(-5.4, -3.4, 10.5, 13.5, 0), (-4.4, -4.2, 10, 13.5, 0), (-6, -2.6, 11, 13.5, 0), (-5, -3, 10.5, 14, 0)], l=A, lw=2.5),
                    ac("oval", None, [(9.6, -4.2, 8, 8, 0), (8.8, -3.4, 8, 8, 0), (10.2, -5, 8.4, 8.4, 0), (9.4, -3.8, 8, 8, 0)], l=B, lw=1.5),
                    ac("oval", A, [(12.0, 6.4, .28, .28, 0), (.8, .9, .28, .28, 0), (12.4, .8, .28, .28, 0), (1.2, 6.5, .28, .28, 0)], al=90, lp=1)]
    return [], []


def deco(D, i, cover=False):
    st, acts = recipe(D, i, cover)
    out = list(st)
    for n, a in enumerate(acts):
        x, y, w, h, rot = a["ps"][i % len(a["ps"])]
        d = S(a["k"], x, y, w, h, f=a["f"], n="!!a%d" % n)
        for key in ("al", "l", "lw", "lp"):
            if key in a:
                d[key] = a[key]
        if rot:
            d["rot"] = rot
        out.append(d)
    return out


def head(D, text, ctx, icon=True):
    hm, tf = D["head"], D["tfont"]
    iw = 1.4 if icon else 0
    if hm == "band":
        tw = W - 1.4 - iw
        o = [S("text", .7, .17, tw, 1.05, t=text, sz=fit(text, tw, 1.0, 40, 24), c="#FFFFFF", b=True, va="m", ft=tf, a="fade", n="!!title")]
        top = 1.95
    else:
        x = .95 if hm == "swiss" else .7
        tw = W - x - .7 - iw
        o = [S("text", x, .5, tw, 1.15, t=text, sz=fit(text, tw, 1.1, 48 if hm == "swiss" else 44, 26), c=D["tc"], b=True, va="m", ft=tf, a="fade", n="!!title")]
        if hm == "pop":
            o.append(S("rect", x, 1.72, 3.0, .16, f=D["sh"], a="wipe", n="!!line"))
        else:
            o.append(S("rect", x, 1.72, 1.2 if hm == "swiss" else 1.5, .1 if hm == "swiss" else .09,
                       f=[D["acc"], D["acc2"]] if D["cardm"] == "glass" else D["acc"], a="wipe", n="!!line"))
        top = 2.15
    if icon:
        o.append(S("text", W - 1.9, .12 if hm == "band" else .3, 1.3, 1.15 if hm == "band" else 1.2, t=ctx["ic"], sz=44, ha="c", va="m",
                   c="#FFFFFF", a="pop", lp=2, n="!!icon"))
    return o, top


def cardbox(D, x, y, w, h, text=None, **o):
    cm, out = D["cardm"], []
    if cm == "pop":
        out.append(S("round", x + .1, y + .1, w, h, f=D["sh"], a=o.get("a")))
    kw = dict(f=D["card"], al=D.get("calpha"), l=D["cline"], lw=3 if cm == "pop" else 1.25, a=o.get("a"))
    if text is not None:
        kw.update(t=text, sz=o["sz"], c=D["ct"], ft=D["font"], va="m", pl=o.get("pl", .15))
    out.append(S("rect" if cm == "swiss" else "round", x, y, w, h, **kw))
    return out


def hlbox(D, x, y, w, h):
    cm = D["cardm"]
    if cm == "swiss":
        return [S("rect", x, y, w, h, f=D["hl"], a="fly")]
    if cm == "pop":
        return [S("round", x + .1, y + .1, w, h, f=D["sh"], a="fly"), S("round", x, y, w, h, f=D["hl"], l=D["cline"], lw=4, a="fly")]
    return [S("round", x, y, w, h, f=D["hl"], ang=45, a="fly")]


def picbox(D, x, y, w, h, sl, si, full, a="fly"):
    src = sl["img"] if full and isinstance(sl.get("img"), str) else None
    o = []
    if D["cardm"] == "pop":
        o.append(S("round", x + .1, y + .1, w, h, f=D["sh"], a=a))
    o.append(S("pic", x, y, w, h, si=si, src=src, l=D["cline"] if D["cardm"] == "pop" else D["acc"], lw=3, a=a))
    return o


FP = {"orb": [(8.6, 1.0), (12.2, 4.9), (7.9, 5.5)], "center": [(1.6, 1.0), (W - 2.6, 1.2), (W - 3.2, 5.6)],
      "band": [(8.2, .7), (12.0, 5.6), (8.4, 5.2)], "swiss": [(W - 1.7, .5), (W - 1.8, 5.6), (W - 3.2, 6.0)],
      "pop": [(W - 3.7, .4), (11.6, 5.4), (.5, .5)]}


def cover(D, sl, ctx, si):
    cv, T, A, B, tf = D["cov"], sl["title"], D["acc"], D["acc2"], D["tfont"]
    sub, im, ic = sl.get("sub") or "", sl.get("img"), ctx["ic"]
    if cv == "center" and im:
        cv = "orb"
    pic = picbox(D, 8.4, 1.5, 4.3, 4.3, sl, si, ctx["full"], "zoom") if im else None
    o = []
    if cv == "swiss":
        w = 7.2 if im else 8.6
        o = [S("rect", .95, 1.9, 2.2, .14, f=A, a="wipe", n="!!line"),
             S("text", .95, 2.2, w, 2.9, t=T, sz=fit(T, w, 2.8, 72, 30), c=D["tc"], b=True, va="m", ft=tf, a="fade", n="!!title"),
             S("text", .95, 5.3, w, .9, t=sub, sz=24, c=D["bc"], ft=D["font"], a="fade"),
             S("text", .95, 6.3, 8, .4, t="PRESENTATION  /  " + D["site"].upper(), sz=12, c=B, b=True, ft=D["font"], a="fade")]
        o += pic or [S("text", 9.6, 1.6, 3.2, 3.2, t=ic, sz=110, ha="c", va="m", c="#FFFFFF", a="zoom", lp=2, n="!!icon")]
    elif cv == "band":
        o = [S("text", .7, 2.0, 6.4, 3.0, t=T, sz=fit(T, 6.4, 2.9, 48, 26), c="#FFFFFF", b=True, va="m", ft=tf, a="fade", n="!!title"),
             S("rect", .8, 5.15, 1.5, .09, f=B, a="wipe", n="!!line"),
             S("text", .7, 5.35, 6.4, 1.2, t=sub, sz=20, c="#DBEAFE", ft=D["font"], a="fade")]
        o += pic or [S("oval", 9.3, 1.3, 3.4, 3.4, f=[A, B], ang=45, a="zoom"),
                     S("text", 9.3, 1.3, 3.4, 3.4, t=ic, sz=84, ha="c", va="m", c="#FFFFFF", a="zoom", lp=2, n="!!icon"),
                     S("text", 8.6, 5.0, 4.8, .6, t="PRESENTATION", sz=18, c=A, b=True, ha="c", ft=D["font"], a="fade")]
    elif cv == "pop":
        w = 7.0 if im else 9.4
        o = [S("round", 1.12, 2.12, w, 3.0, f=D["sh"], a="zoom"),
             S("round", 1.0, 2.0, w, 3.0, f="#FFFFFF", l=D["cline"], lw=4, a="zoom"),
             S("text", 1.3, 2.15, w - .6, 2.7, t=T, sz=fit(T, w - .6, 2.6, 54, 26), c=D["tc"], b=True, va="m", ft=tf, a="pop", n="!!title"),
             S("text", 1.0, 5.4, w, .9, t=sub, sz=22, c=D["bc"], b=True, ft=D["font"], a="fade")]
        o += pic or [S("text", 10.3, 3.4, 2.4, 2.4, t=ic, sz=96, ha="c", va="m", c="#FFFFFF", a="pop", lp=2, n="!!icon")]
    elif cv == "center":
        o = [S("text", W / 2 - 1.2, 1.0, 2.4, 1.6, t=ic, sz=80, ha="c", va="m", c="#FFFFFF", a="zoom", lp=2, n="!!icon"),
             S("text", 1.5, 2.7, W - 3, 1.9, t=T, sz=fit(T, W - 3, 1.8, 54, 28), c=D["tc"], b=True, ha="c", va="m", ft=tf, a="fade", n="!!title"),
             S("rect", W / 2 - .8, 4.75, 1.6, .09, f=[A, B], a="wipe", n="!!line"),
             S("text", 1.5, 5.0, W - 3, 1.0, t=sub, sz=22, c=D["bc"], ha="c", ft=D["font"], a="fade")]
    else:
        w = 7.0 if im else 8.4
        o = [S("round", .9, 1.45, 3.2, .5, f=A, al=28, l=A, lw=1, t="✦  PRESENTATION", sz=13, c=A if D["cardm"] == "glass" and D["bt"] == "#FFFFFF" else D["bt"], b=True, ha="c", va="m", pl=0, ft=D["font"], a="fade"),
             S("text", .9, 2.15, w, 2.5, t=T, sz=fit(T, w, 2.4, 56, 28), c=D["tc"], b=True, va="m", ft=tf, a="fade", n="!!title"),
             S("rect", .95, 4.8, 1.6, .09, f=[A, B], a="wipe", n="!!line"),
             S("text", .9, 5.05, w, 1.0, t=sub, sz=22, c=D["bc"], ft=D["font"], a="fade")]
        o += pic or [S("oval", 9.5, 1.7, 3.4, 3.4, f=[A, B], ang=45, al=90, a="zoom"),
                     S("text", 9.5, 1.7, 3.4, 3.4, t=ic, sz=84, ha="c", va="m", c="#FFFFFF", a="zoom", lp=2, n="!!icon")]
    for k, (fx, fy) in enumerate([] if im else FP.get(cv, FP["orb"])):
        o.append(S("text", fx, fy, 1.1, 1.1, t=ctx["ics"][k], sz=38, ha="c", va="m", c="#FFFFFF", a="pop", lp=1))
    if ctx["logo"]:
        o.append(S("pic", W - 1.45, H - 1.4, 1.0, 1.0, lg=1, a="fade"))
    return o


def layout(D, sl, i, si, ctx):
    t, cm, items = sl["type"], D["cardm"], sl.get("bullets", [])
    if t == "title":
        return deco(D, i, True) + cover(D, sl, ctx, si)
    sh = deco(D, i)
    hd, top = head(D, sl["title"], ctx)
    sh += hd
    bf, bt, x0, tf = D["acc"], D["bt"], (.95 if cm == "swiss" else .7), D["tfont"]
    if t == "agenda":
        for k, b in enumerate(items):
            y = top + k * .66
            sh += [S("oval", x0, y + .06, .5, .5, f=bf, t=str(k + 1), sz=16, c=bt, b=True, ha="c", va="m", pl=0, ft=tf, a="zoom"),
                   S("text", x0 + .75, y, 10.8, .62, t=b, sz=fit(b, 10.8, .6, 24, 14), c=D["bc"], va="m", ft=D["font"], a="fly")]
        return sh
    if sl.get("img") and t != "chart":
        txt = "\n".join("▸ " + b for b in items)
        sh.append(S("text", x0, top, 6.5, 4.5, t=txt, sz=fit(txt, 6.5, 4.5, 24, 14), c=D["bc"], ft=D["font"], a="fade"))
        return sh + picbox(D, 7.6, top, 5.0, 4.5, sl, si, ctx["full"])
    if t == "bullets":
        if cm == "swiss":
            for k, b in enumerate(items):
                y = top + k * 1.2
                sh += [S("text", .95, y, 1.3, 1.0, t="%02d" % (k + 1), sz=28, c=D["acc2"], b=True, va="m", ft=tf, a="fade"),
                       S("text", 2.4, y, 9.6, 1.0, t=b, sz=fit(b, 9.6, 1.0, 26, 14), c=D["bc"], va="m", ft=D["font"], a="wipe"),
                       S("rect", .95, y + 1.07, 11.4, .015, f=D["acc2"], al=40, a="wipe")]
        else:
            n = len(items)
            if n <= 2:
                cells = [(.7, top + k * 2.15, W - 1.4, 1.9) for k in range(n)]
            else:
                cw = (W - 1.4 - .4) / 2
                cells = [(.7 + (k % 2) * (cw + .4), top + (k // 2) * 2.35, cw, 2.1) for k in range(n)]
            for k, (x, y, w, h) in enumerate(cells):
                sh += cardbox(D, x, y, w, h, items[k], sz=fit(items[k], w - 1.15, h, 22, 13), pl=1.05, a="fly")
                sh.append(S("oval", x + .25, y + h / 2 - .32, .64, .64, f=bf, t=str(k + 1), sz=20, c=bt, b=True, ha="c", va="m", pl=0, ft=tf, a="zoom"))
    elif t == "two":
        left = items[:-1] if len(items) > 1 else items
        call = items[-1] if len(items) > 1 else ""
        txt = "\n".join("▸ " + b for b in left)
        sh.append(S("text", x0, top, 7.1, 4.4, t=txt, sz=fit(txt, 7.1, 4.4, 24, 14), c=D["bc"], ft=D["font"], a="fade"))
        sh += hlbox(D, 8.2, top, 4.45, 4.4)
        sh += [S("text", 8.2, top + .25, 4.45, 1.5, t=ctx["ic"], sz=60, ha="c", va="m", c="#FFFFFF", a="zoom"),
               S("text", 8.5, top + 1.85, 3.85, 2.3, t=call, sz=fit(call, 3.85, 2.3, 24, 14), c=D["hlt"], b=True, ha="c", va="m", ft=D["font"], a="fade")]
    elif t in ("stats", "conclusion"):
        nums = sl.get("nums") or []
        cw, y0, ch = (W - 1.4 - .8) / 3, top + .1, (3.4 if t == "stats" else 2.9)
        for k in range(min(3, len(items))):
            x = .7 + k * (cw + .4)
            num = nums[k] if k < len(nums) and nums[k] else ("✓" if t == "conclusion" else "%02d" % (k + 1))
            sh += cardbox(D, x, y0, cw, ch, a="fly")
            sh += [S("text", x, y0 + .2, cw, 1.4, t=num, sz=fit(num, cw, 1.3, 54, 24), c=D["acc"], b=True, ha="c", va="m", ft=tf, a="zoom"),
                   S("text", x + .2, y0 + 1.65, cw - .4, ch - 1.85, t=items[k], sz=fit(items[k], cw - .4, ch - 1.85, 20, 13), c=D["ct"], ha="c", ft=D["font"], a="fade")]
        if t == "conclusion":
            yb = y0 + ch + .4
            sh += hlbox(D, .7, yb, W - 1.4, 1.0)
            sh.append(S("text", .7, yb, W - 1.4, 1.0, t="Thank you!  Questions are welcome", sz=28, c=D["hlt"], b=True, ha="c", va="m", ft=tf, a="fade"))
    elif t == "timeline":
        n, y0 = len(items), top + .9
        sh.append(S("rect", 1.1, y0 + .33, W - 2.2, .07, f=D["acc"], a="wipe"))
        for k, b in enumerate(items):
            cx = 1.1 + .37 + k * ((W - 2.2 - .74) / max(1, n - 1))
            sh += [S("oval", cx - .37, y0, .74, .74, f=bf, t=str(k + 1), sz=22, c=bt, b=True, ha="c", va="m", pl=0, ft=tf, a="zoom"),
                   S("text", min(max(cx - 1.35, .35), W - .35 - 2.7), y0 + 1.1, 2.7, 2.6, t=b, sz=fit(b, 2.7, 2.6, 20, 13), c=D["bc"], ha="c", ft=D["font"], a="fade")]
    elif t == "chart":
        data = sl["data"]
        hi, lo = max(data, key=lambda r: r[1]), min(data, key=lambda r: r[1])
        txt = "Highest\n%s: %g\nLowest\n%s: %g" % (hi[0], hi[1], lo[0], lo[1])
        sh.append(S("chart", .7, top, 7.9, 4.6, data=data, cc=D["acc"], c=D["bc"], ft=D["font"], a="wipe"))
        sh += hlbox(D, 8.9, top, 3.75, 4.6)
        sh.append(S("text", 9.1, top + .2, 3.35, 4.2, t=txt, sz=22, c=D["hlt"], b=True, ha="c", va="m", ft=D["font"], a="fade"))
    return sh


def promo(D, i, ctx):
    cm = D["cardm"]
    x0 = .95 if D["head"] == "swiss" else .7
    hd, top = head(D, "Made with " + D["site"], ctx, icon=False)
    sh = deco(D, i) + hd
    if ctx["logo"]:
        sh.append(S("pic", W - 2.3, .22 if D["head"] == "band" else .3, 1.2, 1.2, lg=1, a="zoom", lp=2, n="!!icon"))
    else:
        sh.append(S("text", W - 1.9, .3, 1.3, 1.2, t=ctx["ic"], sz=44, ha="c", va="m", c="#FFFFFF", a="pop", lp=2, n="!!icon"))
    sh.append(S("text", x0, top - .05, 9.4, .9, t=ctx["tag"], sz=22, c=D["bc"], ft=D["font"], a="fade"))
    for k, name in enumerate(TOOLS):
        x, y = x0 + (k % 3) * 4.0, 3.15 + (k // 3) * 1.0
        sh += cardbox(D, x, y, 3.75, .8, "★  " + name, sz=fit(name, 3.4, .75, 17, 12), pl=.15, a="fly")
    dom = re.sub(r"^https?://", "", ctx["url"])
    sh += hlbox(D, 2.4, 5.25, 8.5, 1.05)
    sh.append(S("text", 2.4, 5.25, 8.5, 1.05, t="🔗  " + dom, sz=fit(dom, 8, 1.0, 30, 18), c=D["hlt"], b=True, ha="c", va="m", ft=D["tfont"], a="zoom", url=ctx["url"], lp=2))
    sh.append(S("star", W - 2.4, 4.75, 1.1, 1.1, f=D["acc2"], l=D["cline"] if cm == "pop" else None, lw=3, a="pop"))
    sh.append(S("text", 2.4, 6.4, 8.5, .4, t="Open the link - it's free, no sign-up needed", sz=15, c=D["bc"], ha="c", ft=D["font"], a="fade"))
    return sh


def build_layouts(did, slides, info, topic="", full=False, agenda=False):
    D = dict(DESIGNS[did], site=info["site"])
    ic, ics = topic_icons(topic + " " + " ".join(q["title"] for q in slides[1:5]))
    ctx = dict(ic=ic, ics=ics, logo=info["logo"], full=full, url=info["url"], tag=info["tag"])
    seq = []
    for k, sl in enumerate(slides):
        seq.append((sl, k))
        mids = [q["title"].replace(" (cont.)", "") for q in slides[1:-1][:7]]
        if k == 0 and agenda and len(slides) >= 4:
            seq.append(({"type": "agenda", "title": "Outline", "bullets": mids}, None))
    seq.append(({"type": "promo"}, None))
    out = []
    for i, (sl, si) in enumerate(seq):
        shapes = promo(D, i, ctx) if sl["type"] == "promo" else layout(D, sl, i, si, ctx)
        k = 0
        for d in shapes:
            if d.get("a"):
                d["a"] = D["anim"].get(d["a"], d["a"])
                d["dl"] = round(min(k * D["step"], 3.2), 2)
                k += 1
            lp = d.pop("lp", 0)
            if lp:
                lo = D["loop"] if lp == 1 else (D["loop"] or "pulse")
                if lo == "spin" and d["k"] in ("oval", "text", "pic"):
                    lo = "pulse"
                if lo:
                    d["lo"] = lo
        out.append({"bg": {"c": D["bg"], "ang": D["ang"]}, "shapes": shapes,
                    "kind": "promo" if sl["type"] == "promo" else "agenda" if sl["type"] == "agenda" else "slide"})
    return out


# ----- animation XML: entrance effects + looping emphasis effects -----
FILT = {"fade": (10, 0, "fade"), "wipe_l": (22, 8, "wipe(left)"), "wipe_r": (22, 2, "wipe(right)"), "wipe_u": (22, 1, "wipe(up)"),
        "wipe_d": (22, 4, "wipe(down)"), "blinds": (3, 10, "blinds(horizontal)"), "box": (4, 16, "box(in)"),
        "checker": (5, 10, "checkerboard(across)"), "circle": (6, 16, "circle(in)"), "diamond": (8, 16, "diamond(out)"),
        "dissolve": (9, 0, "dissolve"), "plus": (13, 16, "plus(out)"), "randbars": (14, 10, "randombar(horizontal)"),
        "wedge": (20, 0, "wedge"), "wheel": (21, 1, "wheel(1)")}
FLY = {"fly_b": (4, "ppt_y", "1+#ppt_h/2"), "fly_t": (1, "ppt_y", "0-#ppt_h/2"), "fly_l": (8, "ppt_x", "0-#ppt_w/2"), "fly_r": (2, "ppt_x", "1+#ppt_w/2")}
LOOPS = {"spin": 8, "pulse": 26, "float": 0, "sway": 32}
FALL = {"fade": "<p:fade/>", "push_u": '<p:push dir="u"/>', "push_l": '<p:push dir="l"/>', "zoom": '<p:zoom dir="in"/>',
        "wipe": '<p:wipe dir="r"/>', "diamond": "<p:diamond/>", "split": '<p:split orient="vert" dir="out"/>',
        "cover": '<p:cover dir="l"/>', "pull": '<p:pull dir="r"/>', "circle": "<p:circle/>", "checker": '<p:checker dir="horz"/>',
        "wheel": '<p:wheel spokes="4"/>', "blinds": '<p:blinds dir="horz"/>', "dissolve": "<p:dissolve/>", "plus": "<p:plus/>",
        "randomBar": '<p:randomBar dir="horz"/>', "comb": '<p:comb dir="horz"/>', "strips": '<p:strips dir="rd"/>',
        "newsflash": "<p:newsflash/>", "wedge": "<p:wedge/>"}
NSP = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'


def trans_xml(did):
    return ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" ' + NSP + '>'
            '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
            'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p159">'
            '<p:transition spd="slow" p14:dur="1500"><p159:morph option="byObject"/></p:transition></mc:Choice>'
            '<mc:Fallback><p:transition spd="med">' + FALL[DESIGNS[did]["trans"]] + '</p:transition></mc:Fallback></mc:AlternateContent>')


def _tav(tm, v):
    try:
        float(v)
        val = '<p:fltVal val="%s"/>' % v
    except ValueError:
        val = '<p:strVal val="%s"/>' % v
    return '<p:tav tm="%d"><p:val>%s</p:val></p:tav>' % (tm, val)


def _anim(i, spid, attr, keys, add="", dur=500):
    return ('<p:anim calcmode="lin" valueType="num"><p:cBhvr%s><p:cTn id="%d" dur="%d" fill="hold"/>'
            '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl><p:attrNameLst><p:attrName>%s</p:attrName></p:attrNameLst></p:cBhvr>'
            '<p:tavLst>%s</p:tavLst></p:anim>') % (add, i, dur, spid, attr, "".join(_tav(t, v) for t, v in keys))


def _effect(nid, spid, kind, delay):
    ms, tg = int(delay * 1000), '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>' % spid
    if kind in LOOPS:
        eid = nid()
        if kind == "spin":
            body = '<p:animRot by="21600000"><p:cBhvr><p:cTn id="%d" dur="9000" fill="hold"/>%s<p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>' % (nid(), tg)
        elif kind == "sway":
            body = '<p:animRot by="360000"><p:cBhvr><p:cTn id="%d" dur="2200" autoRev="1" fill="hold"/>%s<p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>' % (nid(), tg)
        elif kind == "pulse":
            body = '<p:animScale><p:cBhvr><p:cTn id="%d" dur="1100" autoRev="1" fill="hold"/>%s</p:cBhvr><p:by x="108000" y="108000"/></p:animScale>' % (nid(), tg)
        else:
            body = ('<p:animMotion origin="layout" path="M 0 0 L 0 -0.05 E" pathEditMode="relative" rAng="0" ptsTypes="AA"><p:cBhvr>'
                    '<p:cTn id="%d" dur="2600" autoRev="1" fill="hold"/>%s<p:attrNameLst><p:attrName>ppt_x</p:attrName>'
                    '<p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:rCtr x="0" y="-2500"/></p:animMotion>') % (nid(), tg)
        return ('<p:par><p:cTn id="%d" presetID="%d" presetClass="%s" presetSubtype="0" repeatCount="indefinite" fill="hold" grpId="1" '
                'nodeType="withEffect"><p:stCondLst><p:cond delay="%d"/></p:stCondLst><p:childTnLst>%s</p:childTnLst></p:cTn></p:par>'
                ) % (eid, LOOPS[kind], "path" if kind == "float" else "emph", ms, body)
    eid = nid()
    inner = ('<p:set><p:cBhvr><p:cTn id="%d" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>%s'
             '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
             '<p:to><p:strVal val="visible"/></p:to></p:set>') % (nid(), tg)

    def fx(f, dur=500):
        return '<p:animEffect transition="in" filter="%s"><p:cBhvr><p:cTn id="%d" dur="%d"/>%s</p:cBhvr></p:animEffect>' % (f, nid(), dur, tg)
    pid, sub = 1, 0
    if kind in FILT:
        pid, sub, f = FILT[kind]
        inner += fx(f)
    elif kind in FLY:
        sub, attr, start = FLY[kind]
        pid = 2
        other = "ppt_x" if attr == "ppt_y" else "ppt_y"
        inner += _anim(nid(), spid, other, [(0, "#" + other), (100000, "#" + other)], ' additive="base"')
        inner += _anim(nid(), spid, attr, [(0, start), (100000, "#" + attr)], ' additive="base"')
    elif kind == "bounce":
        pid = 26
        inner += _anim(nid(), spid, "ppt_x", [(0, "#ppt_x"), (100000, "#ppt_x")], ' additive="base"', 700)
        inner += _anim(nid(), spid, "ppt_y", [(0, "1+#ppt_h/2"), (65000, "#ppt_y-0.05"), (82000, "#ppt_y+0.02"), (100000, "#ppt_y")], ' additive="base"', 700)
        inner += fx("fade", 700)
    elif kind in ("zoom", "pop", "grow"):
        pid, sub = (31, 0) if kind == "grow" else (53, 16)
        keys = [(0, "0"), (70000, "#ppt_%s*1.15"), (100000, "#ppt_%s")] if kind == "pop" else [(0, "0"), (100000, "#ppt_%s")]
        for ax in ("w", "h"):
            inner += _anim(nid(), spid, "ppt_" + ax, [(t, v.replace("%s", ax)) for t, v in keys])
        if kind == "grow":
            inner += _anim(nid(), spid, "style.rotation", [(0, "90"), (100000, "0")])
        inner += fx("fade")
    return ('<p:par><p:cTn id="%d" presetID="%d" presetClass="entr" presetSubtype="%d" fill="hold" grpId="0" nodeType="withEffect">'
            '<p:stCondLst><p:cond delay="%d"/></p:stCondLst><p:childTnLst>%s</p:childTnLst></p:cTn></p:par>') % (eid, pid, sub, ms, inner)


def add_animation(anims, shape, kind="fade", delay=0.0, mode="n"):
    """Queue an effect (entrance: fade/fly/zoom/wipe/pop/grow/bounce/...; loops: spin/pulse/float/sway) for a shape."""
    anims.append((shape.shape_id, kind, delay, mode))


def timing_xml(anims):
    c = [4]

    def nid():
        c[0] += 1
        return c[0]
    fx = "".join(_effect(nid, sp, k, d) for sp, k, d, _m in anims)
    seen, bld = set(), ""
    for sp, k, _d, m in anims:
        if m != "n" and k not in LOOPS and sp not in seen:
            seen.add(sp)
            bld += ('<p:bldGraphic spid="%d" grpId="0"><p:bldAsOne/></p:bldGraphic>' % sp if m == "c" else
                    '<p:bldP spid="%d" grpId="0"%s/>' % (sp, ' animBg="1"' if m == "sh" else ""))
    return ('<p:timing ' + NSP + '><p:tnLst><p:par>'
            '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
            '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
            '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
            '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>' + fx +
            '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
            '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>' + ("<p:bldLst>%s</p:bldLst>" % bld if bld else "") + '</p:timing>')


def build_pptx(layouts, did, logo=None):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn
    from lxml import etree
    from PIL import Image
    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())
    kinds = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL, "star": MSO_SHAPE.STAR_5_POINT,
             "tri": MSO_SHAPE.ISOSCELES_TRIANGLE, "diamond": MSO_SHAPE.DIAMOND, "hex": MSO_SHAPE.HEXAGON, "cross": MSO_SHAPE.CROSS}
    al = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
    va = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
    D = DESIGNS[did]

    def paint(shp, d):
        f = d.get("f")
        if f is None:
            shp.fill.background()
        elif isinstance(f, list):
            shp.fill.gradient()
            shp.fill.gradient_angle = d.get("ang", 45)
            for stp, col in zip(shp.fill.gradient_stops, f):
                stp.color.rgb = rgb(col)
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(f)
        if d.get("al") is not None and f is not None:
            for e in shp._element.spPr.iter(qn("a:srgbClr")):
                etree.SubElement(e, qn("a:alpha")).set("val", str(int(d["al"] * 1000)))
        if d.get("l"):
            shp.line.color.rgb = rgb(d["l"])
            shp.line.width = Pt(d.get("lw", 1))
        else:
            shp.line.fill.background()

    def put_text(shp, d):
        tf = shp.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left, tf.margin_right = Inches(d.get("pl", .12)), Inches(.12)
        tf.margin_top = tf.margin_bottom = Inches(.06)
        tf.vertical_anchor = va[d.get("va", "t")]
        for j, line in enumerate(d["t"].split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = al[d.get("ha", "l")]
            p.space_after = Pt(8)
            r = p.add_run()
            r.text = line
            r.font.size, r.font.bold = Pt(d.get("sz", 18)), bool(d.get("b"))
            r.font.name = d.get("ft") or "Calibri"
            r.font.color.rgb = rgb(d.get("c", "#000000"))

    def add_pic(sl, d, x, y, w, h):
        try:
            if d.get("lg"):
                if not logo:
                    return None
                im = Image.open(io.BytesIO(logo))
                k = min(w / im.width, h / im.height)
                bw, bh = int(im.width * k), int(im.height * k)
                return sl.shapes.add_picture(io.BytesIO(logo), x + (w - bw) // 2, y + (h - bh) // 2, bw, bh)
            src = d.get("src")
            if not src:
                return None
            im = Image.open(io.BytesIO(base64.b64decode(src.split(",", 1)[1]))).convert("RGB")
            ar = w / h
            cw, ch = (int(im.height * ar), im.height) if im.width / im.height > ar else (im.width, int(im.width / ar))
            l0, t0 = (im.width - cw) // 2, (im.height - ch) // 2
            im = im.crop((l0, t0, l0 + cw, t0 + ch))
            im.thumbnail((1600, 1600))
            bio = io.BytesIO()
            im.save(bio, "JPEG", quality=88)
            bio.seek(0)
            pic = sl.shapes.add_picture(bio, x, y, w, h)
            pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
            if d.get("l"):
                pic.line.color.rgb = rgb(d["l"])
                pic.line.width = Pt(d.get("lw", 3))
            return pic
        except Exception:
            return None

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for lay in layouts:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        c, fl = lay["bg"]["c"], sl.background.fill
        if c[0] == c[1]:
            fl.solid()
            fl.fore_color.rgb = rgb(c[0])
        else:
            fl.gradient()
            fl.gradient_angle = lay["bg"]["ang"]
            for stp, col in zip(fl.gradient_stops, c):
                stp.color.rgb = rgb(col)
        anims = []
        for d in lay["shapes"]:
            k, mode = d["k"], "n"
            x, y, w, h = (Inches(d[q]) for q in ("x", "y", "w", "h"))
            if k == "chart":
                cd = CategoryChartData()
                cd.categories = [r[0] for r in d["data"]]
                cd.add_series("Data", [r[1] for r in d["data"]])
                shp = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, cd)
                ch = shp.chart
                ch.has_legend = ch.has_title = False
                ch.font.size, ch.font.name = Pt(14), d.get("ft") or "Calibri"
                ch.font.color.rgb = rgb(d["c"])
                pl = ch.plots[0]
                pl.gap_width, pl.has_data_labels = 60, True
                pl.data_labels.font.size, pl.data_labels.font.bold = Pt(14), True
                pl.data_labels.font.color.rgb = rgb(d["c"])
                pl.series[0].format.fill.solid()
                pl.series[0].format.fill.fore_color.rgb = rgb(d["cc"])
                ch.value_axis.has_major_gridlines = False
                ch.value_axis.visible = False
                mode = "c"
            elif k == "pic":
                shp = add_pic(sl, d, x, y, w, h)
                if shp is None:
                    continue
            elif k == "text":
                shp, mode = sl.shapes.add_textbox(x, y, w, h), "tb"
            else:
                shp = sl.shapes.add_shape(kinds[k], x, y, w, h)
                shp.shadow.inherit = False
                if k == "round":
                    shp.adjustments[0] = .12
                if k == "cross":
                    shp.adjustments[0] = .33
                paint(shp, d)
                mode = "sh" if d.get("t") else "n"
            if d.get("t") and k not in ("chart", "pic"):
                put_text(shp, d)
            if d.get("url"):
                shp.click_action.hyperlink.address = d["url"]
            if d.get("rot"):
                shp.rotation = d["rot"]
            if d.get("n"):
                shp.name = d["n"]
            if d.get("a"):
                add_animation(anims, shp, d["a"], d.get("dl", 0), mode)
            if d.get("lo"):
                add_animation(anims, shp, d["lo"], (d.get("dl", 0) + .6) if d.get("a") else 0, "n")
        sld = sl._element
        sld.append(parse_xml(trans_xml(did)))
        if anims:
            sld.append(parse_xml(timing_xml(anims)))
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()


def cfg_doc():
    try:
        db = get_db()
        doc = (db.config.find_one({"_id": "site"}) if db is not None else None) or {}
    except Exception:
        doc = {}
    doc.pop("_id", None)
    return {**DEFAULT, **doc}


def site_info():
    c = cfg_doc()
    proto = (request.headers.get("X-Forwarded-Proto") or request.scheme).split(",")[0].strip()
    return {"site": c["name"], "url": "%s://%s" % (proto, request.host), "tag": c.get("tag") or DEFAULT["tag"], "logo": bool(c.get("logo"))}


def logo_bytes():
    lg = cfg_doc().get("logo")
    if lg == "/logo.png":
        return base64.b64decode(LOGO_B64)
    m = re.match(r"^/f/([0-9a-f]{24})$", lg or "")
    if m:
        try:
            from bson import ObjectId
            d = get_db().files.find_one({"_id": ObjectId(m.group(1))})
            return bytes(d["data"]) if d else None
        except Exception:
            return None
    return None


@app.post("/api/ppt/outline")
def ppt_outline():
    limit("pptA", 60)
    j = request.get_json(silent=True) or {}
    topic = s(j.get("topic"), 80).strip()
    if not topic:
        raise Err("Enter a topic first")
    n = max(5, min(15, int(num(j.get("n"), 8))))
    slides, notes = make_outline(topic, s(j.get("desc"), 6000), n)
    info = site_info()
    return jsonify(slides=slides, note=" ".join(notes), designs=[
        {"id": k, "name": v["name"], "tags": v["tags"], "thumb": build_layouts(k, slides[:2], info, topic)[:2]}
        for k, v in DESIGNS.items()])


def _deck(j, full=False):
    did = j.get("design")
    if did not in DESIGNS:
        raise Err("Unknown design")
    slides = clean_slides(j.get("slides"))
    if not slides:
        raise Err("There are no slides to build")
    return did, build_layouts(did, slides, site_info(), s(j.get("topic"), 80), full, True)


@app.post("/api/ppt/preview")
def ppt_preview():
    limit("pptB", 600)
    return jsonify(layouts=_deck(request.get_json(silent=True) or {})[1])


@app.post("/api/ppt/build")
def ppt_build():
    limit("pptC", 30)
    j = request.get_json(silent=True) or {}
    did, lays = _deck(j, True)
    fn = re.sub(r"[^A-Za-z0-9]+", "_", s(j.get("topic"), 60)).strip("_")[:40] or "Presentation"
    return Response(build_pptx(lays, did, logo_bytes()), mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    headers={"Content-Disposition": 'attachment; filename="%s_%s_animated.pptx"' % (fn, DESIGNS[did]["name"].replace(" ", ""))})


@app.errorhandler(Exception)
def oops(e):
    if isinstance(e, Err):
        return jsonify(error=e.msg), e.code
    if isinstance(e, HTTPException):
        return jsonify(error="File too large (max 12 MB)" if e.code == 413 else e.description), e.code
    return jsonify(error="Server error: " + str(e)[:120]), 500


# ----- Website (HTML/CSS/JS) -----
HTML = r'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+Telugu&family=Noto+Sans+Devanagari&family=Noto+Sans+Tamil&family=Noto+Sans+Kannada&family=Noto+Sans+Malayalam&family=Noto+Sans+Arabic&display=swap" rel="stylesheet">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;700&family=Poppins:wght@400;700&family=Montserrat:wght@400;700&display=swap" rel="stylesheet">
<title>ToolHub – All-in-One Tools</title>
<style>
:root{--bg:#f4f6fb;--card:#fff;--tx:#0f172a;--mu:#64748b;--ac:#2563eb;--ac2:#7c3aed;--bd:#e5e9f2;--sh:0 4px 18px rgba(37,99,235,.08);box-sizing:border-box;padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}
@media(prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#0b1020;--card:#151b2e;--tx:#e8ecf8;--mu:#93a0bd;--bd:#232b44;--sh:0 4px 18px rgba(0,0,0,.35)}}
:root[data-theme="dark"]{--bg:#0b1020;--card:#151b2e;--tx:#e8ecf8;--mu:#93a0bd;--bd:#232b44;--sh:0 4px 18px rgba(0,0,0,.35)}
*{box-sizing:border-box}
html{scroll-padding-top:env(safe-area-inset-top,0px)}
body{margin:0;background:var(--bg);color:var(--tx);font-family:Inter,system-ui,-apple-system,Segoe UI,Roboto,sans-serif;line-height:1.4}
.wrap{max-width:1000px;margin:auto;padding:16px}
header{display:flex;align-items:center;justify-content:space-between;gap:10px}
.logo{font-weight:800;font-size:20px;display:flex;align-items:center;gap:8px}
.logo i{font-style:normal;width:34px;height:34px;border-radius:10px;background:linear-gradient(135deg,var(--ac),var(--ac2));display:grid;place-items:center;color:#fff}
.tg{border:1px solid var(--bd);background:var(--card);color:var(--tx);border-radius:50%;width:38px;height:38px;font-size:17px;cursor:pointer}
.hero{margin:18px 0;padding:26px 20px;border-radius:22px;background:linear-gradient(135deg,var(--ac),var(--ac2));color:#fff;box-shadow:var(--sh);animation:up .6s both}
.hero h1{margin:0 0 6px;font-size:clamp(22px,5vw,34px)}
.hero p{margin:0 0 16px;opacity:.9}
.search{width:100%;border:0;border-radius:14px;padding:14px 16px;font-size:16px;outline:none;background:#fff;color:#0f172a}
.tick{overflow:hidden;background:var(--card);border:1px solid var(--bd);border-radius:14px;display:flex;align-items:center}
.tick b{background:var(--ac);color:#fff;padding:9px 12px;font-size:12px;letter-spacing:.5px;flex:none}
.tr{overflow:hidden;flex:1}
.tm{display:inline-flex;gap:40px;white-space:nowrap;padding-left:20px;animation:mq 25s linear infinite}
.tm a{color:var(--tx);font-size:14px;cursor:pointer}
.tr:hover .tm{animation-play-state:paused}
.chips{display:flex;gap:8px;overflow-x:auto;padding:16px 0 8px;scrollbar-width:none}
.chip{flex:none;border:1px solid var(--bd);background:var(--card);color:var(--tx);padding:8px 16px;border-radius:99px;font-size:14px;cursor:pointer}
.chip.on{background:var(--ac);color:#fff;border-color:var(--ac)}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(150px,1fr));gap:14px;margin-top:8px}
.card{background:var(--card);border:1px solid var(--bd);border-radius:20px;padding:16px;cursor:pointer;box-shadow:var(--sh);transition:transform .2s,box-shadow .2s;animation:up .5s both}
.card:hover{transform:translateY(-5px);box-shadow:0 12px 28px rgba(37,99,235,.2)}
.card:active{transform:scale(.97)}
.ic{width:48px;height:48px;border-radius:14px;display:grid;place-items:center;font-size:24px;margin-bottom:12px}
.card h3{margin:0 0 4px;font-size:15px}
.card p{margin:0;color:var(--mu);font-size:12.5px}
.tag{display:inline-block;margin-top:10px;font-size:11px;color:var(--ac);background:rgba(37,99,235,.1);padding:3px 8px;border-radius:99px}
.ov{position:fixed;inset:0;background:rgba(5,10,25,.55);display:none;align-items:flex-end;justify-content:center;z-index:9}
.ov.on{display:flex}
.sh{background:var(--card);width:100%;max-width:520px;border-radius:24px 24px 0 0;padding:20px 20px calc(20px + env(safe-area-inset-bottom,0px));max-height:90%;overflow:auto;animation:sl .3s both}
@media(min-width:600px){.ov{align-items:center}.sh{border-radius:24px}}
.sh h2{margin:0 0 14px;display:flex;justify-content:space-between;align-items:center;font-size:19px}
.x{border:0;background:var(--bg);color:var(--tx);width:34px;height:34px;border-radius:50%;font-size:16px;cursor:pointer}
label{display:block;font-size:13px;color:var(--mu);margin:10px 0 4px}
input,select,textarea{width:100%;padding:12px;border-radius:12px;border:1px solid var(--bd);background:var(--bg);color:var(--tx);font-size:16px;font-family:inherit}
.go{width:100%;margin-top:16px;border:0;border-radius:14px;padding:14px;font-size:16px;font-weight:700;color:#fff;background:linear-gradient(135deg,var(--ac),var(--ac2));cursor:pointer}
.res{margin-top:14px;padding:14px;border-radius:14px;background:rgba(37,99,235,.1);font-weight:600;word-break:break-word;display:none}
.empty{text-align:center;color:var(--mu);padding:30px;display:none}
footer{text-align:center;color:var(--mu);font-size:13px;padding:24px 0 8px}
@keyframes up{from{opacity:0;transform:translateY(16px)}to{opacity:1;transform:none}}
@keyframes sl{from{transform:translateY(40px);opacity:0}to{transform:none;opacity:1}}
@keyframes mq{to{transform:translateX(-50%)}}
.hb{display:flex;gap:8px}.lk{display:flex;gap:8px;flex-wrap:wrap;margin-top:12px}.pl{background:rgba(255,255,255,.22);color:#fff;text-decoration:none;padding:7px 14px;border-radius:99px;font-size:13px}
.pop{animation:pp .45s}@keyframes pp{0%{transform:scale(.8);opacity:0}60%{transform:scale(1.05)}100%{transform:none;opacity:1}}
.big{font-size:30px;color:var(--ac);font-weight:800}.bar{height:10px;border-radius:9px;background:var(--bd);margin-top:10px;overflow:hidden}.bar b{display:block;height:100%;width:0;background:linear-gradient(90deg,var(--ac),var(--ac2))}
.cf{position:absolute;bottom:0;font-style:normal;animation:fl 1.2s both}@keyframes fl{to{transform:translateY(-70px) rotate(40deg);opacity:0}}
.lv{animation:pu 1s infinite;background:#ef4444!important}@keyframes pu{50%{transform:scale(1.03)}}
.vc{display:inline-flex;gap:6px;align-items:center;flex-wrap:wrap;justify-content:center;background:var(--card);border:1px solid var(--bd);border-radius:99px;padding:8px 16px;margin-bottom:10px;font-size:14px;box-shadow:var(--sh)}.vc b{color:var(--ac)}
.ar{border:1px solid var(--bd);border-radius:14px;padding:10px;margin:8px 0;display:grid;gap:6px}.ar small{color:var(--mu);word-break:break-all}.r2{display:grid;grid-template-columns:70px 1fr;gap:6px}input[type=checkbox]{width:auto}
.hot{position:relative;background:linear-gradient(var(--card),var(--card)) padding-box,linear-gradient(135deg,var(--ac),var(--ac2)) border-box;border:2px solid transparent}.hot::after{content:'🔥 HOT';position:absolute;top:10px;right:10px;font-size:10px;font-weight:800;background:#ef4444;color:#fff;padding:2px 7px;border-radius:99px}
.vs{display:none;margin:26px auto 6px;max-width:520px;border-radius:24px;padding:20px;background:linear-gradient(135deg,var(--ac),var(--ac2));color:#fff;box-shadow:var(--sh);text-align:center}.vs .t{font-size:12px;opacity:.92;display:flex;gap:8px;align-items:center;justify-content:center;letter-spacing:1px;font-weight:700}.vs .dt{width:9px;height:9px;border-radius:50%;background:#4ade80;animation:lp 1.6s infinite}.vs .r{display:flex;justify-content:space-around;margin-top:14px}.vs .n{font-size:36px;font-weight:800;line-height:1.1}.vs .l{font-size:12px;opacity:.88}
@keyframes lp{0%{box-shadow:0 0 0 0 rgba(74,222,128,.8)}70%{box-shadow:0 0 0 10px rgba(74,222,128,0)}100%{box-shadow:0 0 0 0 rgba(74,222,128,0)}}
.r3{display:grid;grid-template-columns:1fr 78px 78px;gap:6px;margin-top:6px}.sc{border-radius:16px;padding:12px;color:#fff;text-align:center;font-size:13px;animation:up .5s both}.sc b{display:block;font-size:30px}
.gt{margin-top:16px;border-radius:22px;padding:20px;color:#fff;text-align:center;background:linear-gradient(135deg,var(--ac),var(--ac2));box-shadow:var(--sh);animation:pp .5s}.gt .m{font-size:54px;font-weight:800;line-height:1.1;margin:4px 0}.gt .st{display:flex;justify-content:space-around;margin-top:14px;font-size:13px}.gt .st b{display:block;font-size:24px}
.sg{display:grid;grid-template-columns:repeat(auto-fill,minmax(150px,1fr));gap:10px;margin-top:12px}.sk{background:var(--card);border:1px solid var(--bd);border-radius:18px;padding:14px;animation:up .5s both}.sk h4{margin:0 0 4px;font-size:15px}.sk .mk{font-size:30px;font-weight:800;color:var(--ac);margin-bottom:6px}.sk small{display:block;color:var(--mu);font-size:12px;margin-top:6px}
.card.wide{grid-column:1/-1;display:flex;align-items:center;gap:18px;padding:22px 24px;color:#fff;border:0;overflow:hidden;position:relative;background:linear-gradient(120deg,#6366f1,#ec4899,#f59e0b,#6366f1);background-size:300% 300%;animation:up .5s both,bgm 8s ease infinite}
.card.wide h3{font-size:22px;color:#fff;margin:0 0 4px}.card.wide p{color:rgba(255,255,255,.92);font-size:14px;margin:0}.card.wide .ic{width:68px;height:68px;font-size:34px;background:rgba(255,255,255,.22);flex:none;margin:0}
.card.wide .bt{margin-left:auto;background:#fff;color:#6366f1;font-weight:800;padding:12px 22px;border-radius:99px;flex:none;box-shadow:0 6px 18px rgba(0,0,0,.2)}.card.wide .fl{position:absolute;right:190px;top:14px;font-size:30px;opacity:.4;animation:fy 4s ease-in-out infinite}
.card.wide::before{content:'🔥 NEW';position:absolute;top:0;left:0;background:#ef4444;color:#fff;font-size:11px;font-weight:800;padding:4px 12px;border-radius:0 0 14px 0}
@keyframes bgm{50%{background-position:100% 50%}}@keyframes fy{50%{transform:translateY(-12px) rotate(10deg)}}
@media(max-width:600px){.card.wide{flex-wrap:wrap}.card.wide .bt{width:100%;text-align:center;margin:0}.card.wide .fl{display:none}}
.dg{display:grid;grid-template-columns:repeat(auto-fill,minmax(250px,1fr));gap:12px;margin-top:12px}.dc{border:1px solid var(--bd);border-radius:18px;padding:10px;background:var(--card);box-shadow:var(--sh);transition:transform .2s}.dc:hover{transform:translateY(-4px)}.dc .tr2{display:flex;gap:6px}.dc canvas{flex:1;min-width:0;width:50%;border-radius:8px}.dc h4{margin:10px 0 2px}.dc small{color:var(--mu)}
.pv{width:100%;border-radius:14px;box-shadow:var(--sh);display:block}.th{display:flex;gap:8px;overflow-x:auto;margin-top:10px;padding-bottom:6px}.th canvas{width:110px;flex:none;border-radius:6px;cursor:pointer;border:2px solid transparent}.th canvas.on{border-color:var(--ac)}
.logo i img{height:38px;width:38px;object-fit:contain;display:block}.logo i{width:auto;min-width:34px}
</style>
</head>
<body>
<div class="wrap">
<header><div class="logo"><i>⚡</i><span id="nm">ToolHub</span></div><div class="hb"><button class="tg" id="ad" aria-label="Admin">⚙️</button><button class="tg" id="tg" aria-label="Toggle theme">🌓</button></div></header>
<section class="hero"><h1>Every tool you need, in one place</h1><p>Fast, free and simple calculators &amp; utilities.</p><input class="search" id="q" placeholder="🔍 Search tools..."><div class="lk" id="lk"></div></section>
<div class="tick"><b>LATEST</b><div class="tr"><div class="tm" id="tm"></div></div></div>
<div class="chips" id="chips"></div>
<div class="grid" id="grid"></div>
<div class="empty" id="em">No tools found 😕</div>
<div class="vs" id="vcw"><div class="t"><span class="dt"></span>LIVE VISITOR STATS</div><div class="r"><div><div class="n" id="vc">0</div><div class="l">👥 Total visitors</div></div><div><div class="n" id="vt">0</div><div class="l">👁 Page views</div></div></div></div>
<footer>Made with ❤️ · ToolHub</footer>
</div>
<div class="ov" id="ov"><div class="sh" id="sh"></div></div>
<script>
var N=function(i){return{key:i}};
var T=[
{id:'ppt',n:'Pro Animated PPT Maker',d:'Cinematic animated PowerPoint - 20 designs, topic icons, your photos - free, no AI',i:'📊',c:'Creator',g:'#c4b5fd,#6366f1',u:ppt},
{id:'age',n:'Age Calculator',d:'Exact age in years, months, days',i:'🎂',c:'Education',g:'#fde68a,#f59e0b',f:[{k:'dob',l:'Date of birth',t:'date'}],r:function(v){var d=new Date(v.dob),n=new Date();if(isNaN(d))return'Pick a date';var y=n.getFullYear()-d.getFullYear(),m=n.getMonth()-d.getMonth(),x=n.getDate()-d.getDate();if(x<0){m--;x+=new Date(n.getFullYear(),n.getMonth(),0).getDate()}if(m<0){y--;m+=12}return y+' years, '+m+' months, '+x+' days'}},
{id:'sgpa',n:'SGPA to Percentage',d:'Convert SGPA to marks %',i:'🎓',c:'Education',g:'#bfdbfe,#3b82f6',f:[{k:'s',l:'SGPA (out of 10)',t:'number'}],r:function(v){var s=+v.s;if(!(s>0&&s<=10))return'Enter SGPA 0-10';return'Percentage: '+((s-0.75)*10).toFixed(2)+'%'}},
{id:'att',n:'Attendance',d:'Check 75% attendance status',i:'📅',c:'Education',g:'#a7f3d0,#10b981',f:[{k:'a',l:'Classes attended',t:'number'},{k:'t',l:'Total classes',t:'number'}],r:function(v){var a=+v.a,t=+v.t;if(!(t>0))return'Enter total classes';var p=a/t*100,n=Math.max(0,Math.ceil((0.75*t-a)/0.25));return'Attendance: '+p.toFixed(2)+'% — '+(p>=75?'Safe ✅':'Attend next '+n+' classes to reach 75% ⚠️')}},
{id:'emi',n:'EMI Calculator',d:'Monthly loan instalment',i:'💰',c:'Finance',g:'#fecdd3,#f43f5e',f:[{k:'p',l:'Loan amount (₹)',t:'number'},{k:'r',l:'Interest rate (% per year)',t:'number'},{k:'y',l:'Years',t:'number'}],r:function(v){var P=+v.p,r=+v.r/1200,n=+v.y*12;if(!(P>0&&r>0&&n>0))return'Fill all fields';var q=Math.pow(1+r,n),e=P*r*q/(q-1);return'EMI: ₹'+e.toFixed(0)+'/month · Total interest: ₹'+(e*n-P).toFixed(0)}},
{id:'wc',n:'Word Counter',d:'Words, characters, sentences',i:'📝',c:'Text',g:'#ddd6fe,#8b5cf6',f:[{k:'x',l:'Paste your text',t:'area'}],r:function(v){var x=v.x.trim();if(!x)return'Type some text';return'Words: '+x.split(/\s+/).length+' · Characters: '+x.length+' · Sentences: '+(x.split(/[.!?]+/).filter(Boolean).length)}},
{id:'pw',n:'Password Generator',d:'Strong random passwords',i:'🔐',c:'Text',g:'#fed7aa,#f97316',f:[{k:'l',l:'Length (8-64)',t:'number',v:16}],r:function(v){var l=Math.min(64,Math.max(8,+v.l||16)),s='abcdefghijkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789!@#$%&*?',a=new Uint32Array(l);crypto.getRandomValues(a);return Array.from(a,function(z){return s[z%s.length]}).join('')}},
{id:'bmi',n:'BMI Calculator',d:'Body mass index check',i:'❤️',c:'Health',g:'#fbcfe8,#ec4899',f:[{k:'w',l:'Weight (kg)',t:'number'},{k:'h',l:'Height (cm)',t:'number'}],r:function(v){var w=+v.w,h=+v.h/100;if(!(w>0&&h>0))return'Fill both fields';var b=w/(h*h);return'BMI: '+b.toFixed(1)+' — '+(b<18.5?'Underweight':b<25?'Normal':b<30?'Overweight':'Obese')}},
{id:'unit',n:'Unit Converter',d:'Length, weight, temperature',i:'🔄',c:'Utility',g:'#a5f3fc,#06b6d4',f:[{k:'v',l:'Value',t:'number'},{k:'m',l:'Convert',t:'select',o:['km → miles','kg → lbs','°C → °F','cm → inches']}],r:function(v){var x=+v.v;if(isNaN(x)||v.v==='')return'Enter a value';var m=v.m;return x+' '+m.split(' → ')[0]+' = '+(m[0]=='k'&&m[1]=='m'?x*0.621371:m[1]=='g'?x*2.20462:m[0]=='°'?x*9/5+32:x/2.54).toFixed(2)+' '+m.split(' → ')[1]}},
{id:'pct',n:'Marks Percentage',d:'SSC, Inter groups & university',i:'📊',c:'Education',g:'#c7d2fe,#6366f1',u:mk},
{id:'imgc',n:'Image Compressor',d:'Reduce photo size, keep quality',i:'🖼️',c:'Utility',g:'#bbf7d0,#22c55e',u:imgc},
{id:'pdfc',n:'PDF Compressor',d:'Optimize PDF file size',i:'📄',c:'Utility',g:'#fecaca,#ef4444',u:pdfc},
{id:'tv',n:'Telugu → English Voice',d:'Speak Telugu, get English text',i:'🎙️',c:'Voice',g:'#e9d5ff,#a855f7',u:function(b){mic(b,1)}},
{id:'stt',n:'Speech to Text',d:'Speak & download as Word doc',i:'🗣️',c:'Voice',g:'#fef08a,#eab308',u:function(b){mic(b,0)}},
{id:'ff',n:'Form Filler',d:'Auto-fill PDF forms',i:'🧾',c:'Utility',g:'#99f6e4,#14b8a6',u:ff},
{id:'qa',n:'Question Paper Solver',d:'Upload paper, get answers + explanations',i:'🧠',c:'Education',g:'#fbcfe8,#d946ef',u:qa},
{id:'pg',n:'AI Prompt Generator',d:'Turn an idea into a pro prompt',i:'✨',c:'AI',g:'#fde68a,#f472b6',u:pgen},
{id:'img',n:'AI Photo Generator',d:'Realistic photo from your prompt',i:'🎨',c:'AI',g:'#c4b5fd,#ec4899',u:igen},
{id:'th',n:'Thumbnail Maker',d:'YouTube-style thumbnails in seconds',i:'🖼️',c:'Creator',g:'#fdba74,#ef4444',u:thg},
{id:'rsa',n:'Response Sheet Analyser',d:'Auto marks by subject - no AI',i:'📈',c:'Education',g:'#bae6fd,#0ea5e9',u:rsa}
];
var cur='All';
var $=function(i){return document.getElementById(i)};
function draw(){
var q=$('q').value.toLowerCase(),L=tiles().filter(function(t){return(cur=='All'||t.c==cur)&&(t.n+t.d).toLowerCase().indexOf(q)>-1});
$('grid').innerHTML=L.map(function(t,i){var g=t.g.split(',');if(t.id=='ppt')return'<div class="card wide" style="animation-delay:'+i*60+'ms" onclick="openT(\'ppt\')"><span class="fl">📊</span><span class="fl" style="right:120px;animation-delay:1s">✨</span><div class="ic">'+esc(t.i)+'</div><div><h3>'+esc(t.n)+'</h3><p>'+esc(t.d)+'</p><span class="tag" style="background:rgba(255,255,255,.24);color:#fff">'+esc(t.c)+' · 20 designs · Morph animations · Free</span></div><div class="bt">Create PPT →</div></div>';return'<div class="card'+(/^(pg|img|rsa)$/.test(t.id)?' hot':'')+'" style="animation-delay:'+i*60+'ms" onclick="openT(\''+t.id+'\')"><div class="ic" style="background:linear-gradient(135deg,'+g[0]+','+g[1]+')">'+esc(t.i)+'</div><h3>'+esc(t.n)+'</h3><p>'+esc(t.d)+'</p><span class="tag">'+esc(t.c)+'</span></div>'}).join('');
$('em').style.display=L.length?'none':'block'}
function chips(){var cs=['All'];tiles().forEach(function(t){if(cs.indexOf(t.c)<0)cs.push(t.c)});$('chips').innerHTML=cs.map(function(c){return'<button class="chip'+(c==cur?' on':'')+'" data-c="'+esc(c)+'">'+esc(c)+'</button>'}).join('')}
$('chips').onclick=function(e){var c=e.target.dataset.c;if(!c)return;cur=c;[].forEach.call(document.querySelectorAll('.chip'),function(b){b.classList.toggle('on',b.dataset.c==c)});draw()};
$('q').oninput=draw;
function openT(id){var t=tiles().filter(function(x){return x.id==id})[0];if(!t)return;if(t.h){window.open(t.h,'_blank','noopener');return}
if(t.u){$('sh').innerHTML='<h2>'+t.i+' '+t.n+'<button class="x" onclick="closeT()">✕</button></h2><div id="bx"></div>';$('ov').classList.add('on');t.u($('bx'));return}
var h='<h2>'+t.i+' '+t.n+'<button class="x" onclick="closeT()">✕</button></h2>';
t.f.forEach(function(f){h+='<label>'+f.l+'</label>';
if(f.t=='area')h+='<textarea rows="5" id="f_'+f.k+'"></textarea>';
else if(f.t=='select')h+='<select id="f_'+f.k+'">'+f.o.map(function(o){return'<option>'+o+'</option>'}).join('')+'</select>';
else h+='<input id="f_'+f.k+'" type="'+f.t+'" inputmode="decimal" value="'+(f.v||'')+'">'});
h+='<button class="go" id="go">Calculate</button><div class="res" id="res"></div>';
$('sh').innerHTML=h;$('ov').classList.add('on');
$('go').onclick=function(){var v={};t.f.forEach(function(f){v[f.k]=$('f_'+f.k).value});var r=$('res');show(r,t.r(v))}}
function closeT(){$('sh').style.maxWidth='';if(window.rec)try{rec.stop()}catch(e){}$('ov').classList.remove('on')}
$('ov').onclick=function(e){if(e.target.id=='ov')closeT()};
$('tg').onclick=function(){var r=document.documentElement,d=r.dataset.theme?r.dataset.theme=='dark':matchMedia('(prefers-color-scheme:dark)').matches;r.dataset.theme=d?'light':'dark';try{localStorage.setItem('th',r.dataset.theme)}catch(e){}};
try{var s=localStorage.getItem('th');if(s)document.documentElement.dataset.theme=s}catch(e){}
function esc(s){return String(s).replace(/[&<>"]/g,function(c){return'&#'+c.charCodeAt(0)+';'})}
function kb(n){return(n/1024).toFixed(1)+' KB'}
function dl(b,n){var a=document.createElement('a');a.href=URL.createObjectURL(b);a.download=n;a.click()}
function boom(r){r.style.position='relative';for(var i=0;i<14;i++){var e=document.createElement('i');e.className='cf';e.textContent=['🎉','✨','⭐','🎊'][i%4];e.style.left=(5+i*7)+'%';e.style.animationDelay=i*40+'ms';r.appendChild(e);setTimeout(function(x){return function(){x.remove()}}(e),1500)}}
function show(r,s){r.style.display='block';r.classList.remove('pop');void r.offsetWidth;r.classList.add('pop');
var m=s.match(/\d+(\.\d+)?/),p=m?parseFloat(m[0]):NaN;if(isNaN(p)||!/\s/.test(s)||s.length>150){r.textContent=s;return}
var d=(m[1]||'').length,t0=performance.now(),pc=/%/.test(s)&&p<=100,bar=pc?'<div class="bar"><b id="bf"></b></div>':'';
function f(n){var e=Math.min(1,(n-t0)/900),v=p*(1-Math.pow(1-e,3));r.innerHTML=s.replace(m[0],'<span class="big">'+v.toFixed(d)+'</span>')+bar;var q=document.getElementById('bf');if(q)q.style.width=v+'%';if(e<1)requestAnimationFrame(f);else if(pc&&p>=60)boom(r)}
requestAnimationFrame(f)}
var CF={name:'ToolHub',c1:'#2563eb',c2:'#7c3aed',ov:{},ct:[],logo:'/logo.png',tag:'All-in-one tools for everyone',tk:['Welcome to ToolHub','New: Image & PDF compressor','Telugu voice translator is live'],lk:[{t:'Home',u:'#'},{t:'Telegram',u:'https://t.me/'}]};
function tick(){var h=CF.tk.map(function(x){return'<a>📢 '+esc(x)+'</a>'}).join('');$('tm').innerHTML=h+h}
function applyCfg(){var s=document.documentElement.style;s.setProperty('--ac',CF.c1);s.setProperty('--ac2',CF.c2);$('nm').textContent=CF.name;document.title=CF.name;var li=document.querySelector('.logo i');if(CF.logo){li.innerHTML='<img src="'+esc(CF.logo)+'" alt="">';li.style.background='none'}else{li.textContent='⚡';li.style.background=''}var fv=document.querySelector('link[rel=icon]');if(!fv){fv=document.createElement('link');fv.rel='icon';document.head.appendChild(fv)}fv.href=CF.logo||'data:,';tick();
$('lk').innerHTML=CF.lk.map(function(l){return/^(https?:\/\/|#)/.test(l.u)?'<a class="pl" target="_blank" rel="noopener" href="'+l.u.replace(/"/g,'')+'">'+esc(l.t)+'</a>':''}).join('')}
var J={'Content-Type':'application/json'};
function api(u,o){o=o||{};o.credentials='same-origin';return fetch(u,o).then(function(r){return r.json().then(function(j){if(!r.ok)throw j.error||'Error';return j})})}
function tiles(){var r=[];T.forEach(function(t){var o=CF.ov[t.id]||{};if(o.hide)return;var x=Object.assign({},t);['n','d','i','c'].forEach(function(k){if(o[k])x[k]=o[k]});r.push(x)});return r.concat(CF.ct)}
$('ad').onclick=function(){api('/api/admin/me').then(function(){adm(JSON.parse(JSON.stringify(CF)))}).catch(login)};
function login(){$('sh').innerHTML='<h2>🔐 Admin Login<button class="x" onclick="closeT()">✕</button></h2><label>Password</label><input id="pw" type="password"><button class="go" id="lgn">Login</button><div class="res" id="rl"></div>';$('ov').classList.add('on');
$('lgn').onclick=function(){api('/api/admin/login',{method:'POST',headers:J,body:JSON.stringify({pw:$('pw').value})}).then(function(){adm(JSON.parse(JSON.stringify(CF)))}).catch(function(e){show($('rl'),String(e))})}}
function col(){var E={name:$('a1').value||'ToolHub',c1:$('a2').value,c2:$('a3').value,tk:$('a4').value.split('\n').filter(Boolean),lk:$('a5').value.split('\n').filter(Boolean).map(function(x){var p=x.split('|');return{t:p[0].trim(),u:(p[1]||'#').trim()}}),tag:$('a7').value,logo:$('a8').value,ov:{},ct:[]};
[].forEach.call(document.querySelectorAll('.ar'),function(r){var g=function(f){return r.querySelector('[data-f='+f+']').value},o={i:g('i'),n:g('n'),d:g('d'),c:g('c')};
if(r.dataset.b){o.hide=r.querySelector('.hd').checked;E.ov[r.dataset.id]=o}else{o.id=r.dataset.id;o.g=r.dataset.g;o.h=r.dataset.h;E.ct.push(o)}});return E}
function adm(E){E.ov=E.ov||{};E.ct=E.ct||[];
var h='<h2>⚙️ Admin Panel<button class="x" onclick="closeT()">✕</button></h2><label>Site name</label><input id="a1" value="'+esc(E.name)+'"><label>Main colour</label><input id="a2" type="color" value="'+E.c1+'"><label>Second colour</label><input id="a3" type="color" value="'+E.c2+'"><label>Scrolling news (one per line)</label><textarea id="a4" rows="3">'+esc(E.tk.join('\n'))+'</textarea><label>Links under search (Title | URL per line)</label><textarea id="a5" rows="3">'+esc(E.lk.map(function(l){return l.t+' | '+l.u}).join('\n'))+'</textarea><label>Tagline (shown in PDFs and PPT)</label><input id="a7" value="'+esc(E.tag||'')+'"><label>Logo</label><div style="display:flex;align-items:center;gap:10px;flex-wrap:wrap"><img id="lgp" src="'+esc(E.logo||'')+'" alt="" style="height:56px;width:56px;object-fit:contain;border:1px solid var(--bd);border-radius:12px;padding:4px;display:'+(E.logo?'block':'none')+'"><button class="chip" id="lgu">⬆ Upload logo</button><button class="chip" id="lgr">🗑 Remove logo</button><button class="chip" id="lgd">↺ Use original</button></div><input type="file" id="lgf" accept="image/*" style="display:none"><input type="hidden" id="a8" value="'+esc(E.logo||'')+'"><label>Tiles: icon, name, description, category</label>';
function ar(id,o,b,x){return'<div class="ar" data-id="'+id+'"'+(b?' data-b="1"':' data-g="'+esc(x.g)+'" data-h="'+esc(x.h)+'"')+'><div class="r2"><input data-f="i" value="'+esc(o.i)+'"><input data-f="n" value="'+esc(o.n)+'"></div><input data-f="d" value="'+esc(o.d)+'"><input data-f="c" value="'+esc(o.c)+'">'+(b?'<label><input type="checkbox" class="hd"'+(o.hide?' checked':'')+'> Hide this tile</label>':'<small>'+esc(x.h)+'</small><button class="x dl" style="width:auto;padding:0 12px">🗑 Remove tile</button>')+'</div>'}
T.forEach(function(t){var o=E.ov[t.id]||{};h+=ar(t.id,{i:o.i||t.i,n:o.n||t.n,d:o.d||t.d,c:o.c||t.c,hide:o.hide},1)});
E.ct.forEach(function(t){h+=ar(t.id,t,0,t)});
h+='<button class="go" id="al" style="background:#0ea5e9">➕ Add link tile</button><button class="go" id="ap" style="background:#ef4444">➕ Add PDF tile</button><input type="file" id="apf" accept="application/pdf" style="display:none"><button class="go" id="sv">💾 Save all changes</button><button class="go" id="lo" style="background:#64748b">Logout</button><div class="res" id="ra"></div>';
$('sh').innerHTML=h;$('ov').classList.add('on');
function sl2(u){$('a8').value=u;$('lgp').src=u;$('lgp').style.display=u?'block':'none'}
$('lgu').onclick=function(){$('lgf').click()};
$('lgf').onchange=function(){var f=this.files[0];if(!f)return;var fd=new FormData();fd.append('file',f);show($('ra'),'Uploading logo...');api('/api/admin/logo',{method:'POST',body:fd}).then(function(j){sl2(j.url);$('ra').style.display='none'}).catch(function(e){show($('ra'),'Error: '+e)})};
$('lgr').onclick=function(){sl2('')};$('lgd').onclick=function(){sl2('/logo.png')};
function nt(x){var E2=col();E2.ct.push(Object.assign({id:'c'+Date.now().toString(36),n:'New tile',d:'Edit me',i:'🔗',c:'Links',g:'#bfdbfe,#3b82f6',h:'#'},x));adm(E2)}
$('al').onclick=function(){var u=prompt('Link URL (https://...)');if(u)nt({h:u})};
$('ap').onclick=function(){$('apf').click()};
$('apf').onchange=function(){var f=this.files[0];if(!f)return;var fd=new FormData();fd.append('file',f);show($('ra'),'Uploading...');
api('/api/admin/upload',{method:'POST',body:fd}).then(function(j){nt({h:j.url,i:'📄',c:'PDFs',g:'#fecaca,#ef4444',n:f.name.replace(/\.pdf$/i,'').slice(0,40),d:'Tap to open PDF'})}).catch(function(e){show($('ra'),'Error: '+e)})};
[].forEach.call(document.querySelectorAll('.dl'),function(b){b.onclick=function(){b.parentNode.remove()}});
$('sv').onclick=function(){api('/api/admin/config',{method:'POST',headers:J,body:JSON.stringify(col())}).then(function(j){CF=Object.assign(CF,j);applyCfg();chips();draw();closeT()}).catch(function(e){show($('ra'),'Error: '+e)})};
$('lo').onclick=function(){api('/api/admin/logout',{method:'POST'}).then(closeT)}}
function pgen(b){b.innerHTML='<label>Your idea</label><textarea id="pi" rows="3" placeholder="e.g. old man selling tea at sunrise"></textarea><label>Style</label><select id="ps">'+o(['Ultra realistic photo','Cinematic','Portrait','Product photo','Anime','3D render','Digital painting'])+'</select><button class="go" id="g7">✨ Generate prompt</button><div class="res" id="r7"></div><textarea id="po" rows="6" style="display:none;margin-top:10px"></textarea><button class="go" id="c7" style="display:none;background:#64748b">📋 Copy prompt</button><button class="go" id="m7" style="display:none;background:#10b981">🎨 Create photo from this prompt</button>';
$('g7').onclick=function(){var g=this;g.disabled=true;g.textContent='Writing...';api('/api/prompt',{method:'POST',headers:J,body:JSON.stringify({idea:$('pi').value,style:$('ps').value})}).then(function(j){$('r7').style.display='none';var p=$('po');p.value=j.prompt;['po','c7','m7'].forEach(function(i){$(i).style.display='block'})}).catch(function(e){show($('r7'),'Error: '+e)}).then(function(){g.disabled=false;g.textContent='✨ Generate prompt'})};
$('c7').onclick=function(){try{navigator.clipboard.writeText($('po').value);this.textContent='✅ Copied'}catch(e){$('po').select()}};
$('m7').onclick=function(){window.pp=$('po').value;openT('img')}}
function shrink(f){return new Promise(function(ok){var im=new Image();im.onload=function(){var k=Math.min(1,1280/Math.max(im.width,im.height)),c=document.createElement('canvas');c.width=im.width*k;c.height=im.height*k;c.getContext('2d').drawImage(im,0,0,c.width,c.height);c.toBlob(ok,'image/jpeg',.88)};im.onerror=function(){ok(null)};im.src=URL.createObjectURL(f)})}
function igen(b){b.innerHTML='<label>Describe the photo you want</label><textarea id="ip" rows="4" placeholder="Describe what you want to see"></textarea><label>📷 Upload your photo(s) - optional, up to 3</label><input type="file" id="ipf" accept="image/*" multiple><div id="pv" style="display:flex;gap:8px;flex-wrap:wrap;margin-top:8px"></div><p style="font-size:12px;color:var(--mu);margin:8px 0 0">With a photo, describe the change, e.g. "me in a black suit at a beach sunset".</p><button class="go" id="g8">🎨 Generate photo</button><div class="res" id="r8"></div><img id="im" alt="Generated photo" style="width:100%;border-radius:14px;margin-top:12px;display:none"><button class="go" id="d8" style="display:none;background:#64748b">⬇ Download image</button>';
if(window.pp){$('ip').value=window.pp;window.pp=''}
$('ipf').onchange=function(){var v=$('pv');v.innerHTML='';[].slice.call(this.files,0,3).forEach(function(f){var m=document.createElement('img');m.src=URL.createObjectURL(f);m.style.cssText='width:72px;height:72px;object-fit:cover;border-radius:12px';v.appendChild(m)})};
$('g8').onclick=function(){var p=$('ip').value.trim();if(!p)return show($('r8'),'Describe the photo first');var g=this,fs=[].slice.call($('ipf').files,0,3);g.disabled=true;g.textContent='Creating... (up to 30s)';
Promise.all(fs.map(shrink)).then(function(bl){var fd=new FormData();fd.append('prompt',p);bl.forEach(function(x){if(x)fd.append('photos',x,'p.jpg')});return api('/api/image',{method:'POST',body:fd})}).then(function(j){$('r8').style.display='none';var m=$('im');m.src=j.image;m.style.display='block';$('d8').style.display='block';$('d8').onclick=function(){var a=document.createElement('a');a.href=j.image;a.download='photo.png';a.click()}}).catch(function(e){show($('r8'),'Error: '+e)}).then(function(){g.disabled=false;g.textContent='🎨 Generate photo'})}}
function thg(b){b.innerHTML='<label>Title</label><input id="tt1" value="Your Title Here"><label>Small text (optional)</label><input id="tt2"><label>Background photo (optional)</label><input type="file" id="tf" accept="image/*"><label>Colour theme</label><select id="tc">'+o(['Blue-Purple','Red-Orange','Green-Teal','Black-Gold'])+'</select><button class="go" id="g9">Create thumbnail</button><canvas id="cv" width="1280" height="720" style="width:100%;border-radius:14px;margin-top:12px;display:none"></canvas><button class="go" id="d9" style="display:none;background:#64748b">⬇ Download PNG</button>';
$('g9').onclick=function(){var cv=$('cv'),x=cv.getContext('2d'),th={'Blue-Purple':['#2563eb','#7c3aed'],'Red-Orange':['#ef4444','#f59e0b'],'Green-Teal':['#10b981','#06b6d4'],'Black-Gold':['#0f172a','#d97706']}[$('tc').value];
function fin(){var gr=x.createLinearGradient(0,720,0,150);gr.addColorStop(0,'rgba(0,0,0,.8)');gr.addColorStop(1,'rgba(0,0,0,0)');x.fillStyle=gr;x.fillRect(0,0,1280,720);
x.textAlign='center';x.lineJoin='round';x.font='900 100px "Noto Sans Telugu",Arial,sans-serif';var ls=[],c='';$('tt1').value.split(' ').forEach(function(w){var t=c?c+' '+w:w;if(x.measureText(t).width>1120&&c){ls.push(c);c=w}else c=t});ls.push(c);
ls.forEach(function(l,i){var y=570-(ls.length-1-i)*115;x.lineWidth=14;x.strokeStyle='#000';x.strokeText(l,640,y);x.fillStyle='#fff';x.fillText(l,640,y)});
if($('tt2').value){x.font='700 50px "Noto Sans Telugu",Arial,sans-serif';x.fillStyle='#fde047';x.fillText($('tt2').value,640,670)}
cv.style.display='block';$('d9').style.display='block';$('d9').onclick=function(){cv.toBlob(function(q){dl(q,'thumbnail.png')})}}
var f=$('tf').files[0];
if(f){var im=new Image();im.onload=function(){var k=Math.max(1280/im.width,720/im.height),w=im.width*k,h=im.height*k;x.drawImage(im,(1280-w)/2,(720-h)/2,w,h);fin()};im.src=URL.createObjectURL(f)}
else{var g=x.createLinearGradient(0,0,1280,720);g.addColorStop(0,th[0]);g.addColorStop(1,th[1]);x.fillStyle=g;x.fillRect(0,0,1280,720);fin()}}}
function o(a){return a.map(function(x){return'<option>'+x+'</option>'}).join('')}
function mk(b){var U={'JNTUH':.75,'JNTUK':.75,'JNTUA':.75,'Osmania (OU)':0,'Kakatiya (KU)':0,'Telangana Univ':0,'Andhra Univ':0,'Acharya Nagarjuna':0,'SVU':0,'Other':0};
b.innerHTML='<label>Student name (shown in PDF)</label><input id="sn" placeholder="Your name"><label>Level</label><select id="lv">'+o(['SSC (10th)','Inter (Intermediate)','Degree'])+'</select><div id="ex"></div><div id="sbw"><label>Subject-wise marks (optional)</label><div id="sb"></div><button class="go" id="as" style="background:#0ea5e9;margin-top:8px">➕ Add subject</button></div><label>Total marks (editable)</label><input id="tt" type="number" inputmode="decimal"><label id="ol">Marks obtained</label><input id="ob" type="number" inputmode="decimal"><button class="go" id="g4">Calculate</button><div class="res" id="r4"></div><button class="go" id="d4" style="display:none;background:#64748b">⬇ Download PDF report</button><p style="font-size:12px;color:var(--mu)" id="nt"></p>';
function set(){var l=$('lv').value[0],t=$('tt');if(l=='S')t.value=600;if(l=='I')t.value=$('yr').value.indexOf('Both')==0?1000:500;if(l=='D')t.value='';
var c=l=='D'&&$('md').value=='CGPA';$('ol').textContent=c?'CGPA (out of 10)':'Marks obtained';t.disabled=c;$('sbw').style.display=c?'none':'block';$('nt').textContent=c?'CGPA formula differs by university - confirm with your university circular.':'Default totals can differ by board/year - edit the total if yours is different. If you add subjects, totals are calculated from them.'}
function ui(){var l=$('lv').value[0],h='';
if(l=='I')h='<label>Group</label><select id="gp">'+o(['MPC','BiPC','CEC','HEC','MEC','Other'])+'</select><label>Year</label><select id="yr">'+o(['Both years (1000)','1st year (500)','2nd year (500)'])+'</select>';
if(l=='D')h='<label>University</label><select id="un">'+o(Object.keys(U))+'</select><label>Method</label><select id="md">'+o(['Marks','CGPA'])+'</select>';
$('ex').innerHTML=h;set();if($('yr'))$('yr').onchange=set;if($('md'))$('md').onchange=set}
$('lv').onchange=ui;ui();
$('as').onclick=function(){var d=document.createElement('div');d.className='r3';d.innerHTML='<input placeholder="Subject"><input type="number" inputmode="decimal" placeholder="Marks"><input type="number" inputmode="decimal" placeholder="Max" value="100">';$('sb').appendChild(d)};
$('g4').onclick=function(){var l=$('lv').value[0],r=$('r4'),cg=l=='D'&&$('md').value=='CGPA',x=+$('ob').value,t=+$('tt').value,R=[],m={name:$('sn').value.trim(),lv:$('lv').value};
if(l=='I')m.lv+=' - '+$('gp').value+' ('+$('yr').value.replace(/ \(.*/,'')+')';
if(l=='D')m.lv+=' - '+$('un').value;
if(cg){if(!(x>0&&x<=10))return show(r,'Enter CGPA 0-10');var u=$('un').value,p=Math.max(0,(x-U[u])*10);m.cg=x;m.pct=p;m.R=[];window.mkr=m;show(r,u+': '+p.toFixed(2)+'%');$('d4').style.display='block';return}
[].forEach.call(document.querySelectorAll('#sb .r3'),function(e){var v=e.children[1].value;if(v!=='')R.push({n:e.children[0].value.trim()||'Subject',m:+v,x:+e.children[2].value||100})});
if(R.length){if(R.some(function(q){return q.m<0||q.m>q.x}))return show(r,'A subject mark is more than its maximum');x=R.reduce(function(a,q){return a+q.m},0);t=R.reduce(function(a,q){return a+q.x},0);$('ob').value=x;$('tt').value=t}
if(!(t>0)||x<0||x>t)return show(r,'Check marks and total');
m.R=R;m.x=x;m.t=t;m.pct=x/t*100;window.mkr=m;show(r,'Percentage: '+m.pct.toFixed(2)+'%');$('d4').style.display='block'};
$('d4').onclick=function(){mkpdf(window.mkr)}}
function mkpdf(r){var c1=CF.c1,c2=CF.c2,p=r.pct,msg=p>=75?'Excellent work! 🌟':p>=60?'Good job! 👍':p>=40?'Keep improving! 💪':'Work harder - you can do it! 🔥',td='padding:12px 16px;',
rows=(r.R||[]).map(function(q,i){var z=Math.min(100,q.m/q.x*100);return'<tr style="background:'+(i%2?'#fff':'#f4f6fb')+'"><td style="'+td+'font-weight:700">'+esc(q.n)+'</td><td style="'+td+'text-align:center">'+q.m+' / '+q.x+'</td><td style="'+td+'width:190px"><div style="height:12px;border-radius:9px;background:#e5e9f2"><div style="height:12px;border-radius:9px;width:'+z.toFixed(0)+'%;background:linear-gradient(90deg,'+c1+','+c2+')"></div></div></td><td style="'+td+'text-align:right;font-weight:700;color:'+c1+'">'+z.toFixed(0)+'%</td></tr>'}).join(''),
tb=rows?'<table style="width:100%;border-collapse:collapse;border:1px solid #e5e9f2;margin-top:24px"><tr style="background:linear-gradient(90deg,'+c1+','+c2+');color:#fff"><th style="'+td+'text-align:left">Subject</th><th style="'+td+'">Marks</th><th style="'+td+'">Progress</th><th style="'+td+'text-align:right">%</th></tr>'+rows+'</table>':'',
sm=r.cg!=null?'CGPA: <b>'+r.cg+'</b>':'Total: <b>'+r.x+' / '+r.t+'</b>';
pout('<div style="padding:30px 40px"><div style="display:flex;align-items:center;justify-content:space-between;gap:20px"><div><div style="font-size:27px;font-weight:800;color:'+c1+'">📊 Marks Report</div>'+(r.name?'<div style="font-size:23px;font-weight:700;margin-top:6px">'+esc(r.name)+'</div>':'')+'<div style="display:inline-block;margin-top:8px;background:#eef2ff;color:'+c1+';padding:5px 14px;border-radius:99px;font-size:15px;font-weight:700">'+esc(r.lv)+'</div><div style="margin-top:10px;font-size:20px">'+sm+'</div></div><div style="flex:none;width:165px;height:165px;border-radius:50%;background:linear-gradient(135deg,'+c1+','+c2+');color:#fff;display:flex;flex-direction:column;align-items:center;justify-content:center;box-shadow:0 8px 24px rgba(0,0,0,.18)"><div style="font-size:42px;font-weight:800;line-height:1">'+p.toFixed(1)+'%</div><div style="font-size:13px;letter-spacing:1px;margin-top:4px">PERCENTAGE</div></div></div>'+tb+'<div style="margin-top:26px;text-align:center;font-size:23px;font-weight:800;color:'+c2+'">'+msg+'</div>'+cta()+'</div>','marks-report.pdf')}
function rbar(q,h){var t=q.total||1;return'<div style="display:flex;height:'+h+'px;border-radius:9px;overflow:hidden;background:#e5e9f2"><div style="width:'+q.correct/t*100+'%;background:#22c55e"></div><div style="width:'+q.wrong/t*100+'%;background:#ef4444"></div><div style="width:'+q.not_attempted/t*100+'%;background:#94a3b8"></div></div>'}
function rsa(b){var F=null;window.rsr=null;window.rph=null;
b.innerHTML='<div id="dz" style="border:2px dashed var(--ac);border-radius:18px;padding:22px;text-align:center;cursor:pointer;background:rgba(37,99,235,.06)"><div style="font-size:34px">📄</div><div id="dzt"><b>Drag &amp; drop your response sheet PDF</b><br><small style="color:var(--mu)">or tap to choose a file</small></div></div><input type="file" id="rf" accept="application/pdf" style="display:none"><div style="text-align:center;color:var(--mu);margin:10px 0;font-size:13px">— OR —</div><input id="ru" type="url" placeholder="🔗 Paste response sheet link (https://...)"><div class="r2" style="grid-template-columns:1fr 1fr"><div><label>Marks per correct (+)</label><input id="pos" type="number" value="4" step="any"></div><div><label>Negative per wrong (−)</label><input id="neg" type="number" value="1" step="any"></div></div><label>Your name (shown in the PDF)</label><input id="rn" placeholder="Auto-filled from the sheet if found"><label>Your photo (optional - picked from your sheet automatically if found)</label><input type="file" id="rph" accept="image/*"><button class="go" id="g5">🧮 Calculate my marks</button><div class="res" id="r9"></div><div id="ao"></div><button class="go" id="d5" style="display:none;background:#64748b">⬇ Download PDF report</button><p style="font-size:12px;color:var(--mu)">Works with sheets that list Question ID, Chosen Option and Correct Option. Links must be public. Scanned (image-only) PDFs cannot be read.</p>';
var dz=$('dz'),pk=function(f){if(f&&/\.pdf$/i.test(f.name)){F=f;$('dzt').innerHTML='✅ <b>'+esc(f.name)+'</b>'}else show($('r9'),'Please choose a PDF file')};
dz.onclick=function(){$('rf').click()};$('rf').onchange=function(){pk(this.files[0])};
dz.ondragover=function(e){e.preventDefault();dz.style.background='rgba(37,99,235,.2)'};dz.ondragleave=function(){dz.style.background=''};
dz.ondrop=function(e){e.preventDefault();dz.style.background='';pk(e.dataTransfer.files[0])};
function rr(an){var R=window.rsr;if(!R)return;var p=+$('pos').value||0,n=+$('neg').value||0,G=R.grand_total,rd=function(v){return Math.round(v*100)/100};R.pos=p;R.neg=n;R.subjects.forEach(function(q){q.marks=rd(q.correct*p-q.wrong*n)});G.marks=rd(G.correct*p-G.wrong*n);
$('ao').innerHTML='<div class="gt">'+((window.rph||R.photo)?'<img src="'+(window.rph||R.photo)+'" alt="" style="width:64px;height:64px;border-radius:50%;object-fit:cover;border:3px solid #fff;margin-bottom:6px">':'')+'<div style="font-size:13px;letter-spacing:1px">🏆 YOUR TOTAL MARKS</div><div class="m"><span id="gm">0</span><span style="font-size:20px;opacity:.85"> / '+rd(G.total*p)+'</span></div><div style="font-size:12px;opacity:.9">Marking: +'+p+' / −'+n+'</div><div class="st"><div><b id="gc">0</b>✅ Correct</div><div><b id="gw">0</b>❌ Wrong</div><div><b id="gn">0</b>⏭ Skipped</div></div></div><div class="sg">'+R.subjects.map(function(q,i){return'<div class="sk" style="animation-delay:'+(an?i*80:0)+'ms"><h4>'+esc(q.name)+'</h4><div class="mk">'+q.marks+'</div>'+rbar(q,10)+'<small>'+q.total+' questions</small><small>✅ '+q.correct+' · ❌ '+q.wrong+' · ⏭ '+q.not_attempted+'</small></div>'}).join('')+'</div>'+(R.note?'<small style="display:block;margin-top:10px;color:var(--mu)">'+esc(R.note)+'</small>':'');
[['gm',G.marks],['gc',G.correct],['gw',G.wrong],['gn',G.not_attempted]].forEach(function(x){if(an)cnt($(x[0]),x[1]);else $(x[0]).textContent=x[1]})}
$('pos').oninput=$('neg').oninput=function(){rr(0)};
$('d5').onclick=function(){rapdf(window.rsr)};
$('rph').onchange=function(){window.rph=null;var f=this.files[0];if(f)fimg(f,function(u){window.rph=u;rr(0)},300)};
$('g5').onclick=function(){var u=$('ru').value.trim();if(!F&&!u)return show($('r9'),'Upload the PDF or paste the link first');var fd=new FormData(),g=this;if(F)fd.append('file',F);else fd.append('url',u);fd.append('pos',$('pos').value);fd.append('neg',$('neg').value);g.disabled=true;g.textContent='Calculating...';
api('/response-calculator',{method:'POST',body:fd}).then(function(R){$('r9').style.display='none';if(!$('rn').value&&R.candidate)$('rn').value=R.candidate;window.rsr=R;rr(1);$('d5').style.display='block'}).catch(function(e){show($('r9'),'Error: '+e)}).then(function(){g.disabled=false;g.textContent='🧮 Calculate my marks'})}}
function ins(R){var G=R.grand_total,S=R.subjects.map(function(q){var a=q.correct+q.wrong;return{n:q.name,m:q.marks,acc:a?q.correct/a*100:0}}),o=[],att=G.correct+G.wrong,acc=att?G.correct/att*100:0,rate=G.total?att/G.total*100:0,lost=Math.round(G.wrong*R.neg*100)/100;
o.push('You attempted '+att+' of '+G.total+' questions ('+rate.toFixed(0)+'%) with '+acc.toFixed(0)+'% accuracy.');
if(S.length>1){var b=S.slice().sort(function(x,y){return y.m-x.m})[0],w=S.slice().sort(function(x,y){return x.acc-y.acc})[0];o.push('Strongest subject: '+b.n+' ('+b.m+' marks). Needs the most attention: '+w.n+' ('+w.acc.toFixed(0)+'% accuracy).')}
if(G.wrong>0&&R.neg>0)o.push('Wrong answers cost you '+lost+' marks. When you are not sure, skipping a question can protect your score.');
if(G.not_attempted>0)o.push(G.not_attempted+' questions were left unattempted - practising time management can help you reach them.');
o.push(acc>=85?'Excellent accuracy - keep up this consistency!':acc>=65?'Good accuracy - a little revision can push your score even higher.':'Focus on revising the basics and practise more mock tests to improve accuracy.');return o}
function rapdf(R){var c1=CF.c1,c2=CF.c2,G=R.grand_total,nm=($('rn').value||'').trim(),mx=G.total*R.pos,p=mx>0?Math.max(0,G.marks)/mx*100:0,ph=window.rph||R.photo,
dt=new Date().toLocaleDateString('en-IN',{day:'numeric',month:'long',year:'numeric'}),ini=(nm||'S').split(/\s+/).map(function(w){return w[0]||''}).join('').slice(0,2).toUpperCase(),
av=ph?'<img src="'+ph+'" style="width:132px;height:158px;object-fit:cover;border-radius:18px;border:4px solid '+c1+';display:block">':'<div style="width:132px;height:158px;border-radius:18px;background:linear-gradient(135deg,'+c1+','+c2+');color:#fff;font-size:54px;font-weight:800;display:flex;align-items:center;justify-content:center">'+esc(ini)+'</div>',
msg=p>=75?'Excellent performance! 🌟':p>=60?'Good job! 👍':p>=40?'Keep improving! 💪':'Work harder - you can do it! 🔥',
cd=function(t,v,c){return'<div style="flex:1;background:'+c+';color:#fff;border-radius:18px;padding:14px;text-align:center"><div style="font-size:13px;letter-spacing:.5px">'+t+'</div><div style="font-size:34px;font-weight:800;line-height:1.3">'+v+'</div></div>'},
sc=R.subjects.map(function(q){var a=q.correct+q.wrong,ac=a?q.correct/a*100:0,cl=ac>=75?'#22c55e':ac>=50?'#f59e0b':'#ef4444';return'<div style="width:48.5%;box-sizing:border-box;background:#fff;border:1px solid #e5e9f2;border-top:7px solid '+cl+';border-radius:18px;padding:16px;margin-bottom:14px;page-break-inside:avoid;break-inside:avoid"><div style="font-size:19px;font-weight:800">'+esc(q.name)+'</div><div style="font-size:38px;font-weight:800;color:'+c1+';line-height:1.25">'+q.marks+'<span style="font-size:15px;color:#64748b;font-weight:600"> / '+Math.round(q.total*R.pos*100)/100+'</span></div>'+rbar(q,10)+'<div style="display:flex;justify-content:space-between;font-size:15px;margin-top:9px"><span style="color:#16a34a;font-weight:700">✅ '+q.correct+'</span><span style="color:#dc2626;font-weight:700">❌ '+q.wrong+'</span><span style="color:#64748b;font-weight:700">⏭ '+q.not_attempted+'</span></div><div style="font-size:13px;color:#64748b;margin-top:4px">'+q.total+' questions · '+ac.toFixed(0)+'% accuracy</div></div>'}).join(''),
tips=ins(R).map(function(t){return'<div style="display:flex;gap:10px;margin-bottom:10px"><div style="color:'+c1+';font-weight:800">✔</div><div>'+esc(t)+'</div></div>'}).join('');
pout('<div style="padding:28px 40px"><div style="text-align:center;font-size:31px;font-weight:800;color:'+c1+'">📈 Response Sheet Analyser</div><div style="text-align:center;color:#64748b;font-size:14px;margin-top:2px">Report generated on '+dt+'</div>'
+'<div style="display:flex;align-items:center;gap:22px;margin-top:20px;background:#f4f6fb;border-radius:22px;padding:20px"><div style="flex:none">'+av+'</div><div style="flex:1"><div style="font-size:26px;font-weight:800">'+(nm?esc(nm):'Candidate')+'</div><div style="display:inline-block;margin-top:8px;background:#eef2ff;color:'+c1+';padding:5px 14px;border-radius:99px;font-size:15px;font-weight:700">Marking: +'+R.pos+' / −'+R.neg+'</div><div style="margin-top:8px;font-size:16px;color:#475569">'+G.total+' questions · '+p.toFixed(1)+'% score</div></div><div style="flex:none;width:158px;height:158px;border-radius:50%;background:linear-gradient(135deg,'+c1+','+c2+');color:#fff;display:flex;flex-direction:column;align-items:center;justify-content:center;box-shadow:0 8px 24px rgba(0,0,0,.18)"><div style="font-size:38px;font-weight:800;line-height:1">'+G.marks+'</div><div style="font-size:13px;margin-top:2px">out of '+Math.round(mx*100)/100+'</div><div style="font-size:10px;letter-spacing:1px;margin-top:4px">TOTAL MARKS</div></div></div>'
+'<div style="display:flex;gap:14px;margin-top:20px">'+cd('✅ CORRECT',G.correct,'#22c55e')+cd('❌ WRONG',G.wrong,'#ef4444')+cd('⏭ SKIPPED',G.not_attempted,'#64748b')+'</div>'
+'<div style="font-size:20px;font-weight:800;color:'+c1+';margin:24px 0 12px">📚 Subject-wise Performance</div><div style="display:flex;flex-wrap:wrap;justify-content:space-between">'+sc+'</div>'
+'<div style="margin-top:8px;border:2px solid '+c1+';border-radius:18px;padding:18px 22px;page-break-inside:avoid;break-inside:avoid;font-size:16px;line-height:1.55"><div style="font-size:20px;font-weight:800;color:'+c1+';margin-bottom:10px">🧠 Performance Analysis</div>'+tips+'</div>'
+'<div style="margin-top:20px;text-align:center;font-size:23px;font-weight:800;color:'+c2+'">'+msg+'</div><div style="page-break-inside:avoid;break-inside:avoid">'+cta()+'</div></div>','response-analyser.pdf')}
function gf(x,c,ang,X,Y,w,h){if(!Array.isArray(c))return c;var a=(ang||45)*Math.PI/180,dx=Math.cos(a),dy=-Math.sin(a),L=Math.abs(w*dx)+Math.abs(h*dy),cx=X+w/2,cy=Y+h/2,g=x.createLinearGradient(cx-dx*L/2,cy-dy*L/2,cx+dx*L/2,cy+dy*L/2);g.addColorStop(0,c[0]);g.addColorStop(1,c[1]);return g}
function rr2(x,X,Y,w,h,r){x.beginPath();x.moveTo(X+r,Y);x.arcTo(X+w,Y,X+w,Y+h,r);x.arcTo(X+w,Y+h,X,Y+h,r);x.arcTo(X,Y+h,X,Y,r);x.arcTo(X,Y,X+w,Y,r);x.closePath()}
var PG={tri:[[.5,0],[1,1],[0,1]],diamond:[[.5,0],[1,.5],[.5,1],[0,.5]],hex:[[.25,0],[.75,0],[1,.5],[.75,1],[.25,1],[0,.5]],cross:[[.33,0],[.67,0],[.67,.33],[1,.33],[1,.67],[.67,.67],[.67,1],[.33,1],[.33,.67],[0,.67],[0,.33],[.33,.33]]};
function pa(x,d,X,Y,w,h){x.beginPath();if(d.k=='oval')x.ellipse(X+w/2,Y+h/2,w/2,h/2,0,0,7);else if(d.k=='round')rr2(x,X,Y,w,h,Math.min(w,h)*.12);else if(d.k=='star'){var cx=X+w/2,cy=Y+h/2,R=w/2;for(var i=0;i<10;i++){var a=-Math.PI/2+i*Math.PI/5,q=i%2?R*.4:R;x[i?'lineTo':'moveTo'](cx+q*Math.cos(a),cy+q*Math.sin(a))}x.closePath()}else if(PG[d.k]){PG[d.k].forEach(function(p,i){x[i?'lineTo':'moveTo'](X+p[0]*w,Y+p[1]*h)});x.closePath()}else x.rect(X,Y,w,h)}
var FB={'Georgia':'serif','Courier New':'monospace'};
function tx(x,d,X,Y,w,h,k,W){var px=(d.sz||18)*W/960,lh=px*1.2,pl=(d.pl==null?.12:d.pl)*k,pr=.12*k,mw=w-pl-pr;x.font=(d.b?'700 ':'400 ')+px+'px "'+(d.ft||'Poppins')+'",'+(FB[d.ft]||'sans-serif');x.fillStyle=d.c||'#000';x.textBaseline='top';var L=[];
String(d.t).split('\n').forEach(function(p,j){var ln='';p.split(' ').forEach(function(wr){var tt=ln?ln+' '+wr:wr;if(x.measureText(tt).width>mw&&ln){L.push([ln,j]);ln=wr}else ln=tt});L.push([ln,j])});
var gap=8*W/960,np=L[L.length-1][1]+1,tot=L.length*lh+(np-1)*gap,yy=d.va=='m'?Y+(h-tot)/2:d.va=='b'?Y+h-tot-.06*k:Y+.06*k,pj=-1;
x.textAlign=d.ha=='c'?'center':d.ha=='r'?'right':'left';L.forEach(function(q){if(pj>=0&&q[1]!=pj)yy+=gap;pj=q[1];x.fillText(q[0],d.ha=='c'?X+pl+mw/2:d.ha=='r'?X+w-pr:X+pl,yy);yy+=lh})}
function ch2(x,d,X,Y,w,h,k){var D=d.data,mx=Math.max.apply(null,D.map(function(r){return r[1]}))||1,bw=w/D.length,f=14*k/72;x.font='700 '+f+'px "'+(d.ft||'Poppins')+'",sans-serif';x.textAlign='center';x.textBaseline='alphabetic';
D.forEach(function(r,i){var bh=(h-f*3.4)*Math.max(0,r[1])/mx,bx=X+i*bw+bw*.16,by=Y+h-f*1.9-bh;x.fillStyle=d.cc;x.fillRect(bx,by,bw*.68,bh);x.fillStyle=d.c;x.fillText(r[1],X+i*bw+bw/2,by-f*.4);x.fillText(String(r[0]).slice(0,12),X+i*bw+bw/2,Y+h-f*.5)})}
var IMC={};function pic(src){if(!IMC[src]){var i=new Image();i.onload=function(){if(window.rdre)window.rdre()};i.src=src;IMC[src]=i}return IMC[src]}
function rd(cv,L,t){if(!cv)return;var W=cv.width,Hh=cv.height,k=W/13.333,x=cv.getContext('2d');x.globalAlpha=1;x.fillStyle=gf(x,L.bg.c,L.bg.ang,0,0,W,Hh);x.fillRect(0,0,W,Hh);
L.shapes.forEach(function(d){var X=d.x*k,Y=d.y*k,w=d.w*k,h=d.h*k,e=1,al=1,a=d.a,cx=X+w/2,cy=Y+h/2;
if(t!=null&&a){e=Math.max(0,Math.min(1,(t-d.dl)/(a=='bounce'?.7:.5)));if(e<=0)return}
x.save();
if(t!=null&&a)al=/^(wipe|box|circle|diamond|plus|wedge|wheel)/.test(a)?1:a.indexOf('fly')==0?Math.min(1,e*3):a=='appear'?(e>0?1:0):e;
x.globalAlpha=al*(d.ga==null?1:d.ga);
if(e<1&&a){if(a=='fly_b')x.translate(0,(1-e)*Hh*.4);else if(a=='fly_t')x.translate(0,-(1-e)*Hh*.4);else if(a=='fly_l')x.translate(-(1-e)*W*.4,0);else if(a=='fly_r')x.translate((1-e)*W*.4,0);
else if(a=='bounce'){var u2=1-Math.pow(1-e,3);x.translate(0,(1-u2)*Hh*.4-Math.sin(e*Math.PI*2.4)*(1-e)*.08*Hh)}
else if(a=='zoom'||a=='pop'||a=='grow'){var z=Math.max(.01,a=='pop'?e*1.15:e);x.translate(cx,cy);x.scale(z,z);if(a=='grow')x.rotate((1-e)*Math.PI/2);x.translate(-cx,-cy)}
else if(a=='wipe_l'){x.beginPath();x.rect(X,Y,w*e,h);x.clip()}else if(a=='wipe_r'){x.beginPath();x.rect(X+w*(1-e),Y,w*e,h);x.clip()}else if(a=='wipe_u'){x.beginPath();x.rect(X,Y,w,h*e);x.clip()}else if(a=='wipe_d'){x.beginPath();x.rect(X,Y+h*(1-e),w,h*e);x.clip()}
else if(/^(box|circle|diamond|plus|wedge|wheel)$/.test(a)){x.beginPath();x.rect(X+w*(1-e)/2,Y+h*(1-e)/2,w*e,h*e);x.clip()}}
if(t!=null&&d.lo){var lt=t-(a?d.dl+.6:0);if(lt>0){x.translate(cx,cy);if(d.lo=='spin')x.rotate(lt*.7);else if(d.lo=='sway')x.rotate(.1*Math.sin(lt*2.8));else if(d.lo=='pulse'){var q=1+.05*Math.sin(lt*5.5);x.scale(q,q)}else x.translate(0,-Hh*.012*(1+Math.sin(lt*2.4)));x.translate(-cx,-cy)}}
if(d.rot){x.translate(cx,cy);x.rotate(d.rot*Math.PI/180);x.translate(-cx,-cy)}
if(d.k=='chart')ch2(x,d,X,Y,w,h,k);
else if(d.k=='pic'){var src=d.lg?CF.logo:(d.src||(window.pimg||[])[d.si]),im=src&&pic(src);x.save();rr2(x,X,Y,w,h,d.lg?1:Math.min(w,h)*.06);x.clip();
if(im&&im.complete&&im.naturalWidth){var s2=d.lg?Math.min(w/im.naturalWidth,h/im.naturalHeight):Math.max(w/im.naturalWidth,h/im.naturalHeight),dw=im.naturalWidth*s2,dh=im.naturalHeight*s2;x.drawImage(im,X+(w-dw)/2,Y+(h-dh)/2,dw,dh)}else if(!d.lg){x.fillStyle='rgba(128,128,128,.3)';x.fillRect(X,Y,w,h)}x.restore();
if(d.l&&!d.lg){rr2(x,X,Y,w,h,Math.min(w,h)*.06);x.lineWidth=(d.lw||1)*W/960;x.strokeStyle=d.l;x.stroke()}}
else{if(d.k!='text'){pa(x,d,X,Y,w,h);if(d.f!=null){x.save();x.globalAlpha*=(d.al==null?100:d.al)/100;x.fillStyle=gf(x,d.f,d.ang,X,Y,w,h);x.fill();x.restore()}if(d.l){x.lineWidth=(d.lw||1)*W/960;x.strokeStyle=d.l;x.stroke()}}
if(d.t!=null&&d.t!=='')tx(x,d,X,Y,w,h,k,W)}x.restore()})}
function fimg(f,cb,mx){mx=mx||1280;var im=new Image();im.onload=function(){var k=Math.min(1,mx/Math.max(im.width,im.height)),c=document.createElement('canvas');c.width=Math.round(im.width*k);c.height=Math.round(im.height*k);c.getContext('2d').drawImage(im,0,0,c.width,c.height);cb(c.toDataURL('image/jpeg',.84))};im.onerror=function(){cb(null)};im.src=URL.createObjectURL(f)}
function mph(A,B,u){var m={},used={};A.shapes.forEach(function(d){if(d.n&&d.n.indexOf('!!')==0)m[d.n]=d});var sh=B.shapes.map(function(d){var o=Object.assign({},d),p=d.n&&m[d.n];o.a=null;o.lo=null;if(p){used[d.n]=1;['x','y','w','h','rot'].forEach(function(q){o[q]=(p[q]||0)+((d[q]||0)-(p[q]||0))*u})}else o.ga=u;return o});
var out=A.shapes.filter(function(d){return !(d.n&&used[d.n])}).map(function(d){var o=Object.assign({},d);o.a=null;o.lo=null;o.ga=1-u;return o});return{bg:u<.5?A.bg:B.bg,shapes:out.concat(sh)}}
function ppt(b){$('sh').style.maxWidth='920px';var P={i:0,tok:0};window.rdre=function(){if($('pst')&&P.fr)P.fr()};
['Inter','Poppins','Montserrat'].forEach(function(f){if(document.fonts)document.fonts.load('700 20px "'+f+'"').then(function(){if(P.fr)P.fr()})});
b.innerHTML='<p style="font-size:13px;color:var(--mu)">Make a cinematic animated PowerPoint in 3 steps. Built with python-pptx - no AI. It uses your own text, or a clean template outline if you leave the content empty.</p><label>Topic title</label><input id="pt" placeholder="e.g. Rivers of India"><label>Your content (optional)</label><textarea id="pd" rows="5" placeholder="Paste your points here. Tip: end a line with : to start a new slide, e.g. Benefits:"></textarea><label>Number of slides: <b id="pnv">8</b></label><input type="range" id="pn" min="5" max="15" value="8"><p style="font-size:12px;color:var(--mu);margin:4px 0">An outline slide and a free promo slide for this website are added automatically.</p><button class="go" id="pg1">✨ Generate Designs</button><div class="res" id="pr"></div><div id="pst"></div>';
$('pn').oninput=function(){$('pnv').textContent=this.value};
$('pg1').onclick=function(){var g=this,tp=$('pt').value.trim();if(!tp)return show($('pr'),'Enter a topic first');g.disabled=true;g.textContent='Designing...';
api('/api/ppt/outline',{method:'POST',headers:J,body:JSON.stringify({topic:tp,desc:$('pd').value,n:+$('pn').value})}).then(function(j){$('pr').style.display='none';P.topic=tp;P.slides=j.slides;P.designs=j.designs;P.note=j.note;pds()}).catch(function(e){show($('pr'),'Error: '+e)}).then(function(){g.disabled=false;g.textContent='✨ Generate Designs'})};
function pds(){$('pst').innerHTML='<h3 style="margin:18px 0 0">Choose from '+P.designs.length+' designs</h3><div class="dg">'+P.designs.map(function(d,i){return'<div class="dc"><div class="tr2"><canvas width="320" height="180" data-i="'+i+'" data-s="0"></canvas><canvas width="320" height="180" data-i="'+i+'" data-s="1"></canvas></div><h4>'+(i+1)+'. '+esc(d.name)+'</h4><small>'+esc(d.tags)+'</small><button class="go" style="margin-top:10px;padding:10px" data-pick="'+i+'">Select this design</button></div>'}).join('')+'</div>'+(P.note?'<small style="display:block;margin-top:10px;color:var(--mu)">'+esc(P.note)+'</small>':'');
P.fr=function(){[].forEach.call($('pst').querySelectorAll('canvas'),function(c){rd(c,P.designs[+c.dataset.i].thumb[+c.dataset.s],null)})};P.fr();
[].forEach.call($('pst').querySelectorAll('[data-pick]'),function(bn){bn.onclick=function(){P.d=P.designs[+bn.dataset.pick];pedit()}})}
function syncimg(){window.pimg=P.slides.map(function(s){return s.imgd||null})}
function pcol(bd){return P.slides.map(function(s,i){var r=$('pe').querySelector('[data-i="'+i+'"]'),ln=r.querySelector('textarea').value.split('\n').map(function(z){return z.trim()}).filter(Boolean),o=Object.assign({},s,{title:r.querySelector('input').value});delete o.imgd;if(s.imgd)o.img=bd?s.imgd:true;if(s.type=='chart')o.data=ln.map(function(z){var m=z.match(/^(.*?)[:=\-–]\s*(-?\d+(\.\d+)?)\s*%?$/);return m?[m[1].trim(),+m[2]]:null}).filter(Boolean);else o.bullets=ln;return o})}
function pedit(){syncimg();$('pst').innerHTML='<div style="display:flex;justify-content:space-between;align-items:center;margin-top:18px"><h3 style="margin:0">'+esc(P.d.name)+'</h3><button class="chip" id="pback">← Change design</button></div><canvas id="pc" width="960" height="540" class="pv"></canvas><div style="display:flex;gap:8px;margin-top:10px;align-items:center"><button class="chip" id="pp">◀</button><span id="pi" style="flex:1;text-align:center;font-size:13px;color:var(--mu)"></span><button class="chip" id="pnx">▶</button><button class="chip on" id="ppl">▶ Replay</button></div><div class="th" id="pth"></div><details style="margin-top:12px" open><summary style="cursor:pointer;font-weight:700">✏️ Edit content &amp; add pictures</summary><div id="pe"></div></details><button class="go" id="pdl">⬇ Download PPTX</button><div class="res" id="pr2"></div><p style="font-size:12px;color:var(--mu)">The preview is a close approximation. Morph transitions and animations play in PowerPoint (2019 / 365) slideshow mode; other apps may simplify them.</p>';
$('pe').innerHTML=P.slides.map(function(s,i){var v=s.type=='chart'?s.data.map(function(r){return r[0]+': '+r[1]}).join('\n'):s.bullets.join('\n');return'<div class="ar" data-i="'+i+'"><b>Slide '+(i+1)+' · '+s.type+'</b><input value="'+esc(s.title)+'">'+(s.type=='title'?'<textarea style="display:none"></textarea>':'<textarea rows="3">'+esc(v)+'</textarea>')+(s.type=='chart'?'':'<div style="display:flex;align-items:center;gap:8px;flex-wrap:wrap"><button class="chip pb">📷 Add picture</button><input type="file" accept="image/*" class="pf" style="display:none"><img class="pm" alt="" style="height:44px;border-radius:8px;display:'+(s.imgd?'inline-block':'none')+'" src="'+(s.imgd||'')+'"><button class="chip px" style="display:'+(s.imgd?'inline-block':'none')+'">✕ Remove</button></div>')+'</div>'}).join('')+'<small style="color:var(--mu)">The outline slide and the website promo slide are added automatically.</small>';
[].forEach.call($('pe').querySelectorAll('.ar'),function(r){var i=+r.dataset.i,pf=r.querySelector('.pf');if(!pf)return;var pm=r.querySelector('.pm'),px=r.querySelector('.px');r.querySelector('.pb').onclick=function(){pf.click()};
pf.onchange=function(){var f=this.files[0];if(!f)return;fimg(f,function(u){if(!u)return;P.slides[i].imgd=u;pm.src=u;pm.style.display='inline-block';px.style.display='inline-block';syncimg();ref()})};
px.onclick=function(){delete P.slides[i].imgd;pm.style.display='none';px.style.display='none';syncimg();ref()}});
function lab(){var L=P.L[P.i];$('pi').textContent='Slide '+(P.i+1)+' of '+P.L.length+(L.kind=='promo'?' (website promo)':L.kind=='agenda'?' (outline)':'');[].forEach.call($('pth').children,function(c,j){c.className=j==P.i?'on':''})}
function main(){rd($('pc'),P.L[P.i],null);lab()}
function thumbs(){$('pth').innerHTML='';P.L.forEach(function(L,j){var c=document.createElement('canvas');c.width=220;c.height=124;rd(c,L,null);c.onclick=function(){go(j)};$('pth').appendChild(c)});main()}
P.fr=function(){if(P.L)thumbs()};
function tween(A,B,done){var tok=++P.tok,t0=performance.now();(function f(n){if(P.tok!=tok)return;var u=Math.min(1,(n-t0)/800);u=u<.5?2*u*u:1-Math.pow(-2*u+2,2)/2;rd($('pc'),mph(A,B,u),null);if(u<1)requestAnimationFrame(f);else done()})(t0)}
function play(){var L=P.L[P.i],tok=++P.tok,t0=performance.now(),mx=Math.max.apply(null,L.shapes.map(function(d){return d.a?d.dl+.7:0}))+2.5;(function f(n){if(P.tok!=tok)return;var t=(n-t0)/1000;rd($('pc'),L,t);if(t<mx)requestAnimationFrame(f);else rd($('pc'),L,null)})(t0)}
function go(j){if(!P.L||j<0||j>=P.L.length||j==P.i)return;var A=P.L[P.i],B=P.L[j];P.i=j;lab();tween(A,B,play)}
function ref(first){api('/api/ppt/preview',{method:'POST',headers:J,body:JSON.stringify({design:P.d.id,topic:P.topic,slides:pcol(false)})}).then(function(j){P.L=j.layouts;if(P.i>=P.L.length)P.i=P.L.length-1;thumbs();if(first)play()}).catch(function(e){show($('pr2'),'Error: '+e)})}
var tm;$('pe').oninput=function(){clearTimeout(tm);tm=setTimeout(ref,500)};
$('pback').onclick=function(){P.tok++;P.fr=function(){[].forEach.call($('pst').querySelectorAll('canvas'),function(c){rd(c,P.designs[+c.dataset.i].thumb[+c.dataset.s],null)})};pds()};
$('pp').onclick=function(){go(P.i-1)};$('pnx').onclick=function(){go(P.i+1)};$('ppl').onclick=function(){if(P.L)play()};
$('pdl').onclick=function(){var g=this;g.disabled=true;g.textContent='Building PPTX...';fetch('/api/ppt/build',{method:'POST',headers:J,body:JSON.stringify({topic:P.topic,design:P.d.id,slides:pcol(true)})}).then(function(r){if(!r.ok)return r.json().then(function(j){throw j.error||'Error'});return r.blob()}).then(function(bl){dl(bl,(P.topic.replace(/[^\w\- ]+/g,'').trim().replace(/\s+/g,'_')||'Presentation')+'_'+P.d.name.replace(/\s+/g,'')+'_animated.pptx')}).catch(function(e){show($('pr2'),'Error: '+e)}).then(function(){g.disabled=false;g.textContent='⬇ Download PPTX'})};
ref(true)}}
function imgc(b){b.innerHTML='<label>Choose image</label><input type="file" id="fi" accept="image/*"><label>Quality: <span id="qv">60</span>%</label><input type="range" id="qr" min="10" max="95" value="60"><button class="go" id="g1">Compress &amp; download</button><div class="res" id="r1"></div>';
$('qr').oninput=function(){$('qv').textContent=this.value};
$('g1').onclick=function(){var f=$('fi').files[0];if(!f)return show($('r1'),'Choose an image first');var im=new Image();
im.onload=function(){var c=document.createElement('canvas');c.width=im.width;c.height=im.height;var x=c.getContext('2d');x.fillStyle='#fff';x.fillRect(0,0,c.width,c.height);x.drawImage(im,0,0);
c.toBlob(function(q){show($('r1'),'Saved '+Math.max(0,Math.round((1-q.size/f.size)*100))+'% ('+kb(f.size)+' to '+kb(q.size)+')');dl(q,'compressed.jpg')},'image/jpeg',$('qr').value/100)};im.src=URL.createObjectURL(f)}}
function pdfl(cb){if(window.PDFLib)return cb();var s=document.createElement('script');s.src='https://cdnjs.cloudflare.com/ajax/libs/pdf-lib/1.17.1/pdf-lib.min.js';s.onload=cb;s.onerror=function(){alert('Could not load PDF engine - check internet')};document.head.appendChild(s)}
function pdfc(b){b.innerHTML='<label>Choose PDF</label><input type="file" id="fp" accept="application/pdf"><button class="go" id="g2">Optimize PDF</button><div class="res" id="r2"></div><p style="font-size:12px;color:var(--mu)">Lossless in-browser optimizer. Scanned or image-heavy PDFs need a server (Ghostscript) to shrink a lot.</p>';
$('g2').onclick=function(){var f=$('fp').files[0];if(!f)return show($('r2'),'Choose a PDF first');pdfl(function(){f.arrayBuffer().then(function(a){return PDFLib.PDFDocument.load(a,{ignoreEncryption:true})}).then(function(d){return d.save({useObjectStreams:true})}).then(function(u){var q=new Blob([u],{type:'application/pdf'});show($('r2'),'Saved '+Math.max(0,Math.round((1-q.size/f.size)*100))+'% ('+kb(f.size)+' to '+kb(q.size)+')');dl(q,'optimized.pdf')}).catch(function(){show($('r2'),'Could not read this PDF')})})}}
var LG={'te-IN':['Telugu','te'],'en-IN':['English','en'],'hi-IN':['Hindi','hi'],'ta-IN':['Tamil','ta'],'kn-IN':['Kannada','kn'],'ml-IN':['Malayalam','ml'],'ur-IN':['Urdu','ur']};
function lo(sel){return Object.keys(LG).map(function(k){return'<option value="'+k+'"'+(k==sel?' selected':'')+'>'+LG[k][0]+'</option>'}).join('')}
function lds(u,t,cb){if(window[t])return cb();var s=document.createElement('script');s.src=u;s.onload=cb;s.onerror=function(){alert('Could not load file - check internet')};document.head.appendChild(s)}
function trn(x,a,c){if(a==c||!x.trim())return Promise.resolve(x);var w=x.split(/\s+/),ch=[],cur='';
w.forEach(function(s){if((cur+' '+s).length>400){ch.push(cur);cur=s}else cur=cur?cur+' '+s:s});if(cur)ch.push(cur);
return Promise.all(ch.map(function(q){return fetch('https://api.mymemory.translated.net/get?q='+encodeURIComponent(q)+'&langpair='+encodeURIComponent(a+'|'+c)).then(function(r){return r.json()}).then(function(j){var t=j.responseData.translatedText;if(/MYMEMORY WARNING/.test(t))throw 0;return t})})).then(function(r){return r.join(' ')})}
function pban(){var u=location.origin,lg=CF.logo?'<img src="'+esc(CF.logo)+'" style="width:84px;height:84px;object-fit:contain;display:block;margin:0 auto 8px">':'<div style="font-size:46px">⚡</div>';return'<div style="background:linear-gradient(135deg,'+CF.c1+','+CF.c2+');color:#fff;padding:30px 40px;text-align:center;border-radius:0 0 30px 30px">'+lg+'<div style="font-size:38px;font-weight:800;letter-spacing:.5px">'+esc(CF.name)+'</div><div style="font-size:15px;opacity:.92;margin-top:2px">'+esc(CF.tag||'')+'</div><a href="'+esc(u)+'" style="display:inline-block;margin-top:14px;background:rgba(255,255,255,.24);color:#fff;text-decoration:none;padding:8px 22px;border-radius:99px;font-size:15px">🔗 '+esc(u.replace(/^https?:\/\//,''))+'</a></div>'}
function cta(){var u=location.origin;return'<div style="margin-top:26px;text-align:center;background:#f4f6fb;border:2px dashed '+CF.c1+';border-radius:18px;padding:16px"><div style="font-size:16px">🚀 Try it free at</div><a href="'+esc(u)+'" style="font-size:22px;font-weight:800;color:'+CF.c1+';text-decoration:none">'+esc(u.replace(/^https?:\/\//,''))+'</a></div>'}
function pout(body,nm){var d=document.createElement('div'),w=document.createElement('div');
d.style.cssText='width:794px;background:#fff;color:#0f172a;font:18px/1.8 "Noto Sans Telugu","Noto Sans Devanagari","Noto Sans Tamil","Noto Sans Kannada","Noto Sans Malayalam","Noto Sans Arabic",sans-serif';d.innerHTML=pban()+body;
w.style.cssText='position:fixed;left:0;top:0;width:0;height:0;overflow:hidden;z-index:-1';w.appendChild(d);document.body.appendChild(w);
var fl=['Telugu','Devanagari','Tamil','Kannada','Malayalam','Arabic'].map(function(n){return document.fonts?document.fonts.load('18px "Noto Sans '+n+'"',d.textContent).catch(function(){}):0});
fl=fl.concat([].slice.call(d.querySelectorAll('img')).map(function(i){return i.complete?0:new Promise(function(r){i.onload=i.onerror=r})}));
lds('https://cdnjs.cloudflare.com/ajax/libs/html2pdf.js/0.10.1/html2pdf.bundle.min.js','html2pdf',function(){Promise.all(fl).then(function(){setTimeout(function(){html2pdf().set({margin:0,filename:nm,enableLinks:true,html2canvas:{scale:2,scrollX:0,scrollY:0,windowWidth:794,useCORS:true},jsPDF:{unit:'mm',format:'a4'},pagebreak:{mode:['css','legacy']}}).from(d).save().then(function(){w.remove()}).catch(function(){w.remove();alert('PDF failed - please try again')})},300)})})}
function topdf(x,nm,tt){var c1=CF.c1,c2=CF.c2,sh=esc(location.origin.replace(/^https?:\/\//,''));
pout('<div style="padding:30px 40px 24px"><div style="display:flex;align-items:center;gap:12px;margin-bottom:16px"><div style="width:7px;height:36px;border-radius:5px;background:linear-gradient(180deg,'+c1+','+c2+')"></div><div style="font-size:27px;font-weight:800;color:'+c1+'">'+esc(tt||'Your Translation')+'</div></div><div style="background:#f4f6fb;border:1px solid #e5e9f2;border-left:6px solid '+c2+';border-radius:18px;padding:22px 26px;white-space:pre-wrap;word-break:break-word">'+esc(x)+'</div><div style="text-align:center;color:#64748b;font-size:13px;margin-top:26px">Created with '+esc(CF.name)+' · '+sh+'</div></div>',nm)}
function cnt(el,v){var t0=performance.now();(function f(n){var e=Math.min(1,(n-t0)/900);el.textContent=v%1?(v*e).toFixed(2):Math.round(v*e);if(e<1)requestAnimationFrame(f)})(t0)}
function mic(b,tr){var SR=window.SpeechRecognition||window.webkitSpeechRecognition;
b.innerHTML='<label>I will speak in</label><select id="lg">'+(tr?'<option value="te-IN">Telugu</option>':lo('te-IN'))+'</select><label>Get output in</label><select id="ol2">'+(tr?'<option value="en-IN">English</option>':lo('en-IN'))+'</select><button class="go" id="mc">🎤 Start</button><label>What you said</label><textarea id="tx" rows="5"></textarea>'+(tr?'<label>English</label><textarea id="en" rows="4"></textarea>':'<button class="go" id="d1">⬇ Download PDF</button><button class="go" id="d2" style="background:#64748b">⬇ Download Word</button><div class="res" id="r5"></div>');
if(!tr){var out=function(cb){var r=$('r5');r.style.display='block';if(!$('tx').value.trim()){r.textContent='Speak something first';return}r.textContent='Preparing...';
trn($('tx').value,LG[$('lg').value][1],LG[$('ol2').value][1]).then(function(t){r.textContent='Downloading...';cb(t)}).catch(function(){r.textContent='Translation failed - check internet or try later'})};
$('d1').onclick=function(){out(function(t){topdf(t,'speech.pdf',$('lg').value==$('ol2').value?'Your Speech Text':'Your Translation')})};
$('d2').onclick=function(){out(function(t){dl(new Blob(['<html><meta charset="utf-8"><body><p>'+esc(t).replace(/\n/g,'<br>')+'</p></body></html>'],{type:'application/msword'}),'speech.doc')})}}
if(!SR){$('tx').value='Voice input is not supported in this browser. Please use Chrome.';return}
var on=0;$('mc').onclick=function(){if(on){rec.stop();return}
window.rec=new SR();rec.lang=$('lg').value;rec.continuous=!tr;rec.interimResults=false;
rec.onstart=function(){on=1;$('mc').textContent='⏹ Stop (listening...)';$('mc').classList.add('lv')};
rec.onend=function(){on=0;$('mc').textContent='🎤 Start';$('mc').classList.remove('lv')};
rec.onresult=function(e){var s='';for(var i=e.resultIndex;i<e.results.length;i++)s+=e.results[i][0].transcript+' ';$('tx').value+=s;
if(tr)trn($('tx').value,'te','en').then(function(t){$('en').value=t}).catch(function(){$('en').value='Translation needs internet'})};
rec.start()}}
function qa(b){b.innerHTML='<p style="font-size:13px;color:var(--mu)">Upload a question paper (photo or PDF). Answers and explanations come in the same language as the paper.</p><input type="file" id="qf" accept="image/*,application/pdf"><button class="go" id="g6">Get answers</button><div class="res" id="r6"></div><div id="an" style="white-space:pre-wrap;margin-top:12px;line-height:1.7"></div><button class="go" id="d6" style="display:none;background:#64748b">⬇ Download answers as PDF</button>';
$('g6').onclick=function(){var f=$('qf').files[0];if(!f)return show($('r6'),'Upload a file first');var fd=new FormData(),g=this;fd.append('file',f);g.disabled=true;g.textContent='Solving... up to a minute';
fetch('/api/solve',{method:'POST',body:fd}).then(function(r){return r.json()}).then(function(j){if(j.error)throw j.error;$('r6').style.display='none';$('an').textContent=j.answer;$('d6').style.display='block';$('d6').onclick=function(){topdf(j.answer,'answers.pdf','Your Answers')}}).catch(function(e){show($('r6'),'Error: '+(typeof e=='string'?e:'cannot reach the server. Open the website from your Flask link (https://...), not the saved HTML file'))}).then(function(){g.disabled=false;g.textContent='Get answers'})}}
function ff(b){var P=['name','father','dob','mobile','email','address'],pf={};try{pf=JSON.parse(localStorage.getItem('pf'))||{}}catch(e){}
b.innerHTML='<p style="font-size:13px;color:var(--mu)">Save your details once, then upload a fillable PDF form.</p>'+P.map(function(k){return'<input id="p_'+k+'" placeholder="'+k+'" value="'+esc(pf[k]||'')+'" style="margin-bottom:6px">'}).join('')+'<label>Upload PDF form</label><input type="file" id="ffi" accept="application/pdf"><button class="go" id="g3">Auto-fill &amp; download</button><div class="res" id="r3"></div>';
var M={name:/name|applicant|candidate|student/i,father:/father|parent|guardian/i,dob:/dob|birth/i,mobile:/mobile|phone|contact/i,email:/mail/i,address:/address|village|town/i};
$('g3').onclick=function(){var f=$('ffi').files[0],v={};P.forEach(function(k){v[k]=$('p_'+k).value});try{localStorage.setItem('pf',JSON.stringify(v))}catch(e){}
if(!f)return show($('r3'),'Upload a PDF form first');
pdfl(function(){f.arrayBuffer().then(function(a){return PDFLib.PDFDocument.load(a)}).then(function(d){var n=0;d.getForm().getFields().forEach(function(fl){if(!(fl instanceof PDFLib.PDFTextField))return;var nm=fl.getName();for(var k in M){if(M[k].test(nm)&&v[k]&&!(k=='name'&&/father|parent/i.test(nm))){fl.setText(v[k]);n++;break}}});
return d.save().then(function(u){dl(new Blob([u],{type:'application/pdf'}),'filled.pdf');show($('r3'),n?'Filled '+n+' fields successfully':'No matching fields found - form may be scanned (needs OCR server)')})}).catch(function(){show($('r3'),'Could not read this PDF')})})}}
fetch('/api/visit',{method:'POST'}).then(function(r){return r.json()}).then(function(j){$('vcw').style.display='block';$('vt').textContent=j.total.toLocaleString('en-IN');var t0=performance.now();(function f(n){var e=Math.min(1,(n-t0)/1200);$('vc').textContent=Math.round(j.unique*e).toLocaleString('en-IN');if(e<1)requestAnimationFrame(f)})(t0)}).catch(function(){});
api('/api/config').then(function(j){CF=Object.assign(CF,j);applyCfg();chips();draw()}).catch(function(){});
chips();draw();applyCfg();
</script>
</body>
</html>
'''

LOGO_B64 = "iVBORw0KGgoAAAANSUhEUgAAAQAAAAEACAYAAABccqhmAADUlUlEQVR42uz9d7ylV3UfjH/X3s8pt06VBELC9CKJPqIZsIZiEC5gbA3YcSFObFySvCGM4pT39Uh5k1/yi4jjFCfBiWMndgzMUITpCJhRQUKoazSSUBeSRmXanbn9PM9e6/1jt7Wfc+7MqCEJ7vCRLU25c+45z157re/6FmD1x4/MDxGh7dvFbtu2s1rh108QkZNF5Nn79u07+d577z15TuTZ4edOGvVn/NcTs/ru/mj+oNW34Efi5NN2wGwhcuqwbzhU47Rb7z582v0PLvzUgSNuXdPwpqUB95aWgcYJXO3Q61pM9CtMj5mGub7mhA0TR07a2L3oTaetuRzAHUS0AADbtok57zwQACYiWX3TVwvA6o+nx61viIjDv1f3z+Et19y0/30zc+5vPXKYT1hw05hfBhoBFpaXAFj467wCEcFA4AYDCDMm+j1YAFN9YIwO4dkbxu5+1gn9L7/uRVN/vHYt3am7gi1bcrFZ/bFaAFZ/PAXt/q5dsJs3UyMi9sb7Bh/cc8/Mx+55uH7tocE4jsw3IGvFdrqOiMgYgoEYEUCcAEQgFogwIIARAkDMLGARodpVFVmsnarQGeyfPXVjb9drz1j3hdNe0P1rIlr0hWC73bJly2ohWC0Aqz+eqlv/zhk583u79/2Xe/ZXm+470GCAinv9PlvjrMARxEAgMDAgIYgIIAYQAYlAHEMAEBFIAIGABOjAgMhw0zghJ7ZLBuvHavzEs3DLK188/S9f9fLuXw8GvNoNrBaA1R8/zB87d0oVbv0N39yz9I+u+f6Bcx+cH+8sSdX0epWBOMPM4dMVGDb+U/Y3PCACI8YXAiAVAEPkHwb2hQECkADGFweBGK6XnFDjqo3j8zjjRZ3r3vW2kz463qWLgW1mm5yH80NRWv2xWgBWfzwJP+Jtu3dO3nPJDYf/45699sX7ZpelO9YVEBlhgUj54RohkAhc+BkjgAkdgIOAgfBzqWaARMDMECFUsTMQf7aFwctLA9hm2bz42YyfetOzPv6al/TOdVx2Jqs/VgvA6o8ncN7fsmOH2bFli/vO3Ufef8X1Bz9914GJLvd7da9HHWGGgEDsb/Z4uxPywWbxP0MssLCACDgUAFqpAIBQSWggBP7rMoPIgJl49tAMXnBS12w6o/+pd71p7ceIaK+IWKLVkeCZ8qNafQue/of/rPN22YvP39JccdvsBd+6+sjW2/dbmVrXZ2vQadhP+H4x508rxf/FQ02h0se2PxUIgmEBBQxAmOGnB0ndgi8U5L80ASIGYMDCmfXrpvHAgaV67qqFDy0tN68VkbcT0QPbReyW1SKwWgBWfzz+w3/erl324vM3N9+88eAFF123sPXuA8at2TBtWJwRR4AxgBAkHP54sP3MHw5wnAvCIUYcE/IfATMDIhD1y/7A+4JB8fcLYHw9QMOMqal+Z2lZmsuuqV9SLz28MxSB+1eLwGoBWP3xeA//ebvs+edvbi7aPfPxXbsXPnbvYVOv2zhVOdcQxWtdGAGqU0fX/6uE/1EA/yIYSKEKiAYMJHcCLvy7AGBC6iQEABmPMxAZGAIcMzo9Uzk31ly5Z/7Fdf3gt0XkHUR03+o48PT/sUrxfJof/m/eNHPBJTctfuy+WdOs2zDZ4XD4GfDzOwBKV7v/b+ZwuInAZCBkIEIQJohQ+JMMEYaw/4fQGhModweUsAH/80yAGAJR+DsgMBVXrtNvrruNXvypv9n7LRE5lYjc9u1iVz/R1Q5g9cdjOPy7bjtywa4bl7bunbf1mo3jHdcwjDFhQid/gNOMb0DMYc/vRwJhAeCBQUO+G0DkAiA3BQglIR128r/HCDx3ACV24LsB8V0DAzAEAWArqpyMN9ffsfhi9/m934qdwOo48PT9sboFeJod/h2B03/JbbMXXHzTYOvdM9JMrB2raueKmziRdhAOvxhAOGwB8s4fCIc5HOKI+EvY9YMpHufigSAGKgGE/a+RwgKIPMgIZghL/m/4DUNdO1e5gT3tec1tv/ZLJ7+DiO5fHQdWR4DVH8c4/Ft2+MP/zZtnLvj2TUtb7zqMZmL9RNWw87dyOK8SELkIzIEBApdVXRTSpw54+qkRNwG18ICRvz+ihPHrpyco8gQIZGGd7Ta776pe8ufb935LRJ5DRG7nTlntOFcLwOqPldr+HVvIfWvPzMcv2bO09Z7DaCbX9qumqX3DL5Rufog6kxnsT7f9sVo9UQUhn2fJHAI1GoA84Bd/RoT8WVerRA8o+q/B8F0IkVSgvrv1zuol//NTD+wUkVM2b6Zmp6wWgdURYPXHyJn/oj0HP37pzfXH7puz9eSa8U7jXPEJGSfppk4HnQIHgAES49d5Xt0DAoGdr/JG8kEHABYGcWzdAWJNHiLYiPyzHyn8MWeII1gAURHshEAQUBAVifjRw4aGxDk0tnHVC56zcPuWn/uJX1y/nnavYgKrHcDqDwwDfpfeUn/s/vluPb1+slO3Dj8AsBnu3v2tzZCjKfRbv+Z3/iuXf4F4aGCo78/XhoTDn3+d/FASAcbwK9ZIJabjbru3/+Ivfm3vtx48WP/MltVxYLUArB5+ffhnL9i1e3HrPYdRT6ztd1zTwBqCJYKh8pQyAS4gAPGfdGxjW04UQMFwvxMgRHnSJwMik27/1OqHMYIDosBSagvyaJDhR40ukOjRQr2+im3V77ur9tQnfP2ih7+4/5D84ubN1Hzi6qs7q0/C6gjw43n4d+2y52/e3Hz71iMX7LxpeesPZtBMbZionGtQUV6dMzKTL7bp/kLOH50RBtgfej9/k77uA3pvwurQfz3SXw/h8DsHgQGIwgpQYDg/KF4L4E1EQEFfIOwLTsAF4pf0IiRKXYr/uYrnDh7E619h8dPveNaWk0/qfPbqq6/ubNq0qV59KlY7gB+vm3+zb/t37p7f+oMZbqbWj1dNUxe3shMBi0L3icrDHSu4KgZxh69FQe3/hvrvDPzFmUDSr4cvWPw3UxAeScAEKBYSKfkCMmq0aMz0unW46oaavvK1h7bft7f+wKZNm+qrVzuB1QLw43L4zwpt/9dvPvjxb920uPXeI9RMrp+oGucC2p6PNlHZowmMF+ME8Y8RwLKBEQsjxoN9UKTgdGPbcK5zAdC3fyoSKIk+xGURiX+Ww7hBQqAAFjhmuMBDAHx34Nj/k4uCQFCbifVrcc3NjC9+7aEd993ni8AnPrFaBFYLwI/B4b/4/M3NRTcf/Piltyx/7J4ZaaY2TFp2nG52D65xuP39FF18ncADCCO+P4xhHo/04HiXt7uFeBBNa/ITJQIStLqF1u/THYXvGkIxoKN+7wATRAAnAINpcu007b5Z6AtffuAzt902/4sf+chqEXgqfqwisT/Etv/i8zc3l9x25IJv7Z7/2L1HUE+tH+80roYxlfq9efqPNToeOJPMOfyhF+Pn9EQCCtctGV/XxUlr/5/NgYRF8f/jL8ZZHsUoIoESrL+SSCIkJw4Cpfk//336781YAgARmlq/lm++4zAafmTH7lvmz3nFyyc+u2oxtloAfjRn/vM3N5fcNnPBt/YsbL37iK2n1o11GteUt7QMo7IlCh8PHdLqLwFt6g9H6m7766V9AVEyBY0/z/E4h4pgpX14SQ0i7RGfyi2hRO6h/8rt7yl+PYaYyfVr+fa7Dwu5fdtvu6X+4EteTp+5+mrpvO51aFbtx1dHgB+Zw3/ZHUcu+Naeha13zVA9vW6849qHXx0N0a25Wu0VQv10k6rWPI4HCoQb1c4HFKEsDlSuE4VlRAVS+MDQINH+wQA1AHEoUnFu8ThGfI3MtZlYM06331PJRbsO7rjuhrnf2rSJ6vPO22WJVpdUT/aP1Xf4yTz8YdW3687ZC3btXtp65yFOqz4El97igwiXZgTYCvWdCCQAbibe8GxSySApnX7ABKPNQCIuEBV8ce6PIp8wXpDj4iaPzkLxtfqVI6c+wHsFEkQVnEqoKBzCNnw9SutE42LnIKGrsLJwZFmevW6+ee97nv3RV79u7L/84R+KOe88yGonsDoCPOMO/44dMOdv2dxccueRCy66aXbrPYerZnrDdFVzk+m7aSheoTrTytx+UaYdRv3+5AgsOKouQEZ+PTrq7SAKcoxNvhNOHAMOv8YYQVlGOSIUqIAAThoamx7D3n2uc+llC3+y+4a50894Ff3+6advtyKymka0WgCeeW3/5XfMffyimxY+dtchatZumK7qpgmqvjxga+MNwJt50AqnlqDEN+EnaEipR+lrjjoxKx6juPY/WkXKaGEqEJXxt7tmLifPEaFUngrAsV2GyDMHnWtocv2E3Hz7XL201Pze9644TK9/05rfO2fHahFYLQDPpLb//M3NJXcc+fhFNy997LYZqtdsWNdpmiDpDS0zB6KskdYsHltyM3z7EwGG/QGXXAYKkJDi4rB1jgu0H/qwK2Kv8v7LxaItGhhFPTYgyrRhPzqY4s9Ft+H8LQ1jEwSAG0frTpzq3HX/fF1/e+F3L/7Wfpz1jo2/t2XLahFYBQGfzmAKAZHhd/l98xd8+5alj912iOs16yc6zjXKv6+8KQkYYuW1W/8MylH7DCfnnwIsJFrhEve/L48J5SEsupOjl7oS7mNpAYnD44Wowz/0vYmah4hQDwZYs67fuf9B1JfsGvzu1770yH/5zGe3uC1bdhgRWcWtVjuAp9/N/6fXoPrIJqqvuG/+gq9fP7/1joNUT6+fCJJeU9hvxcfdpMuVsvNu6zB64Mx4dR7nymEkFxSJBgFiEiCofUAj4GfCf/toQGrtHDRJSIYLR+oIwmuJegDm1G2YlcaXogANoxFEwUgEMb/AwDnG5LqxziOHBvUV33W/+6UvPCQ/975n/f6WLbLaCax2AE+r40879uzpfGQT1dfsnf//7bxpaev3D0g9tX6q45hBxoSWXYaGemc8YDYKcEsYgWrCg+w/zdPpwl2B8TfiJ3NxYS64A8fTAfhOJJiJynGcv6MAkRLQfxoxjsTvuHGMyele5+AhU1/9Xfd7X/7CQ3/ymc+S27Jlh9m2TVaf3Seic119Cx7fj3PO2W537Njirrxv4Q++/L2D/+buuU4zuXHaNo0jE253F+drFgWl+X+3jGzbjXAFRhkwS5LUxgkiJfmE21bUnyN1JoQ5b+qjGo/DujCCjZTVekYAK6GpCNRe0oCdAGAHOCSMQZiCIWnMElQmpMG7LIaNJh/DhBOUry2uLpk5rR3jeFCRwcKRul4/5TpnvqH6Lz//gRN/3znfea12AqsdwFPW9mPbNrNjxxZ35b2Lv/uFK/b9m7sXxnhi/ZR1jSMblDx+pz9Mpc1gWW6nE8pfbMiCuCYdDMTlfXmzax5AaOVpxO0f/1XP68TaKDRTe+NBzWCg6j7SPE/BjSjOHWbYNoxjN1B2G4XFWZhX4vsV8wotfFEYn+50Ds12miuvqH/vm189+DkR6ROdRyKrncBqB/AU/Ni5U6q3b6bmqr1LWy/+fn3BtffOuTUb1xrHfkKm4JefbnHmfKtSAM6YUZFpfSCUDPiFBewABDWfVQ26YS29lbR28wcr/HfIC5AA+hkikIvkIMoHmfMWQEIXQAQYxTVIRqNN3gpI6AJMEi0RINZ3H5yBTqgOIH2X7NmIRXJR6ACsEIzyEpAgPa6oi7nDS/XGtdTZ9Hr51M+876RfjsV4tRNY7QB+aD8+cbV0Nm+m5rq9S3/4zd3zF1xz7/xg+oS1pgGTGIIYKhJ1vGmGbnsD796YZL2VrbUjyGdhyMLAwIpHa1NjcMxHnQrjT4SvTex1AKKucu0tEl+BMXHf4Gm7JAYm/H+NN0jSGxhALCAr0IOF0u81YhDzgzlYj0WfAwMkDYIjwCEXFoIFuwaTU93O/kPUXH7p4ENf/Mz+/yoiE0SArGICj+nH6hbgUf7Yvv2m7pZNNLjmwbl//JXrDm279WDVTJ8w3WmcH4iH8bg2jKd+fgTgLtr0BwQh770P8jemUYYfxdqNSgVf/NJlmMcIUI50Sy9DmwoU7X52DdbbCxGNJSimYCQ6xVFm1AZPeYcYhY+wwicNcpVibjA53akWZqW+7hr+nbrZOyFy8m/6RLPVTmC1A3iSb/4tW84YXHr33Nl/c9Whf/H9GYvxdVO24cjmKQ0xfXctKw5c6i4NiDwldF/U0JyJPyWGgDZItyJggdG/HrEJMb7YhL9IWiEC8X/cdvoRGips7TNO7B2DSLjANMoKMPSXZhqRlN+ikIFzDcYnqs7Bw1xfe5X7tS/t2PspEVTnnbfLrm4HVjGAJ+XHOdvF7thC7tK7D5/9rZtmL7xrfqzTnxoXYS4euISsc7gBHYOovK0LMC807BTtfyR7+3mST8L4QCKw8DO+pI2C8TO8QvVJgYKi9/ORYBTTfIRA4eVnsNIlToKRzE+gAET6c+wZjC6cNTOC4GM5z/4SbUrIYxnEqluJqw0AVgzY5UUni8A6rzxMRZAF5AJQaAjL803Tt4vV297Zv+S9H3jWT7Hz4wCdT7z61K52AE9Ihdy5U6odW8hdfveRX/jadQcuvHturNufnhQRMZF4B2i3HE59szGjGXI09DGY4tAK2ret3pC3RDthhvbUe/+CJBw7Fkn6/+EihePa50f+AYwBmaOLlPSfoRbpRzMWkeERtPuNaISaZM4x5VgymcmPA4L+RLdabibcJd+s33bhXz30FyLSp/OJV0NJVzGAJ+THt3dKtXkzNdc9uHj+13fP/+Hds32ZOmFcHDuTOu8QykHGz7nM+YTHWTaev2iWkbgBLZ1/DOAgR8mqK5cJSuBZvnE5uwOFQsGKOhyThOOu378Gk2LD/WH2B9yCwi3titFDI4UkNg06RlYoZ+EPcbrko3goSKBZ1JiDQtGgtwUiOYUsdjXEvkCAPGLg2KE7Ye3S3Ji7/jr5jUHzwEBE/j4RLfsCvdoJrBaAx/hj286d1ebN1Fy7f+mffvl7B//vmw9UzdSJGwxzY4igfPml8OArnHlXmLPiITaR5iN+LQiYQIknNernGTyKfYSkSPzRbXyMBk/jQED+JR5MBb5Z+L2fCxXEhF2AgJONuL/NfSKQqBUgHbfcuAgaB8Qf3NgBmLD6THWz/XUFLbYjFUIo5xz6E107c7iur7sKv9Xr7ztVRD5MRA/LdrG0ajG2WgAey+E/f/Pm5tqHB//0y1fu+1e3HLI8sXHKOm7IUG5xNZ2WWcdvxyZcPbwCRbIh1SFQ4d1fZHWn/jsSeIIZqAvRX9IaP9pzXbjGqcDZ0v0PVhHjHP6SuNcXZfQBCoYkQomNqF9jW5U8/O+UxqNKSDkG5C6DRWAzKjIS8Fyp5DiuMT5edRbnxprLdy68p7IPXSQi7ySiR1Z9BlcLwKM7/Ntv6p6/+YzBtQ8v//qXrjnwr/bsNzJ94hrjhInIH7p426aDTblVTXFdFFpxMtlHz7VuQ2SOnajbLUL3RqKTjifJSLylFVCYgLcIkgklW3AOo7AVBxfAPwuTiDZM1ucEIs7XATOADQCmSX04Q0IrHrsA7RYoqdjkYGIJ70kAGplBInDkR4lYGz1t2fgMBIkdiMcMTLRHZ4GTCBZm1qQXNzmACA03GJvoVIuzY/XOr8y+QgaPfHO1CKwWgEf14+qrpbNpEw2uemjhl7963f4/37OfeOqEdcaJkLTsMYHki6uCNJDJ8oEN5FNzw+82pEAxBN49JTcf1kvE4P5romzYUDqEpnXJkwBN5PDHG5o8j8DfkAgNvAELSrV+EPfo5oSZYGHBGX7Ph5sk6fsBhGyAyEDMjsEGlEJDNPU3bThUEjGrQsZgf9jDP5Zy2kFpWCIBRAyFBgaDukF/otNZnptuLv7G7Cvqem8qAiJiaTWUdHULsNKPT1x9dWfTJqqv3Lv4ni9ffeDP9hwwNHnCOuPApIVzfnVOYCPpn8jx5XBQxRDEBkSeTPiHUlsdkX4h3Q1EBr/fyUtYC7IIXDygUcxDnkmYviapm5P8PwwqwbSEWxi4IFJiQfo6TAYsBuwobAhMjhQLJSNShz1N2VOVg00gHMfNhwEZ04oVD31N3FaoIcHz/EiNFSG7kLzFucTuIx56yXIqkfgY+3EGJKhdjd6EqRpMNJftrF/x2b/Ye5GInEBEbnU7sNoBjPyxU6TaTFRfsXfu7K9ft/8Ld872OhMbpqVhJgENIXlthD664bFqAPzpynbb2sdPC4BEE2riLDEEfJXgGAdjDxtXgBjGzXKPEn/CBPTcx4NnwVAMDaUEOVgWMDg4+4h+1Up0FH9/Kzk4pQcFVaMwOLTzCFsAFk7uSFa9vWmDIfF2Yj9qRKNUUu9TAiljd+CzE1gINTN641VVL0w3V148/8qm2fut1XFgtQCsePNvJqovu/vwe79+7b4v3D43Yfvrp9ixmIToR+hMRu/BmbTZBUVUK8300X3XFEUA+QBJYs0k3B+S9f+x3BBl0DADaMrQU6vrwrzu7bpCN0GlctCGA8UsMGlVGfcOJi3qrfIcYR9PDEP+wJpwexvJLsMiwSa0EC1SIjbFgBDd/gt7vwGO9GIiGJVknvgRbFUUOfvlSRiLWNGuG3bo9KuqWZpqrvvO4iua5R98S0TesVoEVguAOvzS+cgmqq94ZOnsb1y9/8Lbj4zb/oYp8TLTfD9qxx5JhpdY4d6NST7+NLJQyPQjxQDMrTAr/ny8hXW6jj7sFHbhsfUVRgEkggOSLtGqK2jsARDZiP0nwo1JWITJu34IJLz2aA0etxucDjfUnl+SApEjuiB+kWhYtwaZ3e+/Z+tXguJ5ESa+r3GcYWSdQZQFuRR8FkBDpA2KVzJ65iXCe8iuRlWZqlnuN9df6c5oXC4COwPHYxUD+DH9sX272I9sovqGR+qzv3Xd/gtvOdytxjauE5FI7+Uhvz4fcydwxGDy/87h97VTfpwQHBO0jR2TXwQIZbuf7AtYxnCJSgcWKG+9tPYLZ96ErxtZdOH/+q9sIbBFsvBwajCns2kAFTwiSZMQR5io2/cHlMJqsLQOZWS8Qv99KTxUkFiBRSdFK1mK5VUmJQyltDEXMFianHwMgQujh3OMqksV0URz01XVGX/xx/d8U0RO3LyZGvkxxwR+bAuAF/aQu2Wm+e0vXv/QhbsPoDO2flIaGRhBXGeR0qlk7jo7Thx3B8CFFr1hB+fY/8Mcbk0V8RtnY4Lag5OiAEseJ5B/jxOCS68ngor+0DeMUFHius2A2fgwDjGACCw8sMeO01jg64rJFSm4BSH4FKTfF629owJRDIjJOwqlVajXFXC4+Sm4GkUvxPg+xm6E2SUjEM+cDIAiDFzCM7LjkeQa4Fer0VuQCBQk06lTifOV84RGDqCo4wZVlysy/ebma3uv+PM/uvciWZLTaAu5nTvlx7YT/rH8xq8W6Wwiqm+eaT70xRsPfuK6R0SmT1gHIWcik4adQs6lBavp6y7dOBSANYV8B4abDAFzuSPWazHNC5ARNFsRv2HQNy0FpB4kqQVOd2N7XahsvikNA+GVcnvuQOIwBBjUowKxKChlYDzMHLACYvHdSHD0ScSjFnDKag0I5GRjI8H+HFHDoAxMU+HghD+YQL+W5Iaqccs8xjH7ToCX+833r6NX/tl/uOeLsijvoDG658eVMfhj1wF84uqrO5uI6j0H6rM/c+UDf/7d+5d58oQNwmCShCz7h5FJwk3rH2gXOPyOFE8+3lSRBQhFB6a89uLUnmeOfLz5WSliRIdliP8749fhQNvNppyUHcDTNSnhcDhwUPYBis6vjEdEdSQUEoMb58DM/u+N40EoNi5QeFl4RAtPiRmZp/Sy+njCUexqYs6Q+hoSu4oYdx5xCFVkk29xAAwd5w2FAiEj96CUGvtuo+pwZcyYu/EqesEn/+yh6+s5eQ9tIXf11fJjF0/+Y9UBxFXfngP12Z+75sEv3LCfOpMb1oigIX9XWWV/5TTJPgddxHnW5V10moeheXw2u9twXGVlq+/okAPj9/Fx7RWlKy5y/jnBanC+E4dlL+PNDGM1V4fYMYPMLYxgGQkV830EyiACZvJ7+/gtOi78/8R5ACNmGVD4fmzRqXBaX9qo/gO8TZg4z15kCkBiGDfAiYds0uulIH0OAClT1kAE+bD/8370aEAJ4JTgP6Djyr0JKlJBaERA1tnx8TV8+TcPrjH00N/MHZCfm9xAX7/6E1d3Nn1kU71aAH5EV31XPTL33guvefjCa/ejmtq4gYmcia2xDuXg8ExTsZoPTLuEvCvjT+TYjvigRupqE25sE3bw0UCTxa+wWKRopbWYiEbAYlTMFJogpIqSdvHRK7cw32vAUlgC5Vcyrz8y9xAPaqgdotKHRTsYRQyBYAypiLNwX3PgFFAxV+XX0FYQRdBRHXxqZZ7FIuSLbGADMmBDofOfZ2QJZsdhBPoxgc3k5Hp3xTdnq6XFB/73ofvqD647tbPrx2k78GMxAmzfflP3I5s21dceqM/+8vUHL7x2P1dTJ2wUJjFOCAzrF2KUV12xpRa1AssgXmbwCQwEFjCB7YfwTxwj4g0UmHf6QCZyUGC+QSH1CA8v1PhA4cLk2JoLg53AcWAKSr5U4zji/fooEZLSJiAhBZEEhKQ49MYfDAdGE/byGCI++SLD6ufjWBJFQwh6BceNcjgSgAJsSvl2FhhweO8QcMgmMA6duDwmxS6GSogkOS/HfES/28RIwWLUQYAghm1/bFKuv6w+8RsXPvKNuQNy9ubN1HziE1f/WIwDP/IdwPabbupuOeOMwbUH6rO/dM3DF+4+YKvxE9YJC4dgHpvbRDVPp9Y+0HIzgKXAOoXwJ8KOoIT8AlJuI0IeJH2u+GqUVoTQt7nS7SZfgbSPD06/UWijwMME1BUXaiQxqYtWbTgKkaLiJBnJrTRJCUsOyX4FYHIwUgWxEitJNIUyFn4jVaGomISIOiEYfyy9GJlM+H45vR5q6aujFiNlD0LU8KNepzIXiR+yCRkGZMSMjU/z5RfNVHX9wIWH99bvX3Ny56s/DuPAj3QB2CZithANbplrzv3r7+z9l3sO2c7kBn/4OYFWJQ3VXxCUbiwds8WI3vWkfp8y22h528UkHYjW+pmwysteePpu5ZTOYxLhJrfz7GFIIhiyoBjJq9r9VIRarwPJmiszBjm01UWbL8H6KxSSkSahIzwQ4uHyq8LoJeAUl0ADdKOTjFLRQS5AZcGS9L2xhI0ENDciG61q9mThiRgUhxpw9dwGNuPj0/y9nfMdyL4fmyLwI1sAtu3cWZ1P1Nw6J39n+5UP/ts9M130N67hWpwxYmEikwySVk1ou96yp8hGBZ9/gG2S6yeFHGUCjwkWQV4IY1L7nGy71N9BZFrrxdBtkGTCSyToBJpsgtvIIcl9KPjsu3xy4uuJxiIFo5AJZIILb0D0LWzyCzRheaeOSGbcqRWbi0AiUZIt+8MX6LkROEhKSZNkzIiiomKel5QyBM4dWOAdgoi9OpHzpqbgaFA8zL4gC7eKFOvuTeDACedwAIwVMzk5xVfvOtyBe/DCmR/I+9Y+l772o1wEfiQLQKT33nRIfunLN8/9jysfFjd94hrDIgYwSVJb7OHDbZjXWhK1PBmh162l3oOzqHvRYwSGMsCd7iNtBDoqz28Y7suuO6RAyXjTFn04BQos5VBQ3S+nNWHm/kfLMBmK6m630RqMzCYhQgJxDFgTuhJk70JS8KVAJRGLuu2DVkFDqrFDSIeWlYiKVBeBNmEiA4LCIM7FmSUHoGROlkA7maTPWdiMj03x93Ye7gzqe//mkbuX3nfi8/tflZ1S0Y8gMFj96B3+qzsf2UT1jUfkn3zm6ofOu+Yhrtc+a0PlnEs7LBlh0Q8FTmtJf2puA6st/Tn2apgkugneexk3CP6cYXSQoNU3SVlIxV9vkKH6FKWV9gzl7wV53nxae6mDE/f9RQ6hy68x/c1p8+FbfSQTkWgjJuEWzoXD6DOj3j/H8Tun9HqFkLsAIQU4lu87pe9dTeyMpK0MMCZaO46h8UGDNxFnoUDCYsfBs0C/SqTuhRP3IHQ5hszE9Dq+7juHqqp65Atze+V9dDL9SBaBH5ktABGwXcR+ZNOm+vrZ+l9t/94D//rah7i75qT1lRNHyUCPpEDDOWFtlPPptG5FPfSi5vz0vLGk9jqKgVOwZaS6FHx3g5YdRxIIUSDxRGFN+rORKsuSXms6GHGciAlh0ZUnbQIk2HyVoJhR2n6S9m2q2XTRPtwXCsPZt8/EmPEIkookaq9wIEcFoRGLAYf33ER5NJWaKkoIvh8jiBmGjB8xAt04GZ4AbWeWUgUZUo2EqRgR9AcsoQDkHMTYMQgAZyYn1sm1l7rqa5956Atzj8h76UdwO/Aj0QGICO3YAbOFyN28IP/kk9978J9d9bCTdSduRA2hyBv3vnYAMQdUPllkhnGgLYIJvvtqfpQcYl8IgEQDT0ru2l5DcbjCjU69YM44VUD8RXntcYYo8kFTXyuYZqeCkEI+4y6ftPeeShUOXQkEcCQwTkrQESh8DdAy80iJPgphbzSnAdmePI0vIadglK24SV4AeRUrwTw07vFZcSPiSJa4GmmLY+FcTmP27EcpQEUOhdqlsYCSNRrFIkti+v1p/t7OucqYRz6/9/tLv3DyS/tf+cQnru585EcEE6AficMPf/hvnZePfvbG2T/adcdMs+7kEyyDiZRxhJEW6gxfDPQDS+INNzndKKUNGHOWw2oVHCmFmkRvu7QKy6fFGP3VwhzvOIntScWCs/h9uRUTXHckseWieSacpJFAJJiEOoFhgTU2xJK3fHkdBTfhzC8gPScLqbZcrS8RfQGCPx97pN8Eph7It9tVAPuYs5KRJfMY0HC2SQ9uQBL0CDH0NM/plMaEOMqw5AJghOCcy6OaAIZNKmbBZzlZlcemN6UQI29BfAdmfVcQxijrjVd5MJinM15v6nf83AnvP/VV/R+Z7YB5ph/+83btsluI3K2LcsFnrzv0R9+6/VC9/jknVU6GnXxYhU4gbqXVKi7u15koU39je6+46wTteiNF4EVaSJH/Wk5GkVH0XFFWY2rHbaWvz3qxVfLwdQuvx2GRoSqfuofoWcA5xUcKOe5o/90ULioIWsgsYdLjUPQYkFAIiSmNEwh2ZS4W1IYD1Vjhq5K3H8QCuKwJMCL55xEsyNTrghRay+znQCYV69TlMAeGYJQsI8mIJRulmE5nSm78Ljpf+8zDF95x7eGzN31kU331j8A48IwuADt2wJy/eXNz86Jc8KlrDmz99l2zzcZTTuo04oKvXDlkcmyvAzUVHOW8FOgnXmrLMUxDHQQJ+38Ob1nq4FvpOmnFJ76IUCDzcCbcFX7A3krcFK/TA2guxWYJSzL+IP2bQvdaaVjraIk9IXI8/TaVtOMNPLUNeAto02RnjsakFkWEpyB0K6QAQP2n89e04guPiTZfWa2UOxuXw0J9EQ6tehqLFCErRo6p2sqtUsbMyYdRhlzeKKQbU7YbQ/B7EIBITKc7IbdeY7q7Pn/kwluv8EXgmY4JPGMxgE9cfXVnyyaqb5p1F3zm2kNbv33H4ebEnzi5qrkBSG2ySSHI2qY/OWaZBExRaLCjpQaF/ZFXyoXfIZK8/NNWoSDzkIrw4vRMG9LVNjGLwoNLhete/FqGMSJQOGfnkZD/AFklAXFo+eNuXhWlyA0Q4TQOxa9mlLEHgpzXwqZVokTbL44WXJ6MxMokxEhwJ0kXMQcH4RYfIaLx0vL8D5br/rXZMDpw4gZE/oLeCIjzeQWk0otMS8+AIjshFCTRmYYeMDShmEYxUlyrcvgQDcT0euN8yzUL3bo+/IU9l8697/S3Tn51+zlit+x4ZkqJ6Zl5+P2e/7oj9S98cffhz+28a7bZeMqJtmHfxFJQl1ljclSWSHTuCJLTCP7lAmDEoFHSXHKSAit0hp0orQDBgEj8jR2+TmKacd6oGzVUsHg3XRNafiuUSLLplmQHpPVamKc1o1B8MTJilB13IMkwZ5Rfkp8vmFsR445z5Fa8Rcn/eTiBQQVQIBsFcxHDnHb5iEG8oQASGxg2LYCPEm7hHCeCTjQE9d9HOGAxkETyehPCMDGqLHZJYgPo6jsRF0cEvU1Atlm00gpgDUXOJv62xyz8BkZbkIczHfxCjHjFIgt4eWmJXvgKbt79wQ3/9kWbJv7v7eeIPWc7+JkWT/6M6wCipPeqA/V7d1z1yCev2k/NhlOebRtpiMNsaQJIxsxhTRVjtUTdrxRCMmya5wtwMK60wsNJJMVMLSlyK67qCHlBENF1SvNovOriJG+SPoCShRaFA6FpsXEgJtV5RN67zwvgQInNrML4daD35omlWIJy3lAEWtYXDDkDsMmi9PVGbUlyIUHoNhBEUpQ2LJn4k8RUCTU1IM6GJbGwhFCC9Oet5PEihZlw3tunZCWDVASyDlHChiL7PGQptoT66rXO+bP1RaoiC2YDhvODjsSNhn/x4/0Juf26WdvvLPzze69b7v3Ea+jcbeftrETEPZOKQPVMu/k3E9V75uqzP33Vvs9fdQCdqRPWoWFHBXsslf9wm8XVFxVj6ZDCzUGj1qSQfRQUXg3QRf1AlT2+h+C+NugiHFh70f02il7D/hrK0MLf/K6Y4yndfrrF9TNwMhglGQEJREYix+CvrHlIRzEP0E5qRZtpB5+s1FNy2A6I2rOj0DtQJN+ktWULHl0BvNQRasMmKsHINFkSaOaiqMsggJHx4wq6h1i4HLhQSXpgU2koYnyrCE1OTuLGK+ZrdrL1zqvm8cIzJ84FnllF4BlTALbt3Fl9ZBPVNxyuz95x1SMXfncfdyY2bhCXDDzz057nTPWApCVxXllFgU9OyiPv+yMm0Ec5IM/FF4O2+k7ofZQHU/k023iphQpk1A6aRN1Xomixwhghvi0OPBMp7UEJ9MX5I1pm5X2lUsJJO+eoXTDUuyKaEQi0A8r1gSX19xRU4Ii/pA4nu/zQiJJJQxU03uJc/Fr2TShfV8k7VBRpUYBoHG3U9xFjx5okGvKrTv8Bm+LvZCaanJrs3HjFXN043nrblbN4yRumnlFF4GlfAOKq7/zNm5trj9Tv2371I9uv2k/d/oYNLAEn4yi20YkxZBQ7jMBpxgvsNTJhDkXoEvwHzMnjPpNIohyVwn6eKO+4i5uLFLc8ONxEdx4GQ9hLXhPKzQKqKK3IItFIIAmt15uMdJcqJ6Ac8ilhRmW1I8/04+S/F+ZgGxKH8qFWJKYE1kVfvoCwI9ibJ6fi6EEgScHHMP51tJqxhDOEv89GX4PgM1g0bgogZB1eEj0GkI1b4uvSZcKETUUCBDiwDSgTn0LYk9cxiAlsR+dNW5TIC2RCBxbPsk3IBiBwrsHk9GRnz/dmGyLaeuvls3jZm585RaB6uh/+LTtgdmzZ3Nww6y7YccOBrVc+ApnYsF4c2OQb37PGbXiCiCiYZ1Iy70jtaDjssQVMVFBSKzDm/AAVwF/kDuRWVu/kY+sYOewEAhlObalHIf1bTiYj4xBu9Sw0svWFRrrVxR3TdGK4h2jzjyD7TfChsgRfCQMuqcuUV5ncjIaR48tX3gTtUcgYAjdZzZi8/8Xbmdu2pFl1IRFIdeHbN+oWNi2Ls/aIEtWVaesTvq6J5TSsBiOdWYeyIOEqlHgOrTcqfHwNpqYnq5uvmmvAtPXGXTN45Vlrzz39ZrEi8rQGBqun8+Gn80A4n9yVM/UFn7p+Zut37l9ya086wYgw5RsKHq0OaLANNzTCCimTUiQk69pweyB57hXGmeqRQ1hliXjzCDKxCygfgsw9z+yydHOGw2hZWiIaSTdZbnvZu+IoY0zhnDsY9/ap4eW81RBxILLBEYgTXIUwchgFIAplhlxuWgKDMWgDPLhnAwgYC0p+5RxOo2c9mlzHSKUYhe/XiSTGJYDA0suiH9JFI7GOfZdljgI5CLLDiWltAKDAVuFMDmgbikgxZ+hqRkFOzAnU9PbjKvxFew9wg4mJieqWaxcbdth6w0VH8Kp30bnbznp6dwLV0/Xwn7drl63+5dubXY/UF3z62sNbr3pwuV7/7A1VLY5iZReoPbza7TswrBSGdQnRT+g3ShOKdKE6SSIaUXdAbEfTjl7z65PmPdpf685Aq9CM8hKQRBM2uhNnvedHXtZD7eMVVkehSKX5W/SwkG/bKH2ORSnGiWmmXGmnjeRklLwOU0HzQw0C2Slam6VxKnVH2ap8aOfcomVnrD47LEVAUNuklR0MhqnMKd4cQ3JhIhrd80hZODLnAj7dSSiMAqOLUExiZmaM9fvVzVct1ew6W6/56gxed/bac/E0LgLV0/Xwn795c/OdRwZ/+X+unfnV6w/U9cZnre8MpPH3GkGRSCSdj3Q82TO49HQoJZZdAIYZPKOE7LOUkFS7rSUqY8MKoKoAqfR8LSPvMyYoU04Jc6wUvn0plQjI8tgoYCJlKJra6/x6NSBGrA5SLJ5cchBymHlCAQItFikQJcqdhYOIJzgJGb1tQIwnUxODYBg7EfX3J4Vh1P9nTEPShkZrOXKXk75PllIdiNK1qF0ESPnBaa8FE4uktD9jXZhRFC/nGL3xXueW6xZrkf7WK74wgze9b+25OE+qMA7w0+m80dP18F824y74xK69W284yPWpz39Wx4mDng4JEqirLn0bxJJZYOrDTYGegaSihXhWPyRiwNIkGy+EccK4YPtlIsU1BGZGExBWzjmccwEottMRpGNXGIGQQriJjCceuSYx2PzrMAkDkLQpkIDkB4twiY63lEaOrGoLs7Kj9F4YjQFEGTQj8QvYOUjgSBgpiTUcRhIv9yXYkBYU7c8smcL5KKkOIygZuwtHxc0fIRJDxs/nUSGZAMfMkYhOojbmGsJ7EiRaNEtaNUYLsMCZDIAuJ2+DBNiyb+99Hc5dIhHBJnJT5IQgveflMTJwjU+FssZiaWFQv+B029n0zurjb9uy8dzwjJunUxEwT8vDf0Qu+MTlh7de/IPF5uTnnthx3CgnXP//WfS8ni9xpnINFA06uEiuVUIbPdPBtYCwiFCLohiIWgN599w064tR+/fj7fZE7eZduu3zjtEPIqTWXhz39kFAIynGq/QD1HmCUDefFhEV67Lw3saRRgt09EUunGnW3tqckzswM8OFcJFCI8FlUjCbko8f6dTxXWDK2X5lfmLGTsrvrzXby7D+4Hiuw/bXEqER+ZDksx5a2YYSvnfvauzQ7dvOnXsG7pIvHN567VdnPyMia4iIZZs8bc5d9XQ5/Gft2mUv3ry5ueKIXPBfLl/Y+pnrDjWvfsXJFToANxmJDaNnuHnDJ0VGGV/mm52M99hzEk0xZAilFsoCmbRBEIyYISVl3MWVowavKIJ32mxU8tcyQGbSMXttgCjdQIFZUKbaEkPIJZ/9iF57dVxQ8MXo8cBGNPEAS8A0kHtn0lyBuC6TjKxDlGIuAPveVUfdvmnNmAuYaakYSWMQEizCwjomppB5IVCe82N8GMSllh7GKKci33WJZFwkEniMwnMERjE6R+wjVdiohI5BbX2RsyH932G1nXphOoKsKSDyXgiKbizCGO937cN3cfOtT8/+orX0XBF5z44tOw6LCD0dMIHq6XT4Lz/iPv7fvrv4sb++6mB96snrOpNTfbgg7knkTqPseYKSjFD6uynqTGt2lxF/f0tvw20Sjlq3Kcruiqo7UQ56EpxxocVGTnUMyNl86VuQPNOn0A9VAElyHHf8DinDg0ZKhLqkP6rXGDsXjvZl+aEl2IDoS2m62fo+jbIrN5TR9jYRKy7VdWmgdFhy0fNry5wURKLkx5StzRHGjyT0EsVjVK5CmfcQ5/qoAYkAQZWR/7zCyAxRinbwo8lKQ2vHdCFlHoI4wfhYt7p7z+Lg639pzlxcXPz0lh1b3vV02Q5UT/XhP28X7MWbNzdXzcoF//GKhY/99VWzzfjkWOfEjX1YEjSS13p6fQb2CK1KpiuIOACygEajyMj2ncAw7VSGACulw02hHpT20cOuVBnCS/qB+BAZGzj0Js3WMZ6TtH4VBLKRkOLpDmQo6N/9bWkMEiVYBDA2ugCrwmDyTW5YW2kRyATH4nhFh2ASSUUjHAge8dCr3ThauoOItBdMQYNsSkLZjowSFyCLi4gD798EUZP67IzJazimKI4yOcG4VbgLy3JxninCUUPgeSPCzjshKxlCjD+PBilCUlwoVP4lCVhksGKHZvCSGZiemug+cFvd7PxU/c5df7n/grN+bePTYjtQPbWH38/818zLBf/xOwtb/89VR+qqP9GZnmSsme6hkSZV72IIDase4RLLpNYcJ5peG9F2URsExFsukGiQxS0RYZZ0NRv/oBAwRFdFlJZSIoZE6qixNghPBI1jsGsA12DgOPPRkbMFMuCmFW+ReegfNMsEBwNmBysZ3ENkxqk8A3LW3+ox8htUbAzIryDQyMCPGUww1n9nlVRJI5DWppzCunwaMFUAAXXDqJ3zCsMkIKIUOmpVR5MQdtLAnsBzuzjlMfjRyCT7cx8tximRiEVALkaGApU1IOnAoKdMSuI7bPxnLFYVXFHWZZSDCGJrHxWWRuEJOomIcrS7Lx6cKefGJDahCYDh+Hi3euQe19y4E1sv/9R+vPlDT30RoKfy5j9/MzXXLcoFf3zF4tb/deWRutef7nCzgDNePImNawhO/A2gDywp5Na2jDg0oEScgzgKd5xwSHTLn0BDZQseH4C85jeh4Eha+xvkmOw4z1NsNclAxKJuBuBmCROmxgnjhA2TFdb0LCY6gDUCY6vM9EPGNpiRWYdMXh5M8ev72zpyhEychyPpJTjccOA+SHTPkQziJXaCQShSvrwsL1eYOcKYWwIWZoHFeQ5UYwRtAwfU3cENHLhp0KkGmBwXrJ+uMDXdxXiPUmFN2wplxWPiliTq9aFRdkn4ij9gBmQF1gA2kLFABGsAYyx6HYPx8Qq2EkxPd3DpRUdw49UD9Pu9ACCq9p19kpJJzkKBJcla8uTxBKuszI1BKkBxNWnUFsBaryhdWm78sKcquwl6DIJJRW8waOrnv7TbOf0s+fjbf+PEc4FtRradBzr/h78dqJ7Km/+GZbng331naev//u6Rupqc7tRLDusmLNZPdSBcw0byBWlASRKZhymm3w3P5ExJre9BwGhYEXkdCZyjIXJJlqlS0qYXiZsuED/iWM4CNnlmJVSo6wGoPowXrie88pQxvOzEaTx7ncVUx6JrCB0jxZgwDCmU/PYITlJ7sS850CTqCBLwFVRKUa6s30xRrjksANkKjxwG9tzWYGFpCc2RAZaXoRyQ/OtddjVksIRexTj5WRYveUkPL3j+Gjz7xB7WTBp0OwaVVYwiymNIHDdi6x8vvQS6KbKSKNok5egmaNkPJy6Dl/OOj43jpmsPwzkOnywXo58Rk+XOIh7eU4GjLNlSSDQRS2hoS8BqA6B5FIpH6Dc70QmZvWc0saDXqTr33to0g5q2Xrr90ElvOWfth4mIn4oVYfVUHf4bl+WCf3/F4tb//d3DjR1f0xEHCNc4YX0PnUpQNwrxFi1JzSQWRpRmjFqsxZtbEnhnRhBA4j44fnAm2mQlNJgKN13/UwpFTtJaE259wmB5Ds9fU+MdL5vC604Zx9ouwUoN55a9zVUDDEb6BGJ0AJ4uCC3VcRGBhbK4ZQLN8K+LyhGAqbB3X4OLrzqC2+8GZo4AjfMmGUQCMgaNY8AtYO1kjRec2sVrz5jGS184jjVrAIMGzjVwTQNhoHbHbjxJlFuTjKDhDekiRFW+YCxitGcjYE0HdcMjwkMU2Zd1kczwHkdWZugEKZm5mMLzMM70BJXsJGXQalpPJ++RoMwUgyrgQ1VXqgdud80Vn+Nf63bk+SLyc0Q088MuAtUP8/BvAcyOzZub3cvu4398xfzH/vw7hxo7uaEiB7jGoUMNNqybguMG6r0FtdwzaSjgotxv6zY+K94436BHm4cMFXTYuKbSDLckd41zZCDNOGZU9SG840V9/MwrT8KJ/QZusIxmOWQGUIYgh4OBqHXxryDUiS+hFZCpv8xw6Rj96ywCU3Vx//4KX/jWftx2t0Pd9GFN5cN+QqEZLM1j7VSN179qAq87/QSc8qwKve4A7BYxWIxgA+fP5TgGSyr/z9D3SDK6AJKS5SZn4UicEmBpwUHLdpNJaFyNagiTpRXX5g9pxITSeiO8qGgqY9TKMAGahkDiwEmAnSu0S+zO+Dx6rKDTs9UDty823/zLwVsGg+YiEXnXD7sIVD+0w78DZscWcruX5eP//vK5j/3PK/zhR+NFPLLssHZDhfExA3ZNQn9JAdLp4FE2eWBkbzeKe/2gdmEib4FN2Q9A9O8PgJOUrxUCG1pkScYVscUnDnequLydI0Jd15iUQzjn9RvwjhdOwQwW0Cz6XEGP2FPrdpZ0A8nQLuFxgDpyPL/u23lbGRxY7OHzF+3F7u8zrB2DNZRBVgFcfRivelmFd5/1bLzwOR0YWULj5jBY9jegQQRKqcBaRiHxK74mKg1RRxN49K7NqK5cRbk5wsJCA0IVRresoCSmLIACwYlLrEnNucr+pNlSLTI5YzGNng1R8WnCmTZqS5X5VDlIxpE3HPHYkt/ydHud6qE7l5pdfzW7iWrzDRF5NxEdku1iacuT7zP4pBcAIuC8Xbvsji2bm5tqueCPL5//2P+8YqY2U+s77ELv7RjgBiecMA4iFw5fbNGM5/Wbcl73qzCTwRxoqa5JfV92vAk77rDuQmAVciKTtGa9FtKfHL5IMo008BOapsE6zODX33Ii3vTccfD8HJjgD5Ny38mvTR91Oe5D74EzHJVpSCpQVP97SXfzb8Gim8TnvvoQbvz+AN3u2nDr+pvcNQ4dHMJ7zlqHd711A8Y6S3D1LBqE75u03dajO/BPxg9jDJYXBYcP1TCmq/YNKb7ES6Yl4w7UKv6giAPkVt6GrAP9LJi4So05Esyw1m98uGn8KpPg5dNE2ZPRufB5cFJHNiwYG+9U+++rm2/9n5kzB3X99VgEtm/fbrds2eKesQVAROi8PXs6559xxuCmJflH//E7C1v/x3cP12Z6XUec82g0A1ILuh3G2qkK7DiZTFgigJzaNXNmjY2wjYLSjvsmQGAorvigRC1qbksGHHmP7lWxVPr7pifFBNCn8jn24jAps/jVN5+ANz6ng2Z+1oOXlEk9Cega8VpHIxhHe1OP8cs0+t81iwJgUG8Kf/PVg7jiumX0xtclHMQai2ZQY6x7BL/03hPx1k3T4HoOTR3XlqT238f1DKx406/codAxF1Za2iUisNZi9nCDI4cE1vT8cyWmSDPO6T8BmY+mIiF6jIXRkIEJvICUhZjwHxM2AdICYIPJigk6yWT35v+/S3gTBep2TnUWAM4Jev1OdXBv3Vz86dkziezXReSniWjmyS4CT2oB8G3/GYNbavk3//6KxT/40+/ONHZqXcc1dYCfAxOtqTG1xqLXsyrlJT+w1GoInXMlDtCadkWl27bvgTzrQfnQhw9RMkkk/u0JApLhNl3IwCwdxi9sWoM3nNpDszgPa+yI55ee8i2sBv9MfwLfuGIG37r8CLpj69LNZ4gwGAww3T2CX/6FZ+HMV06gWZ4t4sWOVYiOXwPxuL6b1t8JdDoWDz2wjLnDFXqVDXwCqIDT8s+ZIhU4WpcHIJc4B7S27OBYgvmMDiOFJ2tVxsAK4KQZwkEkbCt07JqEi8mnOTn0+rY69KBrLtu+eGaHDsZOYGb7drFbnqRxoHqSrn76xDWofmcT1d9v5J//m10H/+DPbxg4O73WOhfJEgxiG6KiGqxfMwZDPltudHS2jHzQdLx1tLvSST2c0/3C7lwUbVRZgJMEy2dSN3fpblus5ojglufxthf08M4XT0GWZmFMlaino9vvY9+UR2uhH8/hiq1qd2wc39mziC9cdAimWpfUbwCjbhhj5gh+7RdPxmtP66FZWoAhG3b6TysV6/AIgA7u+v5BuKYDdHLGoAZBEwtShsNcODANWUSxH9s1R4U2KiZkXONaGFTGYrlxQ5+XBGViPPBGbS8i5sFMGOtX1cH7XXPxp+dfD8I3AjB4+MkqAk9KAdgG2I9sonr3snzwT64d/Mu/uGHgzPi0leRtXyW3Xmk80Lp2bQeMBtGey2gdNnGwmxpeGXFw64lEGB3/3c4D9CB/jptO8d5QgRFKAY/gDF8Qh8JM7xrBs8eW8XOvejZ6bgCXlQdHbXGPBwh7EkYxiAg6vS7ueEiw/UsPY+DWwlqT0G4SA+tmsOUXTsDrTu9jeWke1pDyMXh038eT3wlkw5GqspidAW65cRGdajKpRpP9X/B0zCvlMAoQChMU5/z7ZFscDITsheiqpHeLHlfKm4l+x2LRMRzn9yMyWZ2IBwGj2Sx5boJNCYkE5wT9sao68EDTXLZj6UxbHboojgNPxnbgCZclnrN9uz2fqLmllvMvuOTAX/zn78zUdmKNiVUyavaJATSADBqM9Rnj4xXYhSSeoXWdZ9+tdEikJX+NhcMqam2kshr2PnLGGM8tD0CiiKARRiNcyI7bzaMPELHAYA6bXzKOU8YBbhzIGNX2PT1sFpJ8lQW2W2HfYhf/68L7MDM3japTpSwECKFZnMF7N6/B286cRr00lxSFaMl6H00haP/zxH5fAUcRRr/Xw+5rZ/DAvYxO1UkMRJF8ukXES3iV3yOzFA5K7R+MMn/Rg4isujVd8BiOHSwResZmD8n494Uuy3MCOOdKhu/HcciRhA9Q6fWr6uAD3Fz2yeUzL/urAxeJyDoi4u3bxT49O4DQ9n9kE9W3iVzwJ9c2W//yplrM2Bq4MECLZNcXxOz4xmFqskJV+epH2lqb8uKahYduogCoJmUZZaynCIDMUdmkG7dsnEGZEsIiQ+QiLWv1K78Gp040+MnnT4EHC2q3TseH1P1QiwBgbIUFN46//Pz9uOu+Psb6PQ+2QmCNweLCEbz51V28d/MGuOXZspcheprGRzFEPPA2P0e45KLDIJmCEAe1ZOT0Z86GUWiSHgdIvVfx04u4P1NMBCpzJ9qgLotPmiYR9G0f88tziM6VyZMQ+XJRIYx5wRlt48lnSvY6VXXoPm6u/PRgU9fOqO3AEzcOPGEF4Levuab6yKZN9W0i/89/vQFb/8PF+2s7tbaSJuhNo0iCTAjmFMAZgBnTUz0Y8gNAdJwDDU39K9g5ifK1y5FchctslJ4mi3zlCqRR6kTmHAGgcV79uXoem146iRP6QLPsYMn+UG7zRzsyRDak60zgr7+wF9feAvTHpiABaLXGoF4e4EWnNNjyM89DxUtwQWVJT1nBkrJ1PloJcA3645PY9dVDuOv7DcZ7XTjXtPISKIF8xcVByT5ZGxCGsSBxO4uCYJOSiVTWYv6a3qRF0K0qjNse5utF7+cgkjKhkOLic5qLMn7LCwsxYMfodm118D5uvvPJ5TMFh1IReKLGgSdkBPjtT1zd+dNNm+rdtfzsv7740P/9x9850JjptZVwEFKyt4syIVzSK8YAcQyYBpNTVXZ4bTFZmJXjjOZfh/ZK8/hT6u+IdB7vSMPJBTiqyaJvXdvdRY8U+sFhZqztNHjVc8ZBroFmqB3PQ/uYWvjH8DXjJWPHpvGFXY/g0msW0RtbAw6Hn8J+e+3YHH7lZ0/BhskaztVZfvsUHv6jYQjJG5EFttPBww8A3/zCI+jQdMpnPN73U0e7a/aol097pWGwnvLPTNBtEAvIseevSI6Y5/CMsWOMVx1U8FHjUiQwIWs3RELbL+mZbC+InDD6/U518B5uLvurpTO/99czXxeRtUQkT8Q48Lg7gBjUedOg+Xv/+bsz/+mvbm7E9se9MpVijfFhk/7mD75dToBG0O0YjE34AmACQcdkb+i8q9VVuwUERQ84X24CeEcqqSZifooeGsEhfc3rkKGyy3CAWL9FqBucsrbCc9Z24JoFv/NH3ic/1ln3eMGz4/v6/qbpTkzjK5cfxFcumUWnt8Fr3wEIWX+JLe/Hz75nPV76ExZLC3Mw1pT+CU/A5uKxdjhHxXqCOot4Ghd++h488lAfk2NWfZ5lroKwou9q48/wbBnJoieNGwmV1nGJdarNQBX+ZMIqWZhhrcV4p4/lpVmIkWj4njZMtiVgS3ZowrD6Zg6MxrHxXnX4/kHz3U8Nzux2j3xDRN5KRMuPdxx4XAXgapHOJqL6NpHf/OPvHvlPf3r9PJuJtQQQsZPU6qQ2XULIinMgsUAjGBu3GOtZsAwKZ5f2B28MocyroNZqShLJIlI/TTSrlxa1tMUAI+q0DYWUG23cGXgiKbsBXnBiD1NdgSwwyNjS2eYJPACPhkabgTH/hPYmpvD1q4/gMxcdAqr1gRiTgcyFhUP4yVd18LZNa7C8OAtjqJhLH+trfyzF4NF0OPF77PfW4EvbH8A1ly1gfPyk4aSmdMtyERiaZJCGVO4jRmoo9NeiFV5vEelSSI8ZY1UHy50e5upF/wSpD4FSxoBg1LtefBTi5de9sU514AeN2/W/F85cXqo/JyK/8nhXhI+5AOwUqcLh/4N/f9WRf/Hfdy84mlpHDKGYZVn42qfDzyAXuwCHsXGLThdYqjms80iRVrKpBzuAzPAiKL5ZTq1mtRU3SQkMloU3aPyTG1Q05QyDocqEjLn0HdPguesn0QFjQNpCO6cR/1ChsBghTr5oGWNhxqbw1Stn8OmvH0Jj1iV2mgm3e1MP8Ky1i3j/u16ADpZQ44naXsjj/BpHL6QxCajbWYOvf/4hfP1zhzExdlJS6UmbQxLiyEhKt8JE+yblUqwFVuplZFfpYvtX3PwUb/8ivsCvE9d2xrE0WMIyOFyEXPBRRGRIPSXRf50yKkBBlt7rdeyhe+rm6u3yXm4OfPPxCogeUwHYtlOqzUTNHpGz/9M18//mz25aBven/HkLAEpsuRGKnmHylFw2oRD439Pr+36HopGCQp/F6XgOgmZokDKQEPIurSm/DmWoU6rakolDCQJJ2nAJNFGbbn8TEMXoK8BCGDOEDeMdkLhMPJIfPvIvrGjGwT232+1i2U7gwp2P4CuXzIOqjbBQ6bcxAbeewXs3n4iT1xMGS3XornSwgDmubcYw5rGSsu/4uhhgNOUgdjfdbg/Nch+f+V/349tfPozx7snJy48ULhRtyUzKE5RsPhrMVBDa+Wi9btSh55Q56J8zbW8wahNlJecw+N/ijUOME3TIYH1vCg8vHPLPow2DAAVHJaEhKVi2cxRljZ6NY8b6/WrfXcv1dz/pNokciDyBxyQgetQFIM78N4q89z9dOXPh/7y55kFnEkZgJByWWCG942qwiHLBlSaAf966ymFybAwmgi4mEjT8h2ZjMLMJgF8AEuODwkLxbUGljBiLh5Hyw6G53zoQI5o0kwrK9IpNyg+GMJgJlhzGu5SdhDT5SDV02e2mPAwrHgIZkYQDlT0Vb7SAFXP+A6gqC9Pt4559jM/vvA9X3+bQ6axXAqdscrE4N4s3vKKLt75uHZaXjyS1axolCEch/hx7zdm+hY91+JNleQhrTZTrGBxKBp1OBx3bx713DPD5v74Lt14vGO8/O9N9UVK9U7FX2wDtaFwWJUqrQhjKzskjWv9H07kQsivSRNXHdGcCB5ePeAOLWGyDQtTrXpAs3qMNHQu8HiYA1USers3C6I33O3N7B83VO5pNVq0IH20n8KgKwLabbup+5Awa3NzIH/y36+b+1f+8ZZnq/jQZBonjwJXOEdJw7Au0g7exdoA0DuSCFTM36HazLRdCi5cMJI3AJqut4be8kgyeMDs/QoSKakRU+IUku+yI/PswCw4RY2Gzq1N8lLOORKAPDpYcOuCUmedBYpfnPw3mSbuuA0JcPoj68A9NhFk7nn6GCIYIla0AW2HAFj844HDFTQdx2fVzODg/hl5/GtQgmadGo4u6rrFhzQLe/+7no6JF3/qTck6OMymZEe28X5GKrksywpCkKAImmZS2MZaozpQgziHj7d2tBchYWGNgyGAwMHjwBw2+e/EDuPKSBczPTGCsPza0DTLKVs3I8LovugGBKKs5pdX6Q/kEFOUumIEZJH9GCnb0FGZODkEwcd1shINVGwGNYENvDRrX4EjjwVaTUP6gTbEc7O0JLtiWmcDEZGbYoDfg4I9hwOiOV9WRvdJctZ3PBGa+lrUDxy8gqh7VzX8GDe4QOfs/XjP7L/7b7nnTjK0RQ4Z8q54tu8gha8NDYktU/VHDMGxg2LvjVsH9FeSCnZZVtyplC+uibfc3slER0kAplaWwUhSXXV+0NpsDsKelRqJaUiK1l3Ve7C3Gf71uZdHreq86AnnX20LuG14L6ZMe5KcUuiSipLmPxYqSw2zLBUc8IUVE4JzB4SXGwcMD3Ll3ATffu4jv3zvAgSOETncdemMVuA5Ot8Zk8EMAWZ7Bu396PU49CRjMLwcAU2CtQWUrH41NwffOKWt1Ufp2HWYauq6cgSjpwEURbAryMCg+I2LyAabGF+amMZAGWFxmLC4w9j8ywL13z+OOW2dxz+01Zg910O+tR6+fD6AuPMwuvYaY3RCoeiCX/RhSQxXWeTFWLgKGevev7eRJmQ4TxcBUzs+MCMSE8TNmMaiRFMw4cWwd6vkGi7wcvp52Y5YsTw4VW9IGAum9FbWVcA7o9Gx1+IGmuWq7ff3S4v5HLSCqHk3bf7fI2f/p+tkL/+z7i53B2LQYghFu3XSOc5y6E1AtwQYr3OzsZ3u/s2XvQFuINggcXFiF8xtjihWXFCsooxgZ+YOhoU5VlN97zBRQzItExSBNSVAPjBHCsnRw1d4B7ukY1EBKHXJOUDvnzTyDMadQfkSMlF7/zL7FY/bvD7PLRhgaUxCCgcGgEczNMw7NNTgw2+DgIYcjiwZOuuh31qE/HsBQF5hm7CDMIBhYIgyWFnH6iyq87cyNaBYPo9OpUHXHMGgM9h1ucOBAjSOHBzgy22Bh0WGwJBgMBK5hsCPUNaNuGnDI0DOGUBGBXdhfixqrQlR5er9zimn4zHzcGotDr2NgrcFgCWiWgfl5xszMMhbmLepFA4M+et11GJ8gCLv0Pml3wBTjBm3YoejIwVI9fcI8bO1NR9lGjPJt8LUlCodMaukTTkAKQw4qwkosTuyvw96FfWjAymG6LDppLAgMVpVRXLo5Be1Ap1dVRx6U5oYL6fVu6dEVgerRHP4/2b144SdudZ353rRQ6HBMBCl0dp2wv3kbgJyB1Aw0/r/R+AJAjmHB6AR0XYIdtYyMtaLC+QWKO5V752xeSYGxF/EEH0UdQjXCzWRCMERInU9FgcSB2CTUl9NI4av8Eo3hr66cB7kBYKqQGksAGzhvQ+RvIwBGbDgcebwwTKk1TzdKA4g03jmWw+uOr1JcaM/DeEMViLqoqIPOOKHv1w9gFzTuTHBMENiQPuQLzFR3Ae9718mY6i2jafrYN8O46daHcPOtM3jo4QazswaLS4K5WcbSsoWxY+j1xoPtlQszesYhTGil/btj001FbOHYBYccBZtLXq3aEKAOBDUo54hvSx2QHQeRxfh4PuBa3cdKw58YfOHwRbFXvDiINHKfC0NcB5KhwpV5pe1GipiPxcYYlUYsmQ/MChcygDT5chN26JsOTuytwyPLh8BG/ZpKRabAB0iZ0mLA0XnIkF9hhmJhYcAMVF1TzT/CzY1fltcPmoe/cbwCoupYq77N5A//f90z94X/dktTLfamxLAzEo0ZUWbtxfUdOfajACPM+xnNJS2jRM5o9wiuGY0mR8Q1zFUx0kv7wjnn0qwvACyZnHFRUC6ODmSZYF3tggOQCMHE0A0ioDsGwXg528ewkjC/+VvfhFtYIcgaZIpxX2EVaUSNIxwVkC4UBE+e8mKq7HHIIZPQBL4Vh+vNhp7VEGFpcQ5nvXECZzy/j6XFAX5w/yxuv/sQBB1setWzsGaqh954BeeAw7MNHn5oGdfvPozb71xEVa2FNRnRI5akiKNkjW7Se0HhNRB7ADjyMnSPZdQajVClcUAaztwO8Yq6RMgN/ns+BBXDbj5xK4HSGzJeIMYYDz7r1j6ZwYTPy5jk8ag/I46foUovFuai24icFzMkY9N7aQM4g7XVGFyvxv7B7AgDRSUSCoVHkNOuo7uwzyow8LQh/2xUXaoWDqK55evmTLJ783bgKEVgxQKwXcRuJmruE3nvf71l7sL/fPNitdhbIwbOOMcpkcbTIUOKLdvcDQT3W98FCIh9pTJhXZdilJIXPylusmR2lHgAJEV3g7PLCwTsJLjDKuPH0NpHz3eFCOWdvQhYalBr5WVgPeVTR0xD4NiEGTns/gU5ZooywSQm69iAU0iQfpp0w1hf3VlKtxqO2SeermyC0i2l44rxH34wpCCxit0WUON0k1JKAW6aBiesWcZP/+SzYNCgbgbYsK6PFzzvRej3CRAH55ocjWUtmKbxrs1rccnlM9hx4T4wr4ch74QkbPJBDpZqRgSGLAReUWmDDbk/gK40UxEvsTYxHkyaYMDiciGJMtuI5Afwtgnx43lUyt1VfC8NDaPzHoMqqb4RIDQ8pBcPIGHwCCAK/BTV+6srirRjUDIIjd1eVLZqGzjfCW/sTIPFYWawCLYIzzmlVCaDykN9RGASWGEICZwYfx6YwOQvTP/RkVcjdm21uL9pbvma3dTpPZy3A9vEjModqFZYzVgicgdE3vsn35+98D/cOFct9NdJxxjjAp00dT1ctnnJaN4FGiMC7589W0fUOoudy0k8YeVGQ7dxbOFjt6FIPaF5YqdafbVIZWAopmo44qrlGUSZGWbD5MCikPwYZKFnN9WexoBQJH+BqEjj9OeyX1Fu+0wCtfwH7WnJvrVldXDQUjW2vycp3juLujmCs86cxnNPtFhcnMfEeIXJCYDdEhYXm/CaSRFbBIIBLAE/+671EBF88jMzoM46322QutHV6pPFBZCqNDlNOAuo8GbQ5qg+4DR3UkYFkPrM5hzakUA5IvXIcVbsjeobJWtFUuQX8nOnVxR5c4Ti4ii5ClL0kqN4EJK6PbXnp5xs3WmAjZ21qBvBglsErGl5X4dLwUJ9Lh53YCmLjwggNo+6Vd9WC/up2fNlPnNy6uDXROSd5xHmRwWSmhGH3xCRE5FT/sdtszv+/e7FzkJ3jVgiw5L91LPlcm5vo8qPOCbBShFCQeHNFPGHNhaAlAKjPlRpU13Ve1w461JuKdOjN0JUcVz0eZU7UK6AfCPRBPHG8IftNxElszs/VKwwDQ68iJTWq9eFcRuQVmSS9/IBCyD12uItGDkRseWP71tTNzh5Q4OzzlwH1ywFE1VG0zT+oAUTCmMyYGaIYI0BxGBhfg5vf+sGvOSFFQaLA/+7Rd/AVRh0w+aAlVyD2Etz1TuZQ8pjsCeB2x+SoisTqbBOJaRBcDfSO/+IIdCK3AWUnAEp4xeo9RmI7iyK7RAVnBEWGi4YYcMwOvohaFWE0HEGz+qvQRed1IBp8YoQg8XBORdEQ56lmoVsLUMbiY5YhG6vqmYf7tQ3fcW+/rt/ue+C/9cSX/Onwxe+aR1+ovMAEVnzp/ctfvk/3zoYP1RNuspWJmbHs8qYkwjgMDy91wWEPqmoJHPx2fjf10S01ASgLvD22CepcsDwWQmH2kU5mitQeMhI5bnpDy2ScfI/FCyulMd8vKlTOymoErgZobjQjoeVowilna3E8FKiZD6W9/qSOhwOrRqFDok0Ei15dRX30PqmtGGzYZNYRa2nQhJu7DwMMh+hqRfx1k1rcOIGC9eE7EMXiniaz9sYSQQb/esY6zNe/YppiKtD/2Ih5P/RIaQsJu2oRTio4BgkZsUejOHj0yDG7+nVIaNg7WZCUYKKe9NGJ17zYVIH5X/eHwJW42V7KxtDRWK2Q3qvkTUCFuQVgepWZ86hnyQ2PQMew5EAvAY8pxSXJOyCwmtkFvS4won9tbBsAibkuykGg6nxIxq8dX5iCMevEZ7JqHYF+zEcARjsj1Wdh+6o67suqz5y81cO/eNNH6F6506pRhcAEdoFWPv/En9u/+If/eW98sr76l5jq8o2EfxiHcaRvdPiAVGvRl3QcW2nJFTxBqyBwbIkMkW29BymW5qWMWi05oJ6cIRaH1Q7CDQ+eM6F1xOtmCyI7Io8fuZ27JjPLYySBlaU3BRueYzOIz6EsXALGBK2B0MMwmO0LtGZhpnT7t41jBPXOLzhFWvRDAbFw3tsPr966F2DU57TQ7frVKagqP8vabYN626V5kw5XDWNV5qgNVqBIy2HjijbJ1W4Y0t/XN9NW2o7UmHYkl0f6/MrCoPqY47JfIzdkEuN3RT1cEJnClQzHHN6N2O2Y7rxI4clFjAyCN5CEAJcYLZEYp1rBN1ep7rn+qa54aLF8xfvk7M2b6ZGlIw4FYDtgNlM1Fy52Hz4wgeq37xs76Dp9HuVSLxlwvohDV7B2DOw6fytEuYeR2rdosSNrD40DxdjcckpDlzoHFw21ZTk0e9gJUsu/bVjsthDNEhjw0pv+JOUkOBKqHwFTyEfnG5uJsCRFPlvVDxMPBwmwiHJNvKKRcpPXcUSxUgyG9pWE+Z9IX0A4nvCSebqj5v1Kbqs2uJAr/Y3vPNT/NICXv3SHk5Zb+EGzfHF9Qw1zx7tXjNZodcTuEhoUqMAi++EmPQW3ikhjqTuKyIERgDL3kHXFg6MYdviso7eA6lNIZqRqM4KBwQxtHMoZCVjQESZTCXFGDeMEhldA03GJnRnACj9xLGowXoNnH4uJBCJg9SM9XYcU7afnJo4kMbEOM8eJVXISHFaYNMYEp/NPCkxhIU6tqIHru/0L/3sA3/c7VlgD6QoANu2idkCsIic8qXbj/zRhXfOuc74uG0CyUFr6qOxJ5TQJ33Y4fD6BxIJCNS3PmnWnhjUiy4hnUKS5lCo9j3r/GNkFClLbxpxf7UEFjld0s/JsCF2OuwIw0Nt1EpHBMd7vxR/Jn7+hYEJqzZbWq86GlDERyP9vf4PtcU2ogQuQ49c+DnHQL+ziNedPgmRASCPUasUxoB+r8JYz+r0bPXXU/hoI+7hMrUbOVDFvylGxZ+HVSZzqdSTvNWgaEoarLZUCxaKdRT1hOyG+A+VR5tGHHe1H1AkohZIWYJXZXNwVJd0WfHnyjdXYWcNcEI1jTF00LDL40z4H4fuisV3CdFWXKJhDgdzk7gmF8rAoTV2cISaB67rvGr31/ZvpfOJYxdgAOD0031U61f2zf/RVx+262ap4/EpJog65Lm1FfWRDYtacpumnpTiJg3HmS2WFyNzMFZJLvawiUUHOg53nGGxigZuRFb4KCTjL5R6cmoX8ZEiGb2VMMhxYpwcYvKHmOzJhohOokmOYb5DGXKpekwZucmIvTKhHgxwykbgeaf0MRgMWq33KKXdSl4E/j2wNuAtLa1E0UbHf5hHfH/DnxGr7yJHeOXQPRuDO4C09cjFMb6XDBvs3Aky1Hmlz1ONqhFToDSqauiWsncgqZJw1NuAyjKiyEKF4YeglQItxTvp2MEysLE7Dev8f6elmhq7NUdAHxIWly44KVKN/TtT9Sr78J2QPZfN/IGIrMU5YBEhIyJmi/+PV3/7B4sfuPaRAXe6PW/hzY2nX0puQ5NlUqnBChRQ35ZLBF8E4efau3gvBAIZLMwOgMbAwGZkNhULhhVONFoJ87akdE09Y+YcdvKXTfGAe168yZx15EnEVzlOdm95lmuFarYcf5j81yTlJORbWA8HmkhTDsXNyOiSxeLguAkfqih7MSqwiQTzUWbiQZms2hBs4ZpFvPwFY5jqS6YGHwtJkBG5C6IJ0qzksVSs4eLPG9K8DlaaAbURSp8B8ljIkUSlDi57/giEQMamJF4KGwRB5jswKVp3lPkq7lUaLRXI4Au0CxoCjyUxA+wisckEnr4pj7oYsMsjkBGBFeM3JxrEjPT1aH6LnHXp8avGt/iqdjphTFIHJ3SngCavaAUe7HbEab2IooOBKkGkLPDixo59uag7fOjmqY03fenA7xGR7DpvlzUf+dNrLIjkU/fP/uqls+O26Y5zG/TSKqnUsmmCk0JRi2oZuwa1NtHfAAhYmndoGgpIeQj3RGnXFa2SRKTg3qXHk0xrDVcCaREAk/h1RiQNy0jZ7nHGdFHJVjTG+MOo1k00IplYWI75xSVRnCnzDUakzkSCknOMsU6N0184BSPNMb8FbdtdAHTqXY5FcZSc2RiTUXtpFUqjV3nDwNxKyJm02X3trgml2Kfk67e/P/U1UhHiAhAdNUxG38HRt/cKa8YVEs2IIr9Ds2GR2aXqDXLOYa3tY201AedcFMTDRQ8B8mOOxK1Z8sukUGuk7LbjOpkF3X5F++5t5LbrjvwjEdm4+fyznPnTj2xqROSkKx+Y+/DuAwOpupURxylAgXRAR5RgRGllWH/ABZKPU5bHjocUpXmPG5UShOVF8ZHONt8kkWsOIrhgimBCWxj533kdEhd1JrdgnL9GcWuLgGADMBmIGSQJA4iRYEZ5xxiOxBtTyHUjuy+BO7FtRUwdZhgJrzd0LazGI6P2+SZkzhUrSzaB3BI8EJQXArTohQM9N34+3GDjNOG5J/e8Q248oMfhz58LQWQgMkofHRNUjEgah0RvJv0wh5Ur56LPUaUnkp6r0Ccpck5QXsYxIlp6ERJ4SMrlJ3MFJADSrb29IstkrAA51RcmULf94TFMXn9vMmMB6vXmQj4Ce0o6gEiXRtIrGPJftwrYl6HsOMzhRRqSYE8vMEzY2JtCL+Ak+kLzXFhGI5I2UKw2In6qiQXZg/AkNiQPseGBdYv3T224/6qFDyBA1nL54fqXbq3XblhE5YRd2d232RPxh+OA/JuQnEKqmgpGu8mrKzO0dm4gmDu87Gm2xQ0e2sk2ffM4cDlaYSUjrQ8v7lEJIywdKBfA4o5gKbTkun2mAoZUtBf14OpLzyjLs4JaWnzPuesxLcxieGtnwM7hWRs7mB6PqT+UJKeP/ocUGxoJnIOih4oYB63gIyKt1iESoJhgwh6VlEx4NEOLy8+f1ScjmW1pVAwcFX9cCq+G2JXawB9QIA+IOKEUFEBKDCkPlc03hiCaxFfQ5KSIa0UiHo0IhyXVaXeYcEJvCuQEDOeNRWMKdnATSgBzLIRqNPaybqgnx4JZUHU79PDdjdxx46EPiIg1IlJd8+DsL+w5xFJ1OjRkrsiljbJea+W2tBQy5Jx4NZOILgr+uFlioCHMHlpON2x84U4DSIQU7SwqdXXk4S9Av+DmynEHLmVibyvvLq+faUSWr9KRR4ZebJsdZ1JG2u9HmSyymnBETc18+WPt6H1BNCNGCTImFS12DidusOh34uUhWVGI4zPgzDHkpEYoSVHnUK1+atP56ACtHwkioBVMNApXZr3aagHAra2MkKi2XN3q+rNPBCvOf384JIlDQKrYK5+DdFEETAUR6ymozK3VjyFF/Bq+kFYCqnXKUXkBMtaYHtZWfdRNDSEHIee3ZaEoB7oQkn+lmNB8sxpfpBiXyZJZmiXad9fiGwFsqACccMfB+sxHFiqiXmWiIaYwDc9SbDKhIiqnhErzBMlIjgmyD8AVNyyBw5/zXmdH9g/Ajf8drFJc/JcxwVklavmz3ZMmmAyzd/Lf5kkySPMUSX5DbGB/gaJhIyBOebxFP7Ygc0ar7bS66Ia2P5Y6E2LHYgIkhfYuQlisd8Tq4bLtLkAELMEyNc152vFG9yo1Nq6ZgEFmwK20sz66FXl+f7lFtTXSUh+E9pkiwzOOV2A/3rAGK4znesZoryis4lIXHzsnwyYsZZQJZ0DvhWzqEjXfMIqHksCLbMKhDUawrEKHkIA81a0Vu3y1sKXYngfQVjioX/XrR8lgleBXKORF6AhKUKKWfiKs/CowNnQnMetqLLkBxqiXzHWowH8FjsIoZrymhll8tiOFQssGxgKWmCwbXjw4PnngluW3mDvQnHnvoDM5gGERplSRVHWLK5HsQ5kNFAIjOW9x0q1gcnUTo7jxJYeerMXhA0tYmvUJO8V6qMXJHyl7KYBQWpF9VYg5ErmGjtr95hsq31bl14ma87aBFALVMygjU3NqihuKQLCmGtnMcLFmVfLUePMSMpAl2UCyssDGtX2QomzTCPTu+JKFaKigUts+naFQfzUssaj9ceQLyIjEXLU6jrgKcsZjNuHPGY4ZBIwORcNrXRMKu4ndZpB0I3oJkClVfSPBOxq5TtbKQTKmpU8ZhiKPn4VZfj4Ogi4DJ/UmAOdQowGHe18CCQ9h/eeE4YJ0PEEzLm76TYGDGWOlmevZe2+e32hu37d89gODyoi1/kikHXjeAwvHWcgDNMLOt70arImPC0V5btSMu6TnIoqqOPb2WuLfwGYBmNm3AGNtAMCUiYaI2v1yBrwoerKFGyYQKhLgE/e9IZfds/htApA0ZSiZdqRKzyBS/y2jxwGNjBtkSWrKf+OSt598ByWvHmVoF40kOGLFoYKUXAqPRWhOvydiWUOYmDBpVhSMpsxKy1VppYc8xa/rfX8sBFxiPQbtTEYztKIyEdMI6HgEWLV3Yy5qACvOvkm8/7DwlRYHMDBSPcsyAHBSIjnaXxGJfJYdoCLIp78vjbkkdSe0I3RclVCBYfnXS6VL8QjKWsZ6fAG0gdKEMF5OUxfrqjEsu2XUEouAS2tjidJzVkvb8N/stEzd4/VkKywfISzMLL7L3D6z+KKHagOYyCGm9A2QmDx8RSMP9i8qoq6Rbx/bkozHG4+4g9pNUfbmi70hd3DggUUP7qiHqiB/iGLKqIeK1UBNaiWkC0BG63KjaNJenopbCCrAVPsOmFhLi+0DqfdLUupMfohQ8OSLzNMo5mjjEJIJKaRkyNHxhhQ7xMSjo8QnHQK6VmfX47g48yvzqjKZZOi2jGar6b1UYxc8uh5gbjVChG+ASqReG6OqyabVWfkCQIq9h3T4V3rto/eXlMUqudhjeF1LwNCIVgo6aDh4poVf6M1DG1xtg9u5qHF+fc5hY2cCPSY4NwjrwAgG+sARB28awtIEz4gSboiK0dgNDxYJ+x9sTjBHlui0mbkBjJCJO9KkSQ+JmJTaDW3DpXdbJq26Uj0u2j8zencaSSOmiwMPL6JZRPKoFylb5czeagEoLGUsV0u3zStYOxv10OWVslf2jboRTWhPk5EFMjpsWjoHIhN8DEokP74P1Lop9Ltiw1wYC87QfNDWJ2iGbDyUJg0iAdfgbFt1NJBu5M4movTDbZA1YT0Z/p5klU4mqC5RSJ1b6GcuLXGmL6jOanSh3BlFzCB1SiM6tPaTZkaQneg43odR2IjRyEp4Jo1jUMNlh5ZedzZQZSoPvlnhdXNkgga367qp0XeEk7qTcM0ALtrYiaaK59i0+E/DnKXqrNy6jdj5uSVA6I1mfgknLy0zQjZUAmRIOMQdq1uoyTe+cDz80SI7vNXRrII47IxNuokigEjGqHaPQYawdFgw88gSep1ueMNk6LaMlNAIwFAU8IQRxTUO0nh/vCbmrcfXJrb42Dnks5koYgcFBxcBnPO2zvHmb7WRlgwMiz+o6soy5FAlYRRae2sVRS5Z85BUYUFQxJyPbtR/i2uvMrU9VpyFg68iEQyjJYOl7Dk/ivG34uHhQM2GKgCZmMPB8caGi8CEP2NCDl4ybVHUVIox2/G9EwMxVhXDiLpn34lYxK0unIkAFhmeGfHmyEMJwpryho+drSku8nYB1EBkGxdNBZgjK7IDlQ+cPncOF2J7pR47piIEJEKixGnFF63niQykabCGupiueqibJS8iihduwEr8aIACw5NwlpPuQjwzVxrC8jx6Zv+ReWGy4ZAYNZOVs6I2Xyj56fHmDSsT0eYEK8yerQgtIgCuh4fumkMFGw5WCdpFbf7RgBkTwB1jTKD+ojRMiDvt9DREIMg30yyibjEaySjT66L2NULRgDDsrbO1P3srNGE4JyWzrSUekeT2E4gyMvxZ6HWoSxr4YSbjSvLUUUXgaEnE0oazA6Id3aByQlLwyQsAMEapMTU1IPweG3wOrSGF5g8Dhtk9CgV1WFqc+5FtAB2D7UCtrcgKnWO7WEQ1qacr05AYaFSk/ahCY9TrGDVxCAk6QjixOw1iQUONZ8iSFJoBRmnOhdYYypEoRISlhRpm0YVLRBThY4i3qYURahRQu9U4swjFQsGtRrI14BVvgoA6HTx8/xKOHKhhqhjmEKtlNIOQozPZQqqKaEFEm+gp0T7K3zRG3/zSAolEhXm0dtVpJClEHXElJykngAz55ipVRt/oe6EVK1YaFasd8gro0ltAyp15GhvEBsm2DC/x6Bg97zFBAEqRa5GRl9thgjYn01qFGJVdAJykn7AcouEp4EaR99ubldanKJkZaDjjNoUhjGSnAkSMiiXd0BJAVBMNX+XYb9GIFb+iNReIEBRwdowPoHShihLxgtIePkhmYMr0sL4zhoYHYDTpTSkYNiq/QYLjdoQglXtdYD/GwyLFa1GVNT4EpFZ5Hf+BkSkBu9Bu6WtThJVkjlIF1UISADDWoJ5j3HfXEdiqp3j/yiikpfRCixAUCw9L9KUn5TKs1iACEOUMQJI4xpjiQy7MLqT8xwaI04qklRuFdGPDJfhjkmNQ5CXk44JQ4CK2EBFcfRgMMgGoQKQLtiISNbm9khJlGjLq1l/5iddRabnoR36/Z1Frs43s8iSSofV0y6WxKXsQGkWSEihApgBYUdLSw3ZAEjiqXg/Hw0/Z7Z0zPVc/5x5rMcnqTI5FlBIpbduovBtTKdesVv3rcTWutSjIUWBGFfWS6SmpwFsGTuhOoUNAjSbZ2psW0haJTv4lh8FMfZ5xpDKVLQ/XcLeU2yJD5AlAww1K0XvSo7huRAs+qIN7bzmEhcM+rQaJgaYAkiSxHeawD0dxDQv7TVKWMZw4z6U2wzJi/SD4myYElEg2jPCAHQdFWETovQlqCMIJWm1egZeg52ok5iBatKmjxm8GFRmN4CiUpiRH2Wcfz+fUHv00KxwOzii7KBHF9C+vTdG+rSIpFsAJjxxT2jZvo7z6oERCQuUeP41bFMhr8dcsjRSCHf09ELiUelMeJikJ0sf/vtLR+6+iWBj/PPXFYK0dBzd10AaoIBaJwSzJKwgOnBy0EjGMBb2egRnvGAE7aCY/JKKGpuShx90m2IdFgFsl0AEqskmTZOI9kGWdmocdkk46XSzNAA/cNoNOpxsUUBYuPE4UjTxCSIZIyZ6CoTJ3PbrRwsAGEYaIQyNe5mxJ7dQzBu8Bw+BkZMIu35JytU23cXDogU2WZBIkyhwk0jHg0kuGY9CjCyxCE+zDKQSQECx70QhCpkESwWgkmj2PPTovmQQsZUeiYwUWH5c4SD0Sxkhi4BH7DshQZsdRal0N2EmLNpvdkrR4KhJ9xNPwcofG5SE/VrkqPPiijTNL2tEkt4CGU4dnue0bELoJykEwbbp/LvSez2DSdiz+e7z+osWcBbWsSn3XZoqCbVYYP6SlaI1muyQGG8wEukyo3cBvREzki0jmtcSi4Ly5aBpppYGA0R03jTlpaoyMtGZcGQEQ6f29bmMStD2KXz9KLKFWWBSMKkmyAXRnDPfedgRLs4Ax1XFw5HU7O4JPLmUX48FKC0MWNpg7QVfatiipVcCMOnAr7X8LoJQoBZBGCygij5oP3ajJcpySYzCpdNgYFCnCCmKJDzuGHZEhLdntY9MDUcENlFQY/EyZV39EFkU9PQq7TmTIkXZIf5acitvgr4jymtTIeihSdBTgM8aXtQ1kY5mgUdLg2PZnO1OK2xGUa87s37MS8hi57jRSwzIcud56HYHy26cKa6oeXLMMljpwRzl1koWBTdJOUPQhdGPjfSwtNZca69yNk70KImnnt2KrnthphROyogJDrX0kAhEtxDQeCgr5eQk8RGIGLs4I7txzEJ1OL7H42k44ogId9QMkQ9yzEkiKTLJMsIGyLMsPeLmq0RrukuAhOaytnFWH4KJ4eysSkbSso0hRbtr7bfG5illFxmm+5hUoqCvdJo8JDqTSZAOcGCNBiWeTRoIUtu0JQSsZnSoATp3bldripIwUv+r0EfPZZlzUZ0o4OuZWjBfI4HGbkFSMgJoJWGBFGcZNuJWMutkldaSkOudIa07FbiTAHR5Tx3DswI3DtB2DEcGyW0DNdR4Xw9qT201MiOpmEaBT46Sf6LB59truzev7FsIk7f1o2t1GDzZF6BOhQL1lxQWP0tOctReBmCRLjaqqYj2X2WNghrEd3LNnBnMHGl8EWJIriujV44ofsm/NTBo5gplIyovXwpIASJGfZg18KKOJlGCIIoCgYPSJBiPZt64m0lGRxTt6E2Lj1iFxlsrNC5Ek5WS0VUu233GHDYIYAZn8JwVy9JvjmHZqR8doENOMmRLlOnY2qakN14shyRx0iWIgKV5j3OODqNiJDxWAZHmecYUM8uWdfjTKoONAoGSIyFeieRpLSJ+FGnPihkfal4aiQ8egklzo4xYrrJ2jNgHRH4IKlW1pZ48CXwEL0DD6qDBtxzBoltHIAI04b2YbrfWRfRhSNB98LkQ1WWN8WnaaF27oXvLcSZ/q691dsp49ZrtL/GYjCUSXFlbKvETxjSEfsUyTio0mtWkIMxqXS0tDhHpWcMvVD6NCbyQE1v6YI6gmauMwnHGv6ZbBICJwElL7xSohpwC6MpFJAmhFzOHQI2EOWg5qyMAAJR03Uk9lhRHpKI+vjLRMD4RrySDpKOHJYwX/yjEgP42lBssgm6OWXRqLtPIcIlXbAIYCmi9HaVFyIKh/7ymPqYQhYRFodNGTx9MCRcfjZK5iEvpOx0Ly2uOYlDyQ6BxVYDzh7ykzTbLvQlqdM2PK9kEC1K6B4zp7UFKL36E4iI4d9dc69/wzNu4xb103ftPzx7kBO8PFQiEcBlJBCOnBzeKN1LoK6fcqxDJbfw8Kq9mH1YPMwfI4ljqHqLwzvR723noYD9xxBP1uP1mFx4fYsQvR2RY28M4NIuMuAJhMMOwAF//qjOTbFtlI721FU5AT6Iik+vJfN5ugpOCTsKpJh9BxCJvQtE8qC0wQ8fi21ru3WDZp5We1QaoiV5EY79PPWTnDYY0Zta/q5R1jPJCjwYBhNLHBpYhgKT/QRiITzWYuRAw/kZAhGS5YG+TXNCJnYYjQQ0jBqRbxhjR+l0+kxg0PRvv534VusfRuMC1WDLV4/smeXdN9ZcTYztHP0MKi45+7sOal0CkbF8Ya489OxrCM9w4MHTZHJZ+LxJxgIycEy0hdd3nhUQoKZfLI/5jpYMz2/EYrjcucrOuZvWyIgq+bhREhZyafXS+d9OrxK00H+O5zu/UtEz1LktI91F8rejPQYjXpaqv2tyZXgRBgoAM0pcxWk1IJF7sN/2s97Ln8ftRHCJXtKBRT2Vepm0WKgxL3xlIKT6DASo6DZwwRMSoqWgr6qk1uvTI8Y650fgyV2v1R93hoFS1iwEopec2KPZNaVRNncCej5C3KzUjf/o9OB7DSuirGmmX6UxBlsYNjGbaETKqEzGAkGW4v2iYu8ZcsZd4Et4hOmnmnDUEKBD/cYSaFh1JS/rUZmSyaEKYLJGcFoFA2WAnAZ3YbMgDZEdqX8rkUKVl6eczzhYBFj9ZQz3n5mXG4kCerHhDThNS6TyRfriL+FuQGPL7WyAmndr8GYJ8honrTcyavfv6khTTMQ6p7Hb0qepotrcH9aEoJ3SQjWQxBisDC+kBoH/jyWRZmUMdicb/DzVc9gK4dz840Kl5cF45RYWwCDIFQpCFJgTLUMJlVFpSM6QAnOSodhTM3qnXMwGmhLEt/bwlWZr1G1Fzk+GpQRqKlxRHQ4BM9OibsMaZlUQASFMszA5JUrH5Lr/0SodTjWf4pI0dtnEdAu2osGZLoU5rHMzGoJHdpGXfStLVejyn/NkXeyhssUp+BieCn8Ej3JWqlfhpRMlcJEnvx/7ggoUc7KbOFpMbvYdJ0UQFoTAMXgkSYgm9ACCR1roEIsNQ02Ph8ope+cu3/IqLGAMA7Th7/D2eMzdYYLNlM05bEkqOAOFPy5Q4AWRRfGErvOIWWPlFyo8tOi6Thi0S2aCoix9JqkUH9Cdyz+yDu/f4BdLs9v0MnhGqnUnkZhTNO7gqoWF1F4IKoCtU7IseSJL0m8MOMGCVGUtq+9AFn/YQpcOBc2gvWY9hXF2Eq6p84u2kjVgqCGQjBMoNU/xAfdoSRwUjLH/9xIf+hyJMLN6ZkQ9fU1FYK/4j7ewqdVFnstKhMSO/840gQPBjCP219QgLNoiGo0gbkLiC64/i0XKs/DyrdhnTnk3YWkprBAv2PyYNp9Ue28JaI7f+wByWrDIyc9JRGE/2ZkZQblxY9V7sHa+MlEUHXWPSqCrUso5ElNOKSdWUm0hG4EWf6jemcdOjG575l+hvANmOwfbudJLrhzA185YmTfmgQfRWFGTUl1Y7imRdtYbl2SrTIIf+11kCW1jGRLktprCBMYPd3HsDCDFBZi6Zx6fc89pvN5eshe6qmLW8p6Cjlw7m6D8eEEdGwGGlFhUrrlqD2bjw+tw4sjWq5DdpamywU4scHeA2xO9VeuSVgKSSoLC3HH2lRypVvX/F9S6GSHK5BBo5Ly6xkH68uD2obRMVaZYb3ca1w6XKlS0VQXQDltDEhB6sdOT7sb2VkpfS8RJmbuPJYFgLRKSLn3k9yrOp6CbAM0PBSKjSQ4B1IwNLSkmx8vtArf+qEPyOi5U984ues2XnCCQSA3nvaSX/8xvVEMlimbB/kbxcWqxMbUyaZkM30yvTrkemXoX0ZkqXlKPDilziv1Qr0tlthcKTC9bt+gC6PwybraoFjH5XEJMkixH9NkzILSYwSFkWXowbCTZgPTeJwe7Za7FwCC89Qus7a3vHJjEPRlPMNGd1rTCEk8mCPJEfgNtBT3iIEC4Mqxngj5s6LVxi2TCbNY0X6jwkAcAIr4+0n8fONDs6BvehDS0xhdjlSZRh36yaOg6aU6ko2GisiGAL9mSMJKm0aKI97oZhysLLzQS3DegpJ4ispXZsDI8iYHMDC4XO0FMxOopaDufVoB+ud4Ihl2IXOWEvsvdhclPgrdk7lM4Shn4MqcPHP9owFsfOgHzs0XId05kANFmHuLdoTT5v7waves+HT27ZtM7/9269rzObNm5tt27bRyyp8/q0b3XUbOkxwzNpJpwzZPEqM1IpW4KqtNxmQKg5/yy+ijVKb3hj237mMm7/3MCbHpkZIatVkPOJrDv38EJIZQRNVuER84nkcL9KHHovLCpbiukBw0Oez5JThBIaq71n0rB8xA1P4CJigXiQFRBpRXAHl+UBPVCGQvFGV0LYSl3bhhFxkqcW2EzVarbiRUN97JBmRutnbqH1un1s3uOo6qLUqjV782j9DaHTXI6OwEyLVE2hlrEutfbIKDyMq0HYYGmEvFzZSeqM0TBzU7Ns2DkDJ28IKYdAMfLy4ZM9gImBxcU5+4pUVveasZ28loofPwlmGiHxW+OnnnUdE5B4Q+Rd75mY//xc3z7OZmizAp9KKa9iNJlW+5HAqmS8aKa2U/3+GZNT1v4JCwl++Dqbbx23XPoKpjX2cetoU5ubnYJOqTyAr3KIrNWLU+rhjWx+RBJ91KgmlF4U+UZDpEsedtCkPXTQVbYN/lBNnIldKayL0Q0K6zYj3QXRpylzcovDFVak8HvrvCN5FwnCMCTHmShDDiv2mDqKS/OcVsor+NmmCFGUbr6LURr0SZRDStu6KvAMyWd1nRvy+vLoOhVaGaeOEUaMFtTwWpJWQKSmfkWiY/agc1TNgS1nhqLMk9KDF7fuM8hZelJCqIoMlDMDUgKVBI96ZatCQG9vY2FNeK1990U9Nfla2i6Ut1MQOA1uI3HYRezLwhfc+ly582cbK8pJzVmym7VIUg2TyY1TgSayGJhYBKJNESoKRWNnL/YHxDxVUKJ/KxysAZONg0McNO+/DwXuXMdEbD7vP+EaFkBId0jh0CFpuuS10MlmRRbOKsALUsydJZqER2Wwsol4zKYYYAkhoTZSDBk5glNMSAGV0GTcOpE51Gnmcd1qKe3ARlx8czRTE4/ACLMi3CkPL1s9qgxPkvym0OLsfQbEIc/qx74iiQUp0TU4mn5FBF3gOTCGIRBmgwIS0Hf33JCMFSqNA7NaKBCV9o1PmCsRTpuFcW+QVxl8pswJN+v0msfmEGZZHUJLj+6SMTD0omoFoKNjCpN8/XMHie+41f74LqQJw7sTBSQ2gAYyT+eXD8tK3msHPf/RF/5SIeIfemsR/OQcQIpJzTpn83V95Ac2MyTIJiNupLKLoqn63yLnllMz1TyIEwz4oQ/wHxwqdTqQ9Y3wUtMncgzQeUwt5tAS33MX3vnE3lg4CvU43+/2DjpmAo7EFx6Jk98emytJjOFTUWtzFr+OjnqloDdtJQyhQbozwnKNh0coxHGge3Y9ADSNBZQgUwEikFRVDh2kw591cjKaK/3AWEaT2uGTJjQYvo0N028H4eD7jR9MPHo0ynGdwWqlJPfbXYkkbr5VIWDKCfplEeJQ/EVYJ1FGKLSKoTJXyLxk+SOTw7Gzzojd0qje/f+O/IqIbdu7cWW3ZQm6oABARbxexRPTQr79y3a/+rZd2DC8cYWNSQFe49UTHrmYPQMlpFqRtXmO1M1qiGgMVMEK5FzcFATyjXK0lZhNUFZZmLL77jTuBpS56thO+pMkectIezhlCLr59uY/RN3VBRFAW0zyqKR4x8LfSYiLgaJWAiGBDll5e7Xg2n9/fRvMHlhyIiqAtQNQahFfsoulJcetqIxXtB/9EjAKtGxcmte6eIWiSM7Iox11DUoSKFCtBIWVnbcIdrKnoeXsARfxKK0LSgCwlXr1okRePtqYbFcqSKOJCypcxg8McdvWUukjVLUVvRgickYDUK96La0XQg5RVuBRbcCidhFF7ZIqxYJSZhswegbDGhq/jACOYW1iq1z+PO69+99j/Ou2nNv6Lndt2Vps3b24wqgOIo8C2nVI9j+jLv3Ha2L9934vGK7ew6IgsQrTOig+HttkTOjrDXTvUFncYtWPEFYswtmyh/TG9Dg4/1ODKr9+Jyo2hssZ/OAGOoRF/8zCXnjDsnibDOEH4bVGB9qjWbK01VbLSzvlBmeetkpM5yYU5se/iw5rFM5HtGOXFj/egH5UKn2y2NOOStHtvq0NpPwalUrJ0mJaY2ZM8KI7yzaiPjNqqPIqfms3zNo2QCGnHLr2JYbXxi54Tolyh9DMkI7Wu+UkjGUHICp8VA6X2rw1q+30WUyaStYTegXIe4McQOe5PNWO5Xmr6G+c6b/3g5O53ffh5H23q/8fswllDIpGq/RPnb6Zm206p3rqW/mD7A/PmwJzZetl9c7WdGO+w49YBzbtuj7S2oBFCCznNrEBpnUFRB55Y/ak0qyEbQsKHfdhuF/vumseVX7sTb3r3izGwc1hualhUamfeYlPl2SI9LKV2xqhRJ4OD0ZqMuS36oaO2oglWTDN7uLXYz7Dh7k+HQRL5CmkulCC9NQE0jfNyO88gN2ecsZAnoiJQGWpBVCL1sehq4lPhaSS5cBNaHkdqoxHfU2kZymYFp1PhnxlRk0jeKhKYkQSq6RmNhXxEqU9Oxxr8k1aDpy8V0UVq+H0mHMsZa3jjsFKj1s5WjYEyThwchdWf+A6SrGB5MN9MrZfqrL916u6f/92XvIuIDsk2MXQ+HbsAAMB5Z8Fhp1RbnkPn7rhvdmOn3//wztvmajM+3pHG21jDBEecaFFOknzaDYUYaSMgYyDGQdhlw3uhLIcVr8ZjE1A2shBq/DdJCtgxgHFIBqQU9v+2N44Hbz2M78qteMO7X+xDEhoHGM/mi+GYRL59hjGQaG9OGQwsdgJx7oQrKKeZDXgU3wShbAkGyl7wLmcLSqv1Sjtl41WZIgwTvOwkkGAidVRz5kk0pVWSMnNFQ45HffBZ8d+Vl70o1+JQamMSj04wMpr+rd5DCuEXnLT7Euyq/R8q3mFDeREiZU9JqtWI+ZFGskMxouW7unSKqAJSLC/OP0fBuzD5/kdMKrAuvfjHFTssJfsM37v6EELBj8+a1ZdREU4sUP7ACciW1PZ57kUTBD5NYBpyxGYsoxksNv11S9XmD/7ETb/80Ve8k4ge2e5RfzfqI65Gt3wkIsLnbxPzS6fgo73u0gnzc/iZ7z245Eyvb7VPfRRCZG00eeKMIaQExExpS1FTqdpSfiNEBXQcq8uOOmcnDlV/AntvOYLv8m14/btfCrHznpRDKoQ55AdE+nHcFHgf/nwcWSTlGaXlIkvBfE/fR0uhkXGQnMhDThI46dlyyLt94w84IQeh+DPE4b2jxKhzUdUWBEIk5YgT1ZeRkkp4EghBK8zRed2rb9KwCOIS5MrrsSgui34LDEcUBiJ127Io72EaGqvi2RqqeaL8RUuq5mhiTfr8m+juP5oaPbR1iGvJlkNy7IbSZ2UeNWgoxEEglFLA/XMSw3MgYd8/gBBjUNeuO7lcvfOc5+z+9X/8unT4t6xw+AGs/KqIiOU8CBHN/PxJYz/7j8+c+KuffE5leXnQWAOQMgMR0dEGeUURyRIGJviBBKQ/3MQICjdOEV4ZREl2CmlcpLKtTnbkHgm1vQk8eOthfO+rt6LL46ishUijmMq+u4BUHogTXWTyHt0oAo9lCtJck5Bsl5yEoykEtdh8aHkHxuhslwxITCKWqBViDGJRoiRdOBNuxF7FmGbvFJbi32se2hofjaR1rKew3KJHJDyGUQyZVihWoHdHjgXRhGJsUrIwO+X2K54VZ7xDq3LqoOQm7Mk/MV9hJbZCVJhyAB9JRbhhiDqesxxFbWLKMNVIiPPAn87CzNbPqdDFLZiMGgc4AeJDY2IYS614s5jMlI3JyJx0DrGgcPwfeVViXbuGuzN28znP2f3hf/aGdxzP4T9qAYidwDYR81M7d1ZbTpn8tb//2u7H3/WifuWWlmuKu38N3kGJRojKZ4+yLXhyYSFSD5kO3PDcAP9PRn+lWE6hUJswGtixCTx42yFcGYqANYTGueREJHFnLaPXPQXrMWq8VQZglN5ziFwqsvrabMMhlmOe+ZJ5pvq+Y+SUhN16aodaYJP2rYleipC8opRWHEQZjvVE6AOyfI6w8kHMW48EBiWad3z/kqU3slmWLpzQdt1Ay4J+9Jyc1aWSrLGOb7eRPwdIWTxFgwAjXJo1I/QYA+LKkWSjiFfBL7Nky0so8vEiZCzXyw1NzVdv/1sn3/V3tr3+nUS0T47j8K84AhSgIBGLiJy1bWf1y6dOnLvjwXm4urf127ctNtTrVFGbLwzlmR5ALxNspOItRs4/DFYBgZQZUip7C+LKFZAEhWH5DlPBvnJcoxqbwEN3zOI7X7wFb3jPS9DrNxgMBulwSWylJSPEHEE2yeRWabePRNBoofaPw4iqXzrRlNNn4p5DKSGj9j8MpmJyFFajWGMCKQJJkLqFyLYzeEJIgC06kMcebBHNfSx8MfZxKeZKvU9CbRRdW6SNfk91F0Kj2HFRUKbwHA63c7Gqbn9/JoY1i881NJJdrkLpl4SHSDLuUG88Mr1HhsU+o+C/Nq9BSvFUEZ6i91cpBs7zagaD5aZaM1e9/gPrb/zgP3nNh4418z/qAqAwAXfWWTurX3r2xLmfum8eru5uvfj2pdqMj3dQS4vWi9QiJScT44EVsYE+bRhkKVlySWoFXZFSnEMk8263TDBQZJowZ3f7fTxy1xFc9je34C3vfTnGpwTLC8sAVcXN7B9Qk3EI5DWbkfKAm2T5xQXTTbfL2VG53DxIkB1TEBYJh69PQS4qFjoeEwCMzYalNm4DogKNsoeCwMGIVXbTZnjX9bhKgi/AFkiOQNFyW7Q5ZqThhtGNwyaDQ4HNq7SMuTFCkYtKxrgoV+zJtBzyaaD52RLxoyNMKwhlmEbe5pukp4Y1H9+UDD3Jl43HjkycDVsXgX5rOYyR1ot/IgFo1G1PVHSMmgrNJIXhbmL+RRoVMcQwmsFyM7Z+sXrVz2/Y/cF//tp3HW/b/6gLQFEEdu6sPnTqxLmfvOcIgXsfu/jOhcb2JipulPcfx5y43J6lli/eJsZAXF4bpbA3VuMDq3z6QC+OsxaRpu+Ep0o88OhY0OmP49D9i9j12evxlp87A2MbJzA/t+jVcgE+Fynb8xXwnuSxJ0o2mnjmslK1D4m64addABcLRqEOv0jPY2xdDdrSmbjPjjiJabnvpIfrSeAC6C1HO5FJd3EU3qvo3hPddAu6vLqNGSMcgoYUkaOIXYpqHDEnaXeHGnykFbsKGurUqEDhLdkQQx9GP6JizSotZiO1QM+VPo9EpQ96E6hDn4J0KaQiE3zaLxzIAnVTN2PrF6vXvH/9Te//J698FxE9snOnVJs3U/NoPs/qUd0FsQhs21n98vOmt37yniOwGPvYt29faGxvvBJ2iZQb0D7VCUT2n/FdaxrNKVd25mx9FYMotGbY5HlS++4nfkC4fU1o17vdDmb3LWHn9mvxxp89HSecOoXZudlMMxEqYRnJIaIZ0Cd1hQhQmm0lMQkl73/OUyXHmdYFRRflOV+FnhnSRB91M8T0XcVbz8xLSqYieS8NlU9DT9CBl0BKI7WuysaJOiNRWrNwBHGNVvrp75Go1cOptayW50ZGX2BCZDPOfI379zp43wXWHhWJqaWDD9ouUYTA+KPU8WlOSHTLo3Q/eQNdiZ1nYm1mF2MJaz1BipxJ3JL2pRELHFN2kqagg0BrdTkYuGZsY1298membnr/P3lFAvwe7eF/1AWgKALwReAzd8+LRW/rRbctNlVvvBLHaCfJxjdfop49+aCHOuHMcOE2w/sdTRBMNE9CNh6NCj3J03a3V2F5rsGlO67D6995Gk49fT0WFubgAtqeREis4stZgXUEz03gmNaJAiVekUAOAYlNwByp2d8DeNFIMhmPFU9CcgQKjLGU4Ripw60YtuEUAvPEXv00OumD1PsFVm63kawUO7joqlSAgyNWdymMtBAOJs6DCdHwSWUR9+7RG8FGc9jhzyV9pmHLo59PEfFpS4GhgAhKkkC4KQNtoPgDKYBQClsZhh/3iNWnQ6U/pBSgcPKqCvCwH3FEPFFMICDDqOtBM3bCoHr1L0zt/pmtL30HEe3bvn37o2r7H1cBGMIEnj9x7oX3zcM1svXb359r7PhEJaze8ZgAG2PECMFmKlN8879LMp+AwBuKsmqX9J+x2U5aTPJoBNjz6ImygYupKsBZXPGVWzB/5Hl42etPxfzgCBy7NDNLyA6o9AZdonU4qey6YUppAXy2iLCi+PzFzj5pabIgidRfHENNOVqyaaT7GABczvELo9hxgnbHtQRQwKVlyrl46tdMGOtciBEb1XZnX4HwearwUy4IMqMowTQkDpI4D0mM/5IV0XaSUnDLqigLSv+LLFKifHnFkZVjwo8UMxgDIfKt9cS0Ze4KN0tue2A0yMlBce4nAppmuZk4Yal6zTmTu9/9D/3hl0cB+D1hBaA9Drz/1IlzP3XHLJZr2fqdOxfqamKi40QgFiBnIM4pUouBMdbbg8dmwQjIBa+4ZDoZXHTIpky69LCY8DVj3DRlDn3gSRTacxION0MfN+68E/OHl/Gqtz0fy3YRTV3DGIvUrAVwyYQ1AAvD+Y9kCFhLqTTK3JTDA1CxuuHCqpOUkQiUGWoK/khzIak4KjX/s8K/KZtumOB941OcY1SYKx/Mx7UDQAhaQXb6KaqWLzQWXtmZN7S5Dc7JSW1dCCkLb4WdaLZjuEhsZCCGCqg7MQ65ldZ4q3A9z9NR1r7RoSnhElx2JolBqJxGbBg9ONQBIYccwJ3vqRT5JZnGRCoZWV8gElBJIwwXxj0JI7MhQu0G9cRz6s5rzlmz++2/+5J3PhGH/3EVgHIckOqXX0Tn/sVts4AxW79zx0Jt+v0ORiCwEUjJNlu5cvqKyCEGtXj8Ek88s4NDgKbTKSVxCkbyi1PUKoAE1fgU7rzmfszNzOEN7z4N1QRjcWkJ1nQCQi/Fqk8MtfbCK7TBI3MQI9uRhvWjKwFDLqDhoVtKIpDYZbR0EvHnYnJPI4wm2LQzxF8kx15QH6MEoFij6QPBwQ8QrNRswfVJ+yfSiDaFRHsKtYgxaHmhqO9ZlEY/rk+T+zTI9x4yWjqs0+qFXVD4Iev5IWDX+AHRxOCY1qevnUZEhkxDhmTb0YOCWk9JYYGeA3lErXWDu6878cVV5xW/uGH3mz/8E+98tKu+J60AtDGBD79k6tw/v/UIWXQ/dsltS7Ud63eSaIY0w4kLvnM6buqmz1gBZ+ucmCRbDGOUtAVpb0o5Gdiz0LJqROBQjU3i4Tvn8O3PXIe3/OwrMX3iJObmFrxTcDRaKEDj8vQIUfGhKz5L8Bj0fx9TKyayNHodOpOicILgO5kEo9qDP2EdaFlO+WQusOPkDegPiXmMp5+GUG1oYqbyykSkXScrsFELSCkPZsw4IE0p5wI+b4fWZ5GNyY7LRYYjZQbpyO9GVAHiVNCscLFCjTTbNq6f1aGCYd1j/gpx6ksxcJaScxIKzx8K2IamceULp+GBGztlYF/4zrEdb/7wT/x9InpEnqDD/4QUgKEi8LLprf/71iOA637skjsXG1uNVaWFbQBWyCmLeCnAsiTiSM6uAbATHlaSqVkq2mSlLF8KSK3x6kETiTRg9MbGsLC/xq4d1+NNZ5+Gk1+4FkcOH/G3gTH5NmMU5pbtXbNBpGb6Hza1jQEKYw667ZJ+ejT2rWlMIV4kdYNyUUz8b3KhjfZmqQZLizlfrrRwlsdWACjoO+JraWkNkmd+tDZXXYqNOZKqhJWMSSruB/8zNhFwRnMZQhoPIYTHItGT0y7fmOzuRJquPDyc28T4ZEXNRrmBopz5l7gMlF2s48UgSSJVlgeTVJKRuaqwrWj+Eb4mhb+gruvBxPMG3dPe19n1rr/30i34h4BsE/NEHf4nrAC0i8Cvv2x661/smQG73scuu325MZ2qSlLixHwzXuxgySeCDTnutpbc0XUoeAD6DUJYqyXfhdBCGQMK4A0ZP8+zoxBHHtaTLOj2Omjma1z8uRuw6e0vwUtfczJmFo7AOYY1prVuRG5zNRik1m/5ocljgxBUXFnpLiOj+KyFaq0UGpFitUhhLBrUE0RgZ7F/pgGMTfMykTyuEUCGCLNtZp4KyxxaHJgWj0F9pHEE4EDDVW1VfAsJUnRi0sJNQMdaXGQlI7eSrfSiTxJdmdN3qTkVBKAyJrkeMbIepWj5E5s9mLeY2CWxmhs460XIr4mdc3AmotaEphk0J55edZ/9NvO5d/3+ab++c8fO6qxdZzHRsKT3aVEAhjqB09du/e83HpGFed567V0LjZkYr1gaf0uQ9Z29rRMQ6IfVcMtbU1Aj4+xkAiU4B5WogI7kqRYeeBusxRqGIQtjgca5QOzwNy0zo6oqCFt872vfx5H9c3jt216MgVnA0vIySCr1kErRgOqdMjNn88n4uoP9mQiH3EATGHAe4U2mSVIqzAzi9+jKYA2E9jitJ/NqMGboCRwMKtx//yLqep2fxSEjEejj+0Azp8EnLkNFu/u21ROcTAuHQWl1kVJvM2sv6feZEZMYENOcQ2grWIKkW7XV8VEJ6JxhE5Wy3qG+oPtGqzFA3KiQWAqMT1KeBIwycCx3K47DmBDQf0JrAcDxdleAZ/H5ivbJzaMGhfw+gS/c0rjeKQvVST/Z/4v3fvQV/5iI5mWbmCf68D/hBUAXgS2ni/2dV9K5n7xt7vX/+qL5t113x1xtpyY7XIvCAcKMb0zgWmeajVBYhSitPgx54MYERxWT/QI0vTKyCSWAYMk3PqwJ8zbRfz0hQa87hluvvA8HHzqEnzz7VVizbhIzh+dgqErsuziTkGgX2iAcUkWADcFBwI1LM25iqXO7BUb2NNS4gowG0Aq68agYsKqLO+6awf5DjA1rDdi5J0YarH0ICrZjsaBJxS8WrqgMTeBY+DzYhINELY/I2MVxzrejEDdOLUl5WtmRghFHsQXTSCAYbhpKSZCswOlLrXp7StdVqbjV85qa1LxfTLCR8Re7WyNwGDTVs2eqV58zdelP/4Mz/rY7FxB5cg7/k8AWyUVg+zmQD2wXe86LJ35h28+u/d6ZL+l23OG5xlfpBsk1I5ZYo5B3lGaY0lqriFIaSrILsyn5QYz66Cmr6owiE5LimXs6kMPY+Dj2/WAOX//kFXj4rnmsmV7nx5ToW6DNrNQNHu2/c75cRo+TmcgIO9EknGSCcaWn/kpb7EReWqFXt1WFfYcIN986i06nGjKzfKyLwNIwBSlxpnidxihPXRTU2wJBj16AOcVPZQqiUIb5cmBLQFJpQkR3NyNowGXy4vD6xaQ53EKiMlXC65LCeC3ItEl73qbuxsgI+/HCQTkrS0PIWKlPMEDDg8acfKB67Ycmr33775z+0598/3b7ZB7+J6UDUEWAw4s/KCLvkbM3/Pm/pcPvu+LWhcaMj1fsnFdhGWRZa/wsNaqP0tSDijdOR5QH80YExSHnZsAk7NEfHyhiTLTroiBS7/d7qBccvn3hNXjNW1+Kl73uFMwvHcFg2ZsuinJnjUQRRAqzSKnATw8LlWBe9CeI2Xj5SfbJU6F4gIfNSimi7clwojwQ/racwBVXHsIbz1yDyhrV1tKjPvrxljcqN0Er3gjZWSmhoJLvUwpx4MYAYJM+LhO/ZugAkhlMS7EnVG4hjAA22IOzN1pCB9msNpmvtKK1c7J1pi6bJLYLmRVw2RkompIq8DZZyFHLKyx+zbC7N+mzLq3BqeXx4MuP8zQJGjS8bm915odO3P323zrjvUS09GS1/T+UAtAqAodE5Jzld0z+/yHuo1fcstCYiX7FMbbEpFIMMhZgF5xQbKKR6gpuDIE5IP6GEgBYxgCF97oKJBSYoETMyxZDcVYnWAOgCeKPysKywTUX3YQD9+/D699xOqpxh/n5JRhbeYomBMoTYijbPmqawKK2FeX2QivpUuMgR/+8YzpPkg/rjkmyGvLWO2dx9XVHsPknJzE/v5AIOo+1B+Dgdw+1ruPkUpQ7FzJG0WVD+0/DoGeMOC/i4WX0WDPqXi+AfEELLB0RLBPer+j667EYGz6LXNJMGhUUsUxhARKTo4IoikR3nChTkEd8dvrlMjvPNTDLzWD9fdUbfmXjTe/8LW/j9WTf/E/qCNAuAttEDBGaD72494/+zpu7/+5NL+9VvLDQWLL+wNvgHaBYKynYMZFdcgQZRzxAuwbHtV9sKa2BKcJJrLd2NhY62ISIYMgzAcVY7xrE3k+u35/EvXv246v/5woc3ruE9dNrYMONRjwc8wQPMfjnuq2Cgw6EUJZOya/Dp+va4JoTNw5CuR0VlqEPzJAZ3iiIwNg1+OJX92P/QUK3V+XDe7zIf2GcEpFIM0QOiq2udoGi8P07zgnKYIIRz3SzIbY6vv/aRYlaPGer2nfWic/kP0/DiTU7dNAoWLJDdfQxPTi25GwkAcvJ3LQFfYgPBPTdalr+emMONgofQbSSDzbiHFKDAz1UkF2GRRimY8B20ByZvrN6099+9k3v+u3XvCPt+X8Ih/+HUgCAaCoCnLNd7N991dqtv/Xmzr9740v6lVscNLbq+lu8IlDlfQOK1VsMUqSWrbeh0m1ISiQ9KsOo5QcXuwzPVqNyRWVMINBkZV9/fByLBwW7dlyLW6/8AabHp9GpKkjDo+glBSYwivgnpKSrUGqvyGaWYRULh0IRUeV0Q0W2Y+vGYwE63T4e2jeJ//PpB0CYQFXZ47YzX/n3BU26cHrdNHKTmdtxUkGwcYBzLOAULwaVpoREJBouRlJ03SylWKgdUzdy4In2ZC3TcaFIQlPPQnxWJF42GJ1sPfRzVCT9ZikypURjxwLbMajtbLNvYnf1lt88Zfc7fuOMtz/RJJ+nfAQYsR1g2ibm776Ktv7lTTMnsVv+1e/dudDYblVx4xLPP0bg6WhtJhQJRMOfMKmSllVnJktp8lhgQoSWZJqniSowHU8ePACrfgXhDq759u3Yt/cwznzH6RifMFhYWExRXMUOOLbGKFl/+iEwJIqoYnIrrw1PRvgU6LjG+PCbEXYGjRN0uuP43jXLWDv9AH71l58D0Dzq2vnO6LFQgqSUYhtlu16wESW3+ToZOlkJQArLbmoBZytRFopcP4haysqIln/F5zBgMy7JsDEi+EU/Z6LHNmiVgcnjDNqfS05u1u5QIozeWIU5OVQ/0rux857fedmN7/zwqzy33wfz/NAO/w+1AKgigLOws/rVM9b+2l/ddOShP+s2W7+9Z6muuv2OE+cPogv6avbMLCmUeFLa/et1VGjkKAA8LB4vSABR2iJIAhNN4pJTCr6IqLBOfgVZjI1P4ge37MfBhy/H63/6DJz83LU4PHcE5LJTawoXTWo8FEq3SJgRyUJejg7WqpCQCmCiEkAPiLPxI0Fal5U3XaSWdvrr8bWd+7G8fC9+/W89F/2xARYXl4PVYvlIF5wEKrfhJgKrQ4aXmcNOYkMXQ0MvPH2v5FFCo/AMHekWlZ0jjTUjm04yem90f6W6wEJAr/arLM6PFSF1R4oCPezKl3CFtEo2wVZOKRhJpSAj+zRQi87FzmFsosKB5fvq/RPf77z5Q8/763d/+DX/gIgObN/+wz/8P/QC0CYL/eoZ0+d+9tZZGMLWb944qG2n6jB7ZyBpvIW3CHvDEB0bJqX9WDx81GoDRTmTS4hQLrzYrAU5DtZc/sP11F2ocOOMtLM49Md6WDzcYOdnrsIr3/gCvOINL8TSYBFLSwOfh6CjsFMqMrViz8IegUof/SGWGhkvthkBbLVZeGk/LiV90kmDTncdvn3ZDPbtvwsf+uBz8MLnTWJ5eQFN3QAjePP67xLxBBtr23oC1UjH1R7lYtfK9shzObLE46jdRivSLN3CK2QDjvoz+UoWDMfAGxwrCZKIlPGtsjNGQP6IWttHKi4l3YgwM8YnO7hv9tb6AXt95wMffsNFv/B7P/l7RHT4h932P6UFoF0EfvFlU+d+9tZZNLXbuuv6xdqOT3SYBmElRKVZqAtEHLIQ16QbI/0mF95DMaETcOnWL4p6QIFN6kTzw+kXEzY8tDEYwqXEGBZGp6pAUuGGS+7C/vtn8KZ3noG1U1M4PHckCXgk5UvZcuJ0koMmwgNt4P0G2g74EbgTQ2mciTuMeANyyFg06taLHY5eL46Nr8ctty7jgn93P96xeQJv/skNOHFDD3AN6qb2wS2wwXhC/HusfPthOK3sTGDsSTCtIGF/+znJGvfgQyBcBp+aaGOe5m5zdNyBWkWAsw9/paPdC8uxEOfOOWla2INzJnA6jPDIUcNz+Y3CmowXd1GVXYFYhkYN3yFJESFEMTuRBBNru7h93/Xuwc6Nnd/4Z+/Y/Z5fe92Wp/rwP2UFoF0EPvCyqXM/ddMsZMBbL75prrG9scrV4aEKqzqNYoui/IrujSn7DrJA6egyIk8t/4Y8U2ZVm2a4SXLklaz5ZgYTMD4xiYfvOYyvbr8cr33zaXjey5+F+aVZLC0to7I2wYMc6TsU0pFZCr4/h1ZYdNQPWuo7jHrhqq7FvTpzYT4SfzjXoNvtYWFhDT7z2cPY+e2D2PTaSbz21evw3OeOYWqKYKtMV/Vnzd92hiqMjRuINLDoJvdbKOtKip8Hcza4lKjCNoXJRg6IHV7jDWn3V/h+qRARt+fvMOYVEfErOyhhRF8T+Rae2WgCOYtwrL7Bd6ihiDDDWIPONON79+5sHuDrq98/7+d3vefXXvdBIpp5qg//U1oA2kXgQ2dMnbvjuhnipvnYpTfOD+zYWJeXBq1PPYNQiTMu2QCrJKly9l9LtlRhrUQWKfeH/ERqUyEIgSfIIwATwbDJPvOhf2Vu0Ov30Cw4fOcbN2HvvQfx6je/CJMTFeYX5lFJRDNV9kFwR6LQacR9uh2aGNWEwyXTtNh5q5uoMApN+X2ZXNU0DUCC/sQ6zBwe4KtfXsCuix7AiSdaPP95Yzjx2RVOPGkcE2MddLoECkNyp1qEcA8UFJzJtIUBglPy6Lhvi32VFOtRozMMJK4wMbRalBZnAiMKIaUoKWWsmizIYgigFleVXBImtAQ5rT5A4wlAGgVkxe2CUqWSATOj06tgxgfYdcuXGnfSvuq3Pvqzl//SP3jb5uZ3+AlX9T0jC0C7CJzzmrVbv3DzXKdrF/7Bt648Utvx8Y5rnLo1TE4QCr54CQ8wBDivMSACbEiKjAalwpL8+Mjzf0O6UeYKmCB+YVJZeCBwTPpLN0pWdAkD1lpYU+Hemx/Cw/cfwKafOg2nvMB7Dza1oCJ/K3NMHjJ+P5yooerip/YsLnn9ZtSKKe9xg38wezMRtPwBEBxnOSTdEgmYG3S6Bv1qLbgBHnxkgPsfcCBZgrWLMMQ+uCl0OwSCpT66VQc1c/LZtwoNN4HEJMXCMLTGOikpMRbNyJs9HXQZ8XNEuQuTsl1I1GIV10Y6W8Vx2sVHpiApcFlUrHoWFJlwo7vUvYkyF01rZWYYZEGaY4f+WBeDzgwu3fOlZmF6b/Vbf//nv3rOR976t7f8DtP27U+Pw/+0KADtIvDzL58415Kxg/m537/0xrnGjk9XztWlM2+bLRapn16k7W+c4ErDTlJBYPFMQApzoAvxUSISdNjlHjhSaI0x2cfQ+YctmUJHUEkIvV4f9bzDJV+8Di995XPxyje+GP2JBgtzs8FFtsqSWdW9ECmJcAAr9eooSoiLxGRpXUB6OSI0FF0NUHb3jRoGdiAy6PW6kH4wsHQ5ZtyEgiEhzQeUFX8UqLICkzMLWHBU5YE6t0SkthehgI/y0KfjW1mSKQsfHaVTJyn3/iuuC49Fk2nZtvnPzGFyuosDgwew6/qv1Oue7zrn/Mrb/vyX/6+f+s0P/j1AROipQPuPCrg+XX74N8d/RF++efGCP/6bg1svunK+riamOq6ps/2/SACtRIWShhvdSUrxMeFgkhNYITC71Bb7fD1OXw9hpYYYaBF9DDl4tYp26/W3ojinGUBBouzBpsXlJazdOIbXvuUlOPnUdVhYmMWgZhip0o6fQuiE4XyDxSy8VChYYMPqST9oJAZwrMA1pDQlDl0PsTLKDGc3JjKD4bP4hJQXXygAnJOITZFwFIRLjNI7x0WjjNz2G6biMBpRnnmtdj9qIjJTMnsLmNYhjNoJkpIBSDlMLDERibJfgpcOqz8TDV9VpxRHSUsEEhvqHQeHqYAlSU4MjhiRCcXLGmBsCrj30B5ce/cVbs3zYX/mV87877/5z8/+KNFZyyK7+IfF8DveH+bp9GJ8JwBs2ybVz542du7W96/9+Nlvnug0i/M1mdagRq0luW7L9EOmON1CylaXIpJuygG0dfOIkvRRbAlNsBqDCVePyf58wTpmfGwMCzMDXPyla3HVrltgeAzj4xNg5nDz+qtw2PQKBeus8M5XW4y0blRpvMlyX38fRm0WSFmuxTWW6rWZfVqy/3qS/jnqrSGt630ES0C39bqI8bCCo7WDL33zSdq8gtG3MrWNP2TkznDk/ZcYi9pBSL2U9MiFDRWD0DSMTq8DmljCdQ9ciivu2tWsf2Fl3/PB1/zXv/uH7/1tIpoX2eWebof/aTMCjBoHdmFn9e6XTZz7zdsWYeXw1i9dPlt3xsc6HDCBMhkoAzcJACLKNhOBA270UEgE523kQUz54AQGGJkoIpJCAuofhAAikmrLW1xUZvaiIhnHbTfcjwfu3YfXvPHlOPX5J2JxeR6D5QEsVcPnCMpRN06nbWfadPCpSBiKqLpRB0+USi7yBKIKLlswSvn741uhjVmlDKqkFTD0Fc/lqH4z6v9bMt5jSZhp5BfOtlrSkhXHi6AYEVYoPQZeschg1X2oDU7KsQrmtiSYXNPF4cFeXHfHZXjoyH31qS9f1/nA333jrl/5h5t/75f4HHvattPk6Xj4n5YFoI0JvPMlY+fuunMZhrH1by491HSmxivXeKQ1mjwguLdGlN4/QJ69R9aDY5b8nlqMCdluBJgqz59O2Y7rhiJmu2sZcIiHjqGSKZdQaXMpRHmzCHpjY1iar3HpN67B8170HLzq9S/B9PQ05ufmvDpRsuuNf23+Js7hJFSYpyI464i+ZUW8UCn+FpakLkyGSwHQM8F2HJyV8qSuaiNRTi1BcuF1/i5Ec8eNiudbAFUsIoQinDQZhkqwgaWYu6gPMPsCrBzzSAVltdd8OjS5cG8NpA5T/H6T1pNgp2S5fp0nnD+r9npQiMBGYJ3qHLVqEwxrCZPTFe45dCN233sFFni22fDcXmfL775l9y/+9k/+8of+3h+a7dvOEzr/6Xn4n7YFQBeBLaeLPeuFdO5ldy6jC9n6mUv2N9XkdDWE4cQVUOALaC43jbqfiEpvfi1MifRgo2woQoIwJYGH9ZnxynFYt+vxFvOTS4NuZdHtTOIHtz+Cffcdwstf+wK8+PRTADfA3Nyir0PeLSW5DEGkWP3HNCERl0cBQCUGkwoxpsBdULbmFNJZIXCiAkrQilpT70POZFA2yeo6jwGeo0IwZQRIxjLkBn9sbAhae6A8Hob6Aga7ODpJ2hIGpy0ACFhQtiiVo0SIU0uXELkh7Bi9sS6kv4jv3bML9x68GY2pm7Un96pf/+i7dv/ib/+kD+3YJubpfPif1gVAFQE+a9vO6i0v7J17xZ01mJutn7vkUFNNr/VGoxF4CjHQqYWX7CufW9+YN08Q4tKgIcaVxXVhPMzhIWKDpH2XlD6jk12US2BkHyZoPmlRMdEfBzvGdd+5BQ/c9SBetenF2HDSNJYWl1APBoE9FtKDkxNRrmA+LiyaaBLYZav0SMH3vbtB9klEEcnFiF54uUDk1j/arMdQ1iBllZLRp4sGS8u2OxqE6qIYQTdZqaFX+Ay3o9ZRaCR0S5ALMKd0qPYalALWI9AR3GGFKqY1/WvMIci1UQHikq/g5JoeDg/24nt7dmJm+RFQJc30Blv9xj969+5zfvct73iqhD0/cgVAdwLnYWf1phd2zr3stmUMatn6pctn62pqquPEHXVOpLCz10GN0WYmOb4Et6C4QRCNkA51EnlFaJK9U47tNoH1JsK5dVSyN+EGZAiTY2M4/MgRXPy1q/DcFz4Hp53xAkxPTmJ+fhGNa8J4QiXRhMoZPYJVRvl+cpQ5GQbYlgdGCBJorka5FpviAJhAi2oRX1iKxBy9tgTnWftoczvJUGTkMTuAvJ9fAWbICSwtcDh3KpEr0o5UM8hFPNMKRuiKBXCNQ7/fQWdCcOe+a3HjvZcDPYdOr2om15nqw//wnbvP+d23vPOZdPifEQWgBAalestL6Nxdty4DQlu/9J2Z2k5PdpyTRAEetUP2K8Oo+/atMEH5tIfxIfqHiMvqrxxiScnO3BQHkfLvSWBbFpBkuX7enUd6cq9jIVTh3tsexIP37cOLX34qXviSU9Af62Fhbt4DiWTVAj3kniluFBtSM7X2RTAKgVeeN+wy5SXoCFwtavdNQ1yLxFGU9ntLKUjleCzHh5iMHDqMIjtcVDq3pBBPoTYjcJgHUDiSaI9/kiLhV6cYayq5UF6/UmCM+tVxjak1PSybGVxxx6W49+CtmFjTQ12bZt0Jneo3/q+37/7A773lncrM4xlx+J8xBaANDJ71st65l91Ww8Jt/cLFB2u7Zm0nzb3RtCLIPcFe2AJ2vmW0ALHNqq50G1D2pIu+BC4vzxNNN3kWRqiKW60ttw5KFMDEDD0BGRdC0Px/93oVuGHcdO0duPeuvTjtFf9fe98eZVdZ5fnb+zvn3nqlUnmRhJcICBgShBaQAGKCQaHthgCpAm3RbkfBEHFQ257u6ZkJtGt6LduZcXrZgtFu34hWFJJuNbbolOMLm860SqEICkFi3qkk9byvc749f3zPc6sCLhuRJOe4XCKVVN3c3G9/e//273EaTjhpAfKsiWa9CS1UiNAmH8JhLa/BUBKnLSnjte+tx8yv12QMV1jMDa/smtv4E0QxYx5XtdsQR53VbWpAHay03QafZfqBLxQoF3eZk3WCC1Hybt8eTDmlAFKyZXnOGK4iMzj+S3jHLAZqAUAUZRVuNCT4Va+IcXWuKEbSS/jl2DCGt38fkzKKnr5Z0E2dzVuYJG+8/fIfX3fz8tc+l3FdZQH4NYrApWek733oFzWqKHrPxm8eaKWz+xIzphY8pouNKRV53YWZ00lNY8lpdPPruM2EAezYavh1QUtPQR9Ogbfuk40oumFcpp0YEklXtYrGWB3/+p2H8dRxc3HWklOwcPFcNFpN1Bt1++H1URaBLxCZjDrTNAa3UVejRwfMQjzfSFtyC1umownN1M7rL6pquT09yeGWgTJzIvBhoT+aQZ0V/TOTtK1Zo+5O5Nm/P+KVZzHr0RUcRyBicsVTo6srQYNG8NCT38UvD/4EHV0JZnXOQm2ikS1Y1JW8ds15H7v+5uXvJqIJ6+F3RB3+I64AxEWAfjqoLjy980+/91h9N3L5wMZvH5DKrD7JRUjEMfIAb7rlEO1I2df+mWNnI+4v+sgLP7KnJX9zGlmscS/KvcNs7AUYgit1ZD1JKNhDRhuMVCmkKsGhvaN4cO8PsfD4+Thr6enom9uLZm0KzUYWUZXdj9F2Q0E+qFDDeuh5h1praa1DsErsrRKoxCZgUBPbnbhMm8Md64/sayahQqmltpa+SJ5xv1ZH4aauO4rt3zWixWuw5nZx5/HPmqnIFZJpzTdm4gCgIrxXTu0JAXRLo1pNoTo1njrwMIa3P4gJOYCu3g4knKA+Xs/mLqomN73rVY9ce8sltxFR4/ky8CwLQNt2YGBgUF1yZsf/eOjJelejVbvzy1vHRXV2i4gm8oh/tNcVHeifrjUXROGh8O2fIxqZVjsPeW9k3HEyyYNynxWUm/dZ4s+t5ZMbTp1xgTa6AhdHFeCCYneSpikAYPev9mHfzv046UWLcfqZJ6N7VheajSbyZl4QuxSCrizxyafo+JbeagjiG9VGUKs2MY1DzGPQkeM2Ge2pRgW8slBg23Q7Vq+vbBovMFNQbOysRBI5ErVd8IU13YxIQ/xTuYghWPshI6bSYGZ09lQx1tqF4Z9/D7vGn4DqFPSkXWAo1Kfq2dxFnclb3vua4WveeuGriag5ODiojsSb/4guAK4IAMjXDw0lF57a8Vf/8lStSvzLv9j84Hhe6e5RECFxA0GbQ69bl3mRbjvRg0yby5HVGCEYUfgRFbnh3Bc2BGijyrHfQIT1oI3SEu3ZLaFVt/6AVlLbUa0AObB92y7s2rEHJ554HF78kpMxa1Y3mq0Gmo2WfT8YIahQotSRZ2B7u8PtALTCik+j3YO4OGZbt6a21V4hNHT6Yj1iMfIzwIQU1qwUKLcc+RA8O3+gjZosZG3Xig5BhpKco7urCiRNPLF3Kx7Z/hCaahSdPRVoMdyBeq2e9S2oJm++fZU7/PuO1Lb/mYelI+wREdq4EenAADUf2tb8zx+8Z+d/v/eB/c2OvrmVXGdW0CGeAsu51fCK9s4uTozC9kOntfbaeo5CJKB1oOfqzKwCnU9AlFJrcvOcZwD5XD2K3HC0zszH2jkBxQw5ye3tGO43IrOmbDbqSKspTj71BLz4tJMwq6cLzWYT9XrdMN+cgMeLnGxp0mLEUvZYk2Xy5XmQ7ZJGiOcGoMTgFmzFR0YrRDb4AtH7VzxsCkYiG6jEIdSDhGwh5eiG19PMNszakH3XpaL3hyWs9EKCcCiepq03oRs+XyhjPy6wDmaglVSho5Owb3I7Hv3VVoxMbkdSBZgVtGgoVqjXG9msuZS8+T2vHb7htkudb7860g//Ed0BtHUCzfXrJVl+Gv31dx+r1xNO/udnvr63VemdlYA0xSw1bxMf5wlIW/tofftIx26+2rLYiinAbj3FbXbR5sNr7y7XRTg3IGXmUSPKk2nSZhYOeGVkacYMdHZ2ItcaTzz+NHY+vROLTzgep572IvTO6kWj0USr2YTO3apRF6i50n7X6mJQxYxWem2/xyf8Rl6FMc+fLYVYLFzvL32foBEpFQ4HFtqcAUXsR4/YZTmkRrW7AgehE0decmy5/64IizCUEnR0J2jmhzC848d4av/PIGmGrp6qEWsJUEmqqNVqWXefJG94x+XDN9x26RUe7T8KDv9R0QHEnQDRRiYayB/6efZn7//00++/7zsHpdI9Czo3w4DpAMyAzjCqMXHqMR1EL67F9IfdgVnaZv55boFtR7VLtI8/2OS1AI7BpmBlxm7tlmtoySBkefyRQ60TB2uRKIvOeVhbRaLWaDYzpJUUxy9eiJNPORGz+3qQtVpoNGqQPANTav8c4c/GbhOQxy67FDQE9rCwGKt01wFAgq22aDFyZCpiB04q7OFOJ/V1DjzauQErz7aDVwCE909EoCixbL1AdRYACYXXWBT4BDzFAZVsOxetc5AGEkXoqKTIMY7tI4/jyT0/QU3G0NGRmsJlqeQJK9RqzaxnLpL+Wy8dfuOfrljlDv/AwNFx+I+KDqAdGKSBQXXB6cnf/ODx2qxWs/aXX/7BpK52zaJcMpeqbT/Aui3lNgiJaNq9F93yThVmlUBUQJsjxh8M/4BJDF3XJusYir+5vxQxjBFn7mMyY+Z72FxKBGKF/AA4M49c4+lf7sDOX+3BvPl9OOnkEzBn7myoCpDXG8izzG4s4jUd+3Hcx3RRG3g3Qy4Be0/DEIgR3+iRjqiQ40hSxF+e7Q6iCHikSIfgmNXttGCKRiYzFiTeLUlrDaYEHR0MTprYO/oYntz1CMZaI1AVoJs7rdJPm8xCZjTqWdY7r5IMrFv+yOvf/cqj8vAfVQUgBgZv3rA1veiMzv/60M8b6One9V82/p8DOunsFuthGzn5WGzAtpoxFwAzoNrOv95xCOkZ1t1EBKMotoYRdk5m5lB8HKJul9ROgOOadi15tOoLktfgh2dmdiJBtaIgAozs2Y8De/ejp7cbJ51wAhYsmIOunl40GzXjCYg4ftwx/EI7rYOtkP3ZRZcfn3qDNlqwtPn4zZDdIs8I+8XfqwjSeZ2GFBYlAA5PQSQKm5jOjhSc5tg3ug1P734UY/URIDUW73D5CU4fQoJGvZn1zk2TG99x8fCa2y9+NRHtOxoP/1E1ArQfvo98RKe33EKtB38uV3/wE49u/PIPailXZgmTsM41dO6CIu0NpnXIq3duNH5mDNxy4/9mcu+gQyhGzJIjIHDtxQ/BtsXX/tZUWiC5sSODzsDWgdd88HOPxrtdP2nLdIxpsVoH9x3Augdp5FkOrTU6OqtYvPg4LF60EF093YAIMocT5GK3GJHDrlDYmmgJ7jfRn8tl9DmLb0+fdSpKESRC/nXFclqVBwWiFtf5KP8+GsDU0HDj9HgHZDoWX4jkC+Er7D3icyhW6EyrUKqFg5M78fTuxzEysQuUCCpJArdb0HYbQyJgxajXm1nnXElef9tlw2tue4VB+49Aht8xXQDc0z8oauMA5Q/+pH7Vhwef2nz/d6YSVekW0sImuFLb0FFr+aWtm7A7pEIRF93ajUluUne81bVF7bU78PCGmUVMwBF1gmeAyoJttUhmvepVBP65BRigJTMzuejCrUe5FOi37NmHFqBrCVpZE0maYP6CeVi4cAHmzO5FNa0gyzSyZubnXuOPGBKAfLS3ey9sd+SAy7jw+OAO21nNXAAIKo9NiAKt2S5Mzd+JsO/KituRkKrMFFlzITgkKWZUKgxChtGpvdh74CkcGN8JUIYkSe3rzK0eBN4bULFCrVbPuhYgGXjHJQ/3v+PCK44mtP+YGAHan40DlG/YsDVdfnbHln97orEW+sm77v/2WKqqPTpsy8xqT0UdpaDYcno7cHHefYero9Myywq/hgK05y9bbx6CyNwiyhFzUlaGuRHbZfdEEcXYqd48scaYllSrKUQE+3bvxb7de9Dd1YUF8+Zhwfzj0N3dA5UyJNfI8pYByxB7JcTW2ohCP8Ns7wiIXMh0LaYFOz2CKwgk5NWUz6QkimPdCviAdn9XxrM5UYxKhZDndYyM7sLukacw2TgAKA2VKjBVbAHRniJgugmGUgpT9UbWsyBNBm57xfD16y644kgU9pQdwGGemzdsTT96y/mth35Wv+pvP/3zzZu+N5WQ6hYiZm1bbeUReB1QZe1otPZmi1SC/pq2rTLnEbMuktvaI25R/mDhRXYf7ccL5KDMfSwdXK5NgCUp26Hkwf3LbgdEW/diSxwCGU6/WDCMozBRB9ppa+3NpNDd3YM5fX2YP2cOerp7QGycf3QrB/Lc7u+DlkHBeSEGrkNhdWi7CZVT0WVHC0hzJO6ZyY5LLO+AfWaAiroRV4jYjgBKJagkgKCBemMch0b34OChnWi0xoGKGGt4u0lhYrBo08lFnUZCCvV6PetaxEn/O14xfO3alx+1gN8xWwAAYMMGMZjATyZ+/8P3btv0j9+up1TtFi3apHO6cE5HUNNmPchubpfpicX+tss1lPtaVADiwDFTANjP9W4NCUtFJslBufY3rdY22gyOuS7RgaNI/mzciQnay5ZIDL+AfaoRCjRbJrZrPw3ROSQHUmZ0dXWjt7cPfbP70NXZhUqaeudirQVZlsFN5kRmixG/HxxRh9sLgAkyffYCEJx9rYRKDE7CAJgJigkpMxIW1OrjODS2F2NjezDVOIQcGRSTyTGkYsF0qKBXYViST63WyLoXUHLD7a945Jq3/96rfxcR3WUBeJ6eoSFJVq6k7Mfb6ld94GO/uP+LD4ypjtmzWUQzCs5B8IAdaR0QcU3eE09r8bZUrGGkw2QDIhxnwH/ocigdyWncfO2ELrbLcL6EWruVVJQ2iwgdb4u0NriECbrQkXc/RKAQ59iZa5SFDFtPcgiZTgDa7MolFzAn6KhU0dPZhd6eHvR096CjoxNKpTacRIcmRTSgnRVbvPu3h5jJM54ZkZrRbUK8nXrYESoQFCkwG3iQKUeiFBQEjdYEpibHMTq2H+MT+9Fs1ZEoASeR6Sk7XMBHw4bZ34G2rNCYambVeTp5/e3Lh6+59bxVx9rhP+oxgPZn5UrKNmyQ9GUvpi0PP5ndNLd35+DHvrRNqr1zvdzPJ8JalZ+5PJxtUCTvZcy4A/T2ZBzCO9xn2xONiPyNHfMOyJtj5P6fJboZDy9zcbFlRQcimWYmGqMRcT3SIDEHzsr/kbXqGGtMYuzQCBQz0iRBR7UTnV1d6O7qRldXF6qqYoJSFUM0e6NQ6NyChubm1RQ2Li4jUdg6ExF7QE+5A29HJK0ztLIaao0J1GrjqE2Nodmags5zo8MgjY5qApE8SkxywSWRmjN63xkCxQqTU/WseyElN7zzouFr3n6eX/UdS4f/mCsAAHDLLdRav34oOfe0ZOPPtstbxsYPfuS+/ztRUWmPJiYOwZZRyCXxtNBOjvwAXTS1eLUcbGSURBGykSsPgMhSOAL0zHlVDK/aC0QYKSQCUaF1dtaDNnKscODjzLGQbShCgXwUy/WsOo7YIONkwy7zvIaJsSmMj+436kGVoCOpIK2kqFY6kKoEHdUOJKoKJkaVE6RJxdziZNpy5R0JzbYhz3PoPEMjy6CzFiTPkGV1ZM0G8lYdzVYdWdbwrEuChlLW6bmgXuQCX4F8Vl8ALd07rVSCxmQj612YJgPvvnD4D//Dy444G6+yAPw7nzvvXJn19w+qM06kT2wbkd2Vjkc3ff6rB9KkY45WRGyUcJbTqtmviig6pRQnzlJ08yO413pX0Vhma9OD/LYBkZOPs7COsHSyYZMiUVyZ9/anIglHxOzCY9pdQOe8WYjPGozCLwSxRbaD/VPjc2BLnkoIpHMPlmZZDVlWw9TkqHfQJc0WcAMSTqCUAiu3yQjvjcAyJPMMeZ7b164B65bEIChldvsqUbZQKQjlxqYL2nQOxD6GG8hsAYgozz57UUBKYWKins1eqJI/es/yh698y9nHxKqvLAAzrQg3DuSDg49UXjyPtmwbkdVMj2669yujKVVnaWP7x54EJN7rr00DHzsHeZkvIslrRCCytF/AJRC3ZxxGV5ZWUcsvofAAIGZritHurc3WeUC3CXKjYuVj1UPBKIwe00aaonxIokASnzUAAilrI4ZAdjJciRyZRpFYRUX5DpM56KIFiuzoZV2AglxawhAkBGblwdVoJWPzFyWScLPtvHIoxajVGq3Zi1R6w7suGvaHf/DYPfzHdAEAgIGBpc0NG7amL55HW7aPyGrWP9l871cOJEnXXE1eRmfJMXCmou6D3uaXD3jFuovy1gRPVvEMfzG3Fqjlk4EDsm8M78Uu18187G7WaP2ooiQ8H9wJ5GK8y523n4ZREMIDiaZIuRUiSBcMPdn6KCqoKAjUFTL2RdD791lVUZyOVGg+qPjvnVW6luDDwDoEpOpIfm3CWhk5mULixEqsHAEr9k2AKRht+AdZy2JShEajlc07sZJet/a8R1731qWrjlQPv7IAPOeYwPmtDRu2pifNoy3b9rZWJ8njm+/58kHFlT7NENaRt18M6gWi0HTBjFfN2XCPeFxge0u5biGUEApiJLZqOwrBIEThphZtpIUOOXdNtULEI7CAlwvIJJ+bKFFBYPuP2jsam/ZdeUGTSJv+wHUQhbBCRIIov0UNuQqAFwkVYr/blUJo+57epQl2VGIbihKKBjCDK5HNYReIIfk0annXfJ287q1nf/d1b1t2/bGI9pcF4FmKwNCQJC8+jr66fVSuJjy2+dObRpKkY7ZmAmsd+ec5pd8MSjn3cczt4XOyX8j06CkX0+0adz8FCM1wnuzNbd1+fJSVPZwmnTaAiG7jYDwO2a8Q3YEPcmMHkEVGqI7HR+3BJNa4hDSgIvmwRfDDZF/MYQrvmSVbxWeedLGHskVCCaLxyTsVgiganeiwKxGb3aXBSYKpyals1vGUrFl3/vAf3Hzua4iodjQ4+ZQF4Dl+zIpwa3rSbNqyZ0KuzuXRzZ++b09a6ZyvTacaHDTJIuXhw28FPf5zHAhBRG3ehIgARW97XXT+cTmH8D/D/TcvhpQyQWkX3GlWcMQh+wDWeNOMMXkMT8KEC5DdBiRmdrZOSdr56YNtsbEHz1qHsS+AgaMQgkbMT1BAtCeB3zCEP5MEm/VnDAKV6LfZ947VjP5/bu1o1rSM2lQt613MyQ23n//wlX+ydBUR1Y71mb8sAM82Dmzdmi7soa/tnJDVBNn06Y1706RrnmZlQqfE79tDJBV5CW+U5AuTR0cS+fXBuwuag+MyQbXyNmVMBFIKeVtugccJKLAJocXn5Dkpr+MWmI6fowGfrbgJbX8GFEBNtyYEBe6++5pEUV8eoiSXImz1EygmFxWCPjjKIuRfL02EeTomSfGfc4YOQLHC5ORkNvsETt7w3ouGV9105hVHi4dfWQB+20XgfIMJHN9DW0bqsrpD8aaPD+5MpXOuZiamSC/vNvu5A7QQteZiRSc2gsxLainkfIWdtVkeOh9Asf28p8SScRsm7daHoVvgqH+QkIweiDgx+g9js+Xae7/2yyMnX3K26W6tRhF+EXQRiAI74nZcxYIgN45IlIMg8PwDBrfN/4ZVSRRAvoLQisgnB7vNCEP5kUxEkChCbWIq611MyY3vuWh41U1nrjrWV31lAfgNgcF5HbRltC6rKwk2fezzO1JU5mjFinUE9oEi0wpp47/HrX08zlvGofcN4DBHe7WhvelZKGrryY+4REEU463M7fdnn1Lscu9soeDcGJbaoEyncmQtLrrEou/WN7/NQVcB0/L1CsjnDP/aBYq6ZkSi8UfTDN7ANA3Si0ucdWp2yiDyUWyiNRKlMDU1mfUej+QN77lweNWbz1x1tHn4lQXgeS4Csztoy8iEXJsqtekj92xPJJmrmcEu044QWc/EfSpTYX5tPyo+4opopk9/dMvNcKJghUUhB8nO4cXIbvdyiGJXI/vyELHmFEVdDYWfM20CiYhFMzAS3UgQo/weINTFIBIXPcqHm//ZAaVtzv8U3h1yWIZYtL82lfcdnyb9//G84VV/vOSKEu1/9ofKt+CZnyGRZCVRNtGSqz7wd49t/ug9O1JO52swsbYfZZHQ/nrXH8p9d8tCfp5WAmjk1k/Q7t7z6NSIERVBR2CiNbnM8/gmjLE1oy5UOrj1TAPeACjRNvdXmyxB18ozoMRam/s5ILeGCZGYyI4BiW/p/d4jPrcG1nAGHRGYT1HRcKtVJmMjLoXXCShlABSGExoJlLVZI6GC+w8TYapWz2YtzpKb3rv8h5e/8Ywry8NfdgDPzXaAzHagJ6Utoy25hllt/ui9OxNNc7XzyPSU3PYuYBp4FUAxtx1wuIC4yHK7qoMV0AhMarFZ/2mIODmuFLz5iIzyjrSe4fUgAiDtLa9NUCgBUULPDDcE0Qxbi3YK1LMhedHvinwJKBoB4tfMTEGHgQBmOjNU8dkCht47NVVrLT61K71m3ZIfXP7GM1ZEcV3l4S8LwHMzDmzdujWdndKWlsjVaSXZ/NF7dlUauluTUuystLVvv606LwLjHOtPW+mrV8aL/X2CEJXtx16jZ9eevqu9IEiL6dzhQzYiAo20b+MD8ObGEoo0Dex798RuNXTYWcSGIn7McCs/BO59W71hl6Hgfo4fLeKv24JEYq3MQtYZtbsyCflxwXUezITGVCPvO1Gl19xy9jdX3XTGH9nDX878ZQF4bp/zzzdkoZToazWRP8wy+eTHP79rUUN6c8VQOrf5g17zHvT74vns4kMo4zixeFsu5Kd4w/n3s3dwIw6rRglsu2hNd3irYh39bIoGiVB0yMekOwMTfVi8T1DccEhhhJeCZ38MS8QjBRH5P8dMMylJ2w+0mEvCjMnJeqt3sU5vvP38H6x60xmrj+SU3rIAHAnjwErKtm6VtJPo6y2Rm6pp8tWPfHZnpZV3acXMWmsreJEAEFrDC5fV49mBbj6Pca3Ygtu1wwRfQOx0bRRwMVUwTjlyST0FoU9b6EnU1ouPEA8iIdbRbd2uOZJngpGkDbBEoYVvBziFzHjDzzhLSCAE2v+vOMHExETrhDN60qvXnvmjVW8qtP26/KSWBeC32AlQa+tWSVOib0605NpGrfXJj39hx/wMs3VCioVyu6piiGQFOitZeXGIJhOINc8IjD+B+R6uI7C3I4cYL3KFxuYMuIPBFN3ErsjkceCFMlJaH4kNL/9V1mnX8AfZi3L8oYtAP4psQZUtSl4vAGv3RSaCi6MOwXUhFBUh082wd/GhKFSVrMtQHCKsmDE5MZEvPLWSXr12yZeueNPp64ioWR7+sgA8z0Vga9qT0ldFZHWlwt/60Ce2JTqZo1NijpN+BRqHZ7xZlxq2HHrNALLCrd4uOIpvauKIzydtY0DhgqbAV5AQWuLYOgwBK7fOdIk/OigObcsuHLYV4SsSRErON4EockaZzmYUz+2n4lhANCOI6rqUCqWYmKhl817EyQ3vPu87K15/2o1ElJWHvywAvzNMgIi+Xxe5hiCb/+7jTyW5mq0TZvYKPM9cs5p5z7Qjn8QrQoWumZxkdybE3ScMB2PNODTbYQO+faY2ARKz5+9yZIEOaQMMSRlrLsdlICuBJhTESIaByBYk5FC0HPinGTPEktqxwKavavIqyABgSigkoqGYMT45kc8/JUne/BcXDV/Wf8priUiXh78sAL9TTGBoSJIOoi0icg0RNn1ww2MpVxdolSRszktUAICIxRYfCC6iZRZKJOet3zaHu3jvduMPmQYAxrLZ6HdLAXaMe4sIjGML5Gl7yMmz+tjRcn0yD0XaASnoHukwGQnt8uHiV6lw+yulMD4+2eo7Cekf//kF/3pZ/ymvi1R95eEvC8DvGhjcmhLRlpbIaoA2fXjDE5VM5uokUSwFU87QTh/WhAdsW2MqZAaGW5ytOYgUvP5hv69uO9ieK0CHh+6CyMgUn2JAaptSMTqiFK3sCqOKpzsKaIYwUCpIpAMvYqaHiTE2NtE647x56VVvOe0rlw2cci0RtWR9efjLAvACGge2bt2apqYTWJ2Q3Pe3H/1FJZd5Ok0ThrfpEht6AZuGx9ZmO8y+xlFImXtXivRdl4nHVvmn2WUbRkAAu/FC/GFjCvFZ7AcTl8tjWnRN1ircVhISQg4NJmV8/iSLh3UgWl6ycwP2Y39e6CzAHGUgGOIRW/6CiAZziAFzRkyAyeqbHJ/ITl7SlQ6882Wbll+7aC0Rtco9f1kAXqCYwFBCpgi8EpJ95+5P7ehotObk1YpSIrnP1yMxSbSOLsCaPHJP0fqOOGLKi7bovbmlNWSaTMCRZRxAKLFSj4Jygdnakuv4pi227/4/DjQEW+OQyG13Bsaht/2GeNNjbyUmzqnIDUbaDwlxCCgAsFKYGBtvnXleX/q6Pznr7ouvX3yrwUuEysP/3D1cvgXP5TiwMrPA4NZ3rX3J79/2lhftmN05pbJmliek7OE096+IsgCcoeIxC9jm3BmHXPGjAHnQLEoBdkrA2OrLewCEUYKj30ZsVYFOfmw7EHcDWztODzIq6wNoZM3KUHqJfWoPRaMAOX4DucW+LXZEUERIyLgFKxskqD2O4BR9Cq58KVaYGJ9onXVBX7r65rPuvuKPT711/WXrEzvzS/lJKzuAFzQmYOWnQyKyIsta9w/eN7Z05ABnlc40kVxHYFyE5jlue3yvR/BBjBk4v9uwI5fgjyfxhF/0IKB2azKOwkXDZB/ZhSkHAXqTT9IUKfUpZP0UrnC73iTyXYbI9Dk/mJRykDKzwvj4eGvJhfPSa9521t2XDZx865q8X93xrTvy8vCXBeCIeAYGKLedwC9E5PLujqe++ZnPjyzbMyJZtZom2kaJE7u5nwo2uoV9uCbvputzAmxLreOoMXJzPyKSTZEaHDvpUET7DaxFtBUfscIcZ1IarQApuJVKFNuNth1/4Cy0g5DaR6a7kCVmhYnx8WzJhXPT1Tefdddl/SetW3PdoBoc7Nfl4S8LwBHXCdgisE9EXk1M37zncweW7d6fZWknJdDBWYciWa7ASIGdu5AQW1/9wNZzd7iyJ9XadUIxGYtwj9WFX+s8+x2/wPkahu2ChOgzMSYbpFvwYSXE/oI3+YNUhC5tN+NczmJ7cPL0aBWqi8R6CQ2lUkyMj7WWXTIvvfbms+9avnrxuuvLw18WgCO9CFjEep+IrGKhBz71+f3n7D1IWWdFJXmug/7eHiRNQOLovxKZbkYuP2ZiCOaiyt74hkczHZgzYSKROs9978Ms4v1czigEmMRrv9jX2+kP2L+GdoCSPI4JT1myIww0SCUYHx9vvXzForT/1pfede6Vx6175aXrk8HB/rLtLwvAkf0QUW4xgb0icoWm/Bufu/fQsn0HdJZWOTEx4GFiL/jtuXEaEc5nfQKco0481xumIUOLESf7lZ/oGfT70+2746+YAsSRb0HRr58jE8LYPtzxnPxIUFAB6ii3z3Y1CePQ6GjrgssXpjfefs5dy1bMcTd/efjLAnD0YAJREVhVTZ7+xqfv2b9s7748TzurSiSzCVgh4pKkWADCbB6z+6iQTyA+uEOsAahJ3hE7a5tW3K4TbV4AUQQASviephhxMPWPEoB8GImtWCb6XApmpb5TiUxEnUTYvASGUoRDowdav3fZwvS6t7/07qWvmrNuzfVl218WgKO/CFyhBQ986pM7l42MUlbpSOx2gD0xB6KnOwkdpl2PW22x8mMFILFrtpwMmMiGemSdf20f4TqOmRy2XciQ9x2IugYptvhO1+9yACnyF+BCzyE+sWf00KHW761YmA6sXXr3+a9ZdOsaKQ//896hlm/B8/s4FpuInHbvvdu/f/eGp48bG+/MK9VU5eKY/sZsnKwXHuLce61tCEc70OdyNM2vVxIsxUy8dlDwsTaHm7WYtF+4EBEFBTYpvbYwaG0st9kGjTivgcTl/fkxQk8j8wCGJajs69JWWJQwY3T0QHb+5YuSN7zznLuWrZxf3vxlB3DsYAKW0PKEiFxeq+WfG/zC/nP2jeRZUk2SPM/87MxtakAd5QkUQzraq3m7K3Cb23BkC25CP4OVGQk8QTiME9GK0hqUBrlw9I3lmW8asmaf46Oj2SuuOCm58bYldy191bzy8JcF4JgrAtqOAz8RkSs6quqBez+775y9IzqrVpMkz3MzAdjUIXMuGGy5/MYlKB4LYkWfvXnRrswz6z0VmYCI992zdJ4o/tunG/nvY3MKfKAIF53D22zCYwdgp3ZmVjg0Otp6xarF6Rvffc5dZ108uzz8ZQE4djEByxPYKyJXVBP6xmc+s2fZ3r2UdXYmSdbKIZIYRD9y5zIFQDzk5mLAnENQ4PO5IVxAUIFmrEMiUFgzUpsENxx50bln9RmiQssKjGwnoQEgD9HoEf9fvEOBcfIZOXSodcGrF6Y3vnPJXS+9pDz8ZQE4xp+INrxXRFZRhb/xqb/fsWzv7mrW3VVNcsmRi3a4mh0BnNeGeK49SZTpB2s4IlRk5koRNNRi7UsjJqBIMerb03cjOaJPM9LBSxA2oiy2/mKfYiRgpTB66GBr+WsWpTfcdvbd57xy3rrrdXn4ywJQPtO2A6TxjU/9w46l+3bnWVdXNSEhZFpBdI6E2C/YlbPUtlRabTy7vN8eu5Kh2Tr/BCwhOP6YzYBEFuUCMmEefrFvtwXaxZA50U4OZ3lmjjAHxwNvcGRIPqNjB/PzVs5NB9YuuXPZK+ff8d+08B2C8vC/EMbR8i14YTzRduC4f9y855v3/MPOpTu2I+vsriSt3Ny2yvLy3YpQLBJvtngC1jZC3I4AJBqwBUBTDm3pxE5zb9iBuR8JjNEIWY8B6xxE2mwBdAhDNb4GOVwegikepvtgHTwFFRMOjh5oXXLVCemN6866+8yL5t7aj0E1KOXNX3YA5dMODOaD4juBV1dYffWTf7/95b/a3sw6uqqJOO2Anc91lAlgDrNV34EN+ce5bzgADs5nsN2DP9B72VF5bDHwhKRI9ON/hyUJmMKhLWtQQGTCx5gYo+Oj+aVXnZi+4fald73kglnrhtYPJSvuWFEy/MoCUD4zjgNhRbhXRC4+eGDqg//8T4dufeIXzVZPd0eqJRgBGksQC7J5jz+2smHL33cyOy8A8sb/phhoAtyBjU1BNQCtva4/5Pq1uw0rWzhyeDWAVf5MjI9nyy6ela55+1l//ZILZv3lmusHVXn4yxGgfH69cYCJSJSCfO5T2z/8z5tHb932WN6q9lRSrcXHdOnI3FOssxC5oE4x3n4sEtGFHa8gN229HQVENJQDGrU2ph2wRcCyfdzPNN2GtoXFgIdKAjcBgIyPH8wvvnJRsubtp39oySVz3vmqVw0l3/pWefhfiE/pCPTCHAe0iOC66wbV69980rqrru29+8yzK2lrqtWqsHPrDdZaZN12mMm6+BhRkCIyYB8lwd4LYgw4hG0REOvWY4M+WEUAIHnQj71FmJUek7JjheEH5hCIIpmYPIRzL+tN+tee8b6ll8155/r15eEvO4Dy+U07ARqgjfwlNZDfd++Ou76+aXTtY8ONVldPZ5JrFzgGPwL4FZ3LFLSzufHny31MGSybjy1g6Kx7zWrRFAbJdcHsQ1k6sokitwQhF0fOBBLIRG1MLlzRSzfcctafn33ZvL/p7x9UGzcOlP59JQZQPr9hJyAiogdoUF134wm3fuVLu1CpjK195EdTurOzB1pn5GZziVp0ZmVCR6zwx8tznfOQv4w52ibEq8CQX6invSZbIMzrMyYkmehaY4wuuXI2rf6Tl/affemcL61fP5TccceKnMorpuwAyuff3wncQXfQ+/hO/b2hA+///Cd/+ac//WFOlbRHKyIV+LjOJIQ8+Kch0KLhHIg8P8A5AmmynkJSkBaTBGUfXGfhosNs0VAMNJp5RmoqWbl6ntzwtrPXLDyz477164eSO+9cmZV/c2UHUD7PUSdAIFlz/aBa/qq5/2n4R5P/suWLT2/8zgMTqlnryjo6U6VFEzvjvRyOz2tWg8JOKghH8WNr4SWWCagcsFj8udPdheDdf2VyalIvPDlNLnrNosff+t5lb6Mu+vbQ+qFkZXn4yw6gfH47z9B6SVbeSdnUmFz6uY8/8aEHv9U491dPalSSjixNSIkIic6h7Oye2w6AbVegJfdJv9qSfIRgJMLakoLsfB8+JKY7MLc+S6PeyolqyennVvEHr3/J1y5bs/BmIto+NCTJypVUHv6yAJTPb/MZHBQ1MEC5iFS/+42Ddz7wT7+6ddujlVkH9uRQnGZpKoq1GMmQuEQgRGo/F+Ft+ARuejDhJVYGpMWyBL37kG41cp3lteTEUztwzsVdT7/2ujP/7NTzq1+AADIoigbKwI6yAJTP81oELEaw5P7P7nj3g0O7bti/c3bPvh0NKKR5kjAS0maDp3UACkUM049D9LaxFTf7fQ5bAZ3luTSaTahEq4XHV3HKEoyuuPKUjy2/esH/IqJd/f2lqKcsAOXzO3lEhDZuBEeF4JwtX9i79uH/t++GvU8nc/bv0hg72AJy0gklOmEGE0hDwC6Nw5J6NEigRfJWJjrL0cobrFLirl7CohclOOm0ZNvLLzr+/kuvO+5/E9H29iJUPkfmU4KAR3L1Nrduvn69MPAtJqKHAawVkff9+MGxVY8/cvDa3dtrK/Zsb/VNjIAnDwlaTUA0Y2pyClobqy8tgmq1Qqoi6JkN9M4B+uZ3oHdBsu30pX0/ftkFJ9x/4kvVF4loyh38/n7oMqOv7ADK5wX0rF8vfPZPQQMbw8EUkeNGd2Yv++kP9536b/+2+7g5c7ovHtndSju66IJWIyNApKtLUdbix/v6sK9zVvdTp51W+dFZ5y5+Ej34LhHV/NjRP6j6y3a/fMrnhV8IhoYk6cegmvEXmNl/kYgcLyKL7f/OeBn0Y1ANDUlyuK+Xz5H9/H+0xr3x+pX7XwAAAABJRU5ErkJggg=="


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=True)
